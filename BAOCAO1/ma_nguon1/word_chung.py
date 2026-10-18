"""Hàm dựng file Word dùng chung (python-docx): kiểu chữ, bảng, hình, chú thích."""
from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
INK = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x52, 0x51, 0x4E)
ACCENT = RGBColor(0x1C, 0x5C, 0xAB)


def _font(run, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(k), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


class Doc:
    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        st = self.d.styles
        for name, size, bold in (("Normal", 12, False), ("Heading 1", 15, True),
                                 ("Heading 2", 13, True), ("Heading 3", 12, True),
                                 ("Title", 20, True)):
            s = st[name]
            s.font.name = FONT
            s.font.size = Pt(size)
            s.font.bold = bold
            s.font.color.rgb = INK if name == "Normal" else ACCENT
            rpr = s.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(k), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            pf = s.paragraph_format
            if name.startswith("Heading"):
                pf.space_before = Pt(14 if name == "Heading 1" else 10)
                pf.space_after = Pt(6)
                pf.keep_with_next = True
        st["Normal"].paragraph_format.space_after = Pt(6)
        st["Normal"].paragraph_format.line_spacing = 1.2
        self.n_hinh = 0
        self.n_bang = 0
        self._footer()

    def _footer(self):
        p = self.d.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            r._element.append(el)
        _font(r, 10, color=MUTED)

    # ----------------------------------------------------------- văn bản
    def h(self, text, level=1):
        return self.d.add_heading(text, level)

    def p(self, *parts, size=None, align=None, after=None, italic=False):
        """Đoạn văn; mỗi phần là chuỗi hoặc (chuỗi, 'b'|'i'|'bi')."""
        par = self.d.add_paragraph()
        for part in parts:
            txt, fmt = (part, "") if isinstance(part, str) else part
            r = par.add_run(txt)
            _font(r, size, bold="b" in fmt or None, italic=("i" in fmt) or italic or None)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if after is not None:
            par.paragraph_format.space_after = Pt(after)
        return par

    def bullets(self, items, style="List Bullet"):
        for it in items:
            par = self.d.add_paragraph(style=style)
            parts = it if isinstance(it, (list, tuple)) else [it]
            for part in parts:
                txt, fmt = (part, "") if isinstance(part, str) else part
                r = par.add_run(txt)
                _font(r, None, bold="b" in fmt or None, italic="i" in fmt or None)
            par.paragraph_format.space_after = Pt(3)

    def cong_thuc(self, text):
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(text)
        _font(r, 12, italic=True, name="Cambria Math")
        par.paragraph_format.space_after = Pt(6)

    def ngat_trang(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ----------------------------------------------------------- hình, bảng
    def hinh(self, path, caption, width_cm=16.4):
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(path), width=Cm(width_cm))
        self.n_hinh += 1
        cp = self.d.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        _font(r, 10.5, italic=True, color=MUTED)
        cp.paragraph_format.space_after = Pt(10)

    def bang(self, header, rows, widths_cm, caption=None, size=9.5, bold_row=None):
        if caption:
            self.n_bang += 1
            cp = self.d.add_paragraph()
            cp.paragraph_format.keep_with_next = True
            r = cp.add_run(f"Bảng {self.n_bang}. {caption}")
            _font(r, 10.5, bold=True)
            cp.paragraph_format.space_after = Pt(3)
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, txt in enumerate(header):
            c = t.rows[0].cells[i]
            c.width = Cm(widths_cm[i])
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            r = c.paragraphs[0].add_run(str(txt))
            _font(r, size, bold=True)
            sh = OxmlElement("w:shd")
            sh.set(qn("w:val"), "clear")
            sh.set(qn("w:color"), "auto")
            sh.set(qn("w:fill"), "E8EEF7")
            c._tc.get_or_add_tcPr().append(sh)
        for k, row in enumerate(rows):
            cells = t.add_row().cells
            for i, txt in enumerate(row):
                cells[i].width = Cm(widths_cm[i])
                par = cells[i].paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                if i > 0:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if _is_num(txt) else WD_ALIGN_PARAGRAPH.LEFT
                r = par.add_run(str(txt))
                _font(r, size, bold=(bold_row is not None and k == bold_row) or None)
        # lặp hàng tiêu đề khi bảng sang trang
        trpr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def hop(self, lines, fill="F2F6FC"):
        """Khung nổi bật (1 ô) cho tóm tắt kết quả."""
        t = self.d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        c = t.rows[0].cells[0]
        c.width = Cm(16.4)
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), fill)
        c._tc.get_or_add_tcPr().append(sh)
        first = True
        for parts in lines:
            par = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            par.paragraph_format.space_after = Pt(3)
            for part in (parts if isinstance(parts, (list, tuple)) else [parts]):
                txt, fmt = (part, "") if isinstance(part, str) else part
                r = par.add_run(txt)
                _font(r, 11, bold="b" in fmt or None, italic="i" in fmt or None)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def save(self, path):
        self.d.core_properties.title = "HUIM cho định vị trong nhà trên UJIIndoorLoc"
        self.d.core_properties.language = "vi-VN"
        # mẫu mặc định của python-docx có <w:zoom> thiếu w:percent (bắt buộc theo schema)
        for z in self.d.settings.element.findall(qn("w:zoom")):
            if z.get(qn("w:percent")) is None:
                z.set(qn("w:percent"), "100")
        self.d.save(str(path))


def _is_num(s) -> bool:
    s = str(s).replace(".", "").replace(",", "").replace("%", "").replace("−", "").replace("-", "").strip()
    return s.isdigit()
