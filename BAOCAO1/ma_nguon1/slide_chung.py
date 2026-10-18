"""Hàm dựng slide dùng chung cho BAOCAO1 (python-pptx, 16:9): chữ, gạch đầu dòng, thẻ số,
hình, biểu đồ cột / đường gốc của PowerPoint (sửa được số liệu trong PowerPoint)."""
from __future__ import annotations

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from hinh_chung import DIAGRAM, vn

NAVY = RGBColor(0x0D, 0x36, 0x6B)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
GRID = RGBColor(0xE1, 0xE0, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xF2, 0xF4, 0xF8)
ICE = RGBColor(0xCD, 0xE2, 0xFB)
MODEL = {"A": RGBColor(0x2A, 0x78, 0xD6), "B": RGBColor(0xEB, 0x68, 0x34),
         "C": RGBColor(0x1B, 0xAF, 0x7A)}
DEV, VAL = RGBColor(0x86, 0xB6, 0xEF), RGBColor(0x1C, 0x5C, 0xAB)
GREY = RGBColor(0x9A, 0x9A, 0x9A)
HEAD, BODY = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, Hh = 13.333, 7.5


def m(x, nd=2):
    return vn(x, nd) + " m"


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        parts = para if isinstance(para, tuple) and isinstance(para[0], tuple) else (
            (para,) if not isinstance(para, tuple) else (para,))
        for part in parts:
            if isinstance(part, tuple):
                t, fmt = part
            else:
                t, fmt = part, {}
            r = p.add_run()
            r.text = t
            r.font.name = fmt.get("font", font)
            r.font.size = Pt(fmt.get("size", size))
            r.font.bold = fmt.get("bold", bold)
            r.font.italic = fmt.get("italic", False)
            r.font.color.rgb = fmt.get("color", color)
        p.space_after = Pt(6)
    return tb


def bullets(slide, x, y, w, h, items, size=15, color=INK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.28))))
        pPr.set("indent", str(-Emu(Inches(0.22))))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        parts = it if isinstance(it, list) else [it]
        for part in parts:
            t, b = (part, False) if isinstance(part, str) else part
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = BODY, Pt(size), b
            r.font.color.rgb = color
        p.space_after = Pt(8)
    return tb


def title(slide, t, sub=None):
    text(slide, 0.6, 0.35, W - 1.2, 0.8, t, size=30, bold=True, font=HEAD, color=NAVY)
    if sub:
        text(slide, 0.6, 1.08, W - 1.2, 0.5, sub, size=15, color=INK2)


def footer(slide, src):
    text(slide, 0.6, Hh - 0.45, W - 1.2, 0.3, src, size=10, color=MUTED)


def card(slide, x, y, w, h, big, small, fill=SOFT, big_color=NAVY, small_color=INK2):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    text(slide, x + 0.2, y + 0.15, w - 0.4, h * 0.5, big, size=36, bold=True, color=big_color,
         font=HEAD, anchor=MSO_ANCHOR.BOTTOM)
    text(slide, x + 0.2, y + h * 0.55, w - 0.4, h * 0.42, small, size=13, color=small_color)


def picture(slide, name, x, y, w=None, h=None):
    p = DIAGRAM / f"{name}.png"
    return slide.shapes.add_picture(str(p), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def style_chart(ch, title_txt=None, legend=True, y_title=None, x_title=None, labels=True, fmt='0.0'):
    ch.font.name, ch.font.size = BODY, Pt(12)
    ch.font.color.rgb = INK2
    if title_txt:
        ch.has_title = True
        ch.chart_title.text_frame.text = title_txt
        tp = ch.chart_title.text_frame.paragraphs[0]
        tp.runs[0].font.size, tp.runs[0].font.bold = Pt(14), True
        tp.runs[0].font.color.rgb = INK
    else:
        ch.has_title = False
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(11)
    va = ch.value_axis
    va.minimum_scale = 0                 # cột / đường luôn bắt đầu từ 0, không phóng đại chênh lệch
    va.major_gridlines.format.line.color.rgb = GRID
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    ch.category_axis.tick_labels.font.size = Pt(11)
    ch.category_axis.format.line.color.rgb = GRID
    if y_title:
        va.has_title = True
        va.axis_title.text_frame.text = y_title
        va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    if x_title:
        ca = ch.category_axis
        ca.has_title = True
        ca.axis_title.text_frame.text = x_title
        ca.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    if labels:
        pl = ch.plots[0]
        pl.has_data_labels = True
        pl.data_labels.number_format = fmt
        pl.data_labels.number_format_is_linked = False
        pl.data_labels.font.size = Pt(10)
        pl.data_labels.font.color.rgb = INK


def bar_chart(slide, x, y, w, h, cats, series, colors, **kw):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, [round(float(v), 3) for v in vals])
    gf = slide.shapes.add_chart(kw.pop("kind", XL_CHART_TYPE.COLUMN_CLUSTERED), Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
    ch = gf.chart
    for s, c in zip(ch.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = c
    ch.plots[0].gap_width = 60
    ch.plots[0].overlap = -10
    style_chart(ch, **kw)
    return ch


def line_chart(slide, x, y, w, h, cats, series, colors, dashes=None, **kw):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, [None if pd.isna(v) else round(float(v), 3) for v in vals])
    ch = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    for i, (s, c) in enumerate(zip(ch.series, colors)):
        s.format.line.color.rgb = c
        s.format.line.width = Pt(2.25)
        s.smooth = False
        s.marker.style = XL_MARKER_STYLE.CIRCLE
        s.marker.size = 7
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = c
        s.marker.format.line.color.rgb = c
        if dashes and dashes[i]:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    style_chart(ch, **kw)
    return ch
