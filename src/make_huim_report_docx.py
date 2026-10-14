"""Build report/Bao_Cao_Thuat_Toan_HUIM.docx - the academic report "Báo Cáo
Thuật Toán HUIM" (theory -> EFIM -> HUIM model for IPS -> Version 1 results ->
Version 2 improvements of the metre error), in Vietnamese.

Format: A4, Times New Roman everywhere, body text 13 pt (tables 10-11 pt so
that wide tables fit), table captions above, figure captions below, formulas
rendered as images in a Times-like math font.

Every chart is drawn here from the result files and is followed by a table of
ALL the values it plots; every number in the text is read from results/ at
build time (the reference list is the Crossref-verified bibliography of
research/papers_table.csv, typed in IEEE style).

Inputs  results/*.json, results/metrics/*.csv, results/predictions/*.csv,
        results/verification/*.json, verification/kiem_chung_thu_cong.xlsx,
        tests/test_efim.py (EFIM running example), the two dataset CSVs
Outputs report/Bao_Cao_Thuat_Toan_HUIM.docx, figures/fig30..fig36 (.png+.csv)

Run: python src/make_huim_report_docx.py
"""
from __future__ import annotations

import io
import os
import re
import subprocess
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from report_builder import Doc, _rfonts
from utils import (FIGURES, METRICS, PROJECT_ROOT, REPORT, RESULTS, find_dataset,
                   get_logger, load_json)

LOG = get_logger("huimreport", "huim_report_docx.log")
OUT = REPORT / "Bao_Cao_Thuat_Toan_HUIM.docx"
BODY = 13
FONT = "Times New Roman"

plt.rcParams.update({
    "font.family": "serif", "font.serif": [FONT, "DejaVu Serif"],
    "mathtext.fontset": "stix", "font.size": 12, "axes.titlesize": 13,
    "axes.labelsize": 12, "xtick.labelsize": 11, "ytick.labelsize": 11,
    "axes.grid": True, "grid.alpha": 0.3,
    "axes.spines.top": False, "axes.spines.right": False,
    "legend.frameon": False, "legend.fontsize": 10.5, "savefig.dpi": 220,
    "savefig.bbox": "tight"})
ACCENT, GREY, LIGHT = "#1f5fbf", "#8a8a8a", "#c8c8c8"
MODEL_COL = {"A_signal": "#1f3b73", "B_discriminative": "#8c2d04",
             "C_rare_discriminative": "#145a32"}
LETTER = {"A_signal": "A", "B_discriminative": "B", "C_rare_discriminative": "C"}


# --------------------------------------------------------------- formatting
def vn(x, d=2) -> str:
    """Vietnamese number format: 1.234,56"""
    return f"{float(x):,.{d}f}".replace(",", "\0").replace(".", ",").replace("\0", ".")


def vi(x) -> str:
    return f"{int(round(float(x))):,}".replace(",", ".")


def pc(x, d=2) -> str:
    return vn(float(x) * 100, d) + "%"


def day(ts: dict) -> str:
    """'30/05/2013 đến 20/06/2013' from a TIMESTAMP_stats block (UTC)."""
    f = lambda s: f"{s[8:10]}/{s[5:7]}/{s[:4]}"
    return f"{f(ts['min_utc'])} đến {f(ts['max_utc'])}"


def mu(p) -> str:
    return f"{float(p):g}".replace(".", ",")


# ------------------------------------------------------------------- loading
class Data:
    """Every result file the report needs, loaded once."""

    def __init__(self) -> None:
        self.S = load_json(RESULTS / "data_summary.json")
        self.hsum = load_json(RESULTS / "huim_summary.json")
        self.O = load_json(RESULTS / "reverse_thinking_observations.json")
        self.coord = load_json(RESULTS / "coordinate_system_verification.json")
        self.PO = load_json(RESULTS / "positioning_optimisation.json")
        self.H = pd.read_csv(METRICS / "huim_runs.csv")
        self.I = pd.read_csv(METRICS / "ips_results.csv")
        self.COST = pd.read_csv(METRICS / "ips_cost.csv")
        self.AB = pd.read_csv(METRICS / "positioning_ablation.csv")
        self.CN = pd.read_csv(METRICS / "positioning_conditions.csv")
        self.Pold = pd.read_csv(RESULTS / "predictions" / "val_predictions_study_knn.csv")
        self.Pnew = pd.read_csv(RESULTS / "predictions" / "val_predictions_optimised.csv")
        self.va = pd.read_csv(find_dataset()["validationData.csv"],
                              usecols=["LONGITUDE", "LATITUDE", "FLOOR", "USERID",
                                       "TIMESTAMP"])
        self.base = self.I[self.I["condition"] == "baseline_all_520"].iloc[0]
        done = self.H[(self.H["status"] == "COMPLETED") & (self.H["n_huis"] > 0)]
        self.primary = (done.sort_values("rel_minutil_pct")
                        .groupby("utility_model").head(1)
                        .sort_values("utility_model").reset_index(drop=True))
        self.totals = (self.H.groupby("utility_model")["total_db_utility"].first())

    def run(self, model, pct):
        r = self.H[(self.H["utility_model"] == model)
                   & np.isclose(self.H["rel_minutil_pct"], pct)]
        return r.iloc[0]

    def cond(self, name):
        return self.I[self.I["condition"] == name].iloc[0]


# --------------------------------------------------------------- the document
class Report(Doc):
    """report_builder.Doc with the format asked for: A4, 13 pt body."""

    def __init__(self) -> None:
        super().__init__()
        st = self.d.styles
        st["Normal"].font.size = Pt(BODY)
        pf = st["Normal"].paragraph_format
        pf.line_spacing = 1.3
        pf.space_after = Pt(6)
        for lvl, size in ((1, 16), (2, 14), (3, 13)):
            h = st[f"Heading {lvl}"]
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            h.paragraph_format.space_after = Pt(6)
        rpr = st.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        rpr.find(qn("w:sz")).set(qn("w:val"), str(BODY * 2))
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)

    # captions at 12 pt (body 13)
    def caption(self, text: str, kind: str = "Bảng") -> str:
        if kind == "Bảng":
            self._tab += 1
            num = self._tab
        else:
            self._fig += 1
            num = self._fig
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = kind == "Bảng"
        r = par.add_run(f"{kind} {num}. ")
        r.bold = True
        r.font.size = Pt(12)
        _rfonts(r)
        r = par.add_run(text)
        r.italic = True
        r.font.size = Pt(12)
        _rfonts(r)
        return f"{kind} {num}"

    def para(self, parts, align="justify"):
        """parts: a string, or [(text, style)] with style in '', 'b', 'i'."""
        if isinstance(parts, str):
            parts = [(parts, "")]
        self.rich(parts, align=align)
        # a lead-in ending in ':' or a bold-only line stays with what follows
        text = "".join(t for t, _ in parts).rstrip()
        if text.endswith(":") or text.endswith(").") and "Bảng" in text \
                or all("b" in s for _, s in parts):
            self.d.paragraphs[-1].paragraph_format.keep_with_next = True

    def toc_title(self, text: str):
        """'Mục lục' as plain bold text, so it is not listed inside the TOC."""
        par = self.d.add_paragraph()
        par.paragraph_format.space_after = Pt(8)
        r = par.add_run(text)
        r.bold = True
        r.font.size = Pt(16)
        _rfonts(r)

    def numbered(self, items, start: int = 1):
        """1. 2. 3. typed as text: python-docx's List Number style would
        continue one numbering through every list of the document."""
        for k, it in enumerate(items, start=start):
            par = self.d.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.6)
            parts = it if isinstance(it, list) else [(it, "")]
            r = par.add_run(f"{k}. ")
            _rfonts(r)
            for text, style in parts:
                r = par.add_run(text)
                r.bold = "b" in style
                r.italic = "i" in style
                _rfonts(r)

    def tbl(self, df: pd.DataFrame, caption: str, widths, size: float = 11):
        assert len(widths) == len(df.columns), caption
        self.caption(caption, "Bảng")
        t = self.d.add_table(rows=1, cols=len(df.columns))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False                      # fixed layout: widths are kept
        tblPr = t._tbl.tblPr
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), "62")             # 0,11 cm instead of 0,19 cm
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        tblPr.append(mar)
        hdr = t.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)
        for j, c in enumerate(df.columns):
            self._cell(hdr.cells[j], str(c), size, bold=True, fill="D9E2F3")
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for j, c in enumerate(df.columns):
                v = row[c]
                self._cell(cells[j], "" if pd.isna(v) else str(v), size,
                           left=j == 0)
        for j, w in enumerate(widths):
            t.columns[j].width = Cm(w)
        # short tables move as one block; every row is unsplittable
        keep_all = len(t.rows) <= 14
        for k, row in enumerate(t.rows):
            el = OxmlElement("w:cantSplit")
            el.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(el)
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
                if (keep_all and k < len(t.rows) - 1) or k == 0:
                    for par in row.cells[j].paragraphs:
                        par.paragraph_format.keep_with_next = True
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(2)

    def _cell(self, cell, text, size, bold=False, fill=None, left=False):
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        _rfonts(r)
        if fill:
            tcPr = cell._element.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill)
            tcPr.append(shd)

    def fig(self, path: Path, caption: str, width_cm: float = 15.5):
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(path), width=Cm(width_cm))
        self.caption(caption, "Hình")

    def eq(self, tex: str, fontsize: float = 15):
        """A display formula drawn by matplotlib mathtext (STIX, Times-like)."""
        f = plt.figure(figsize=(0.01, 0.01))
        f.text(0, 0, f"${tex}$", fontsize=fontsize)
        buf = io.BytesIO()
        f.savefig(buf, dpi=300, bbox_inches="tight", pad_inches=0.06)
        plt.close(f)
        from PIL import Image
        w_px = Image.open(io.BytesIO(buf.getvalue())).size[0]
        width_in = min(w_px / 300 * BODY / fontsize, 6.2)
        buf.seek(0)
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(buf, width=Inches(width_in))


# ------------------------------------------------------------------- charts
def _save(fig, name, data: pd.DataFrame) -> Path:
    p = FIGURES / f"{name}.png"
    fig.savefig(p)
    plt.close(fig)
    data.to_csv(FIGURES / f"{name}.csv", index=False)
    return p


def chart_mining(D: Data) -> Path:
    H = D.H.copy()
    fig, axes = plt.subplots(2, 1, figsize=(7.4, 7.6))
    for model, g in H.groupby("utility_model"):
        g = g.sort_values("rel_minutil_pct")
        c = MODEL_COL[model]
        done = g["status"] == "COMPLETED"
        for ax, col in ((axes[0], "n_huis"), (axes[1], "unique_waps_in_huis")):
            ax.plot(g["rel_minutil_pct"], g[col], color=c, lw=1.3,
                    label=f"Mô hình {LETTER[model]}")
            ax.scatter(g.loc[done, "rel_minutil_pct"], g.loc[done, col], color=c, s=22,
                       zorder=3)
            ax.scatter(g.loc[~done, "rel_minutil_pct"], g.loc[~done, col], s=46,
                       facecolors="white", edgecolors=c, lw=1.4, zorder=4)
    ticks = [20, 10, 5, 3, 2, 1.5]
    for ax in axes:
        ax.set_xscale("log")
        ax.xaxis.set_major_locator(matplotlib.ticker.FixedLocator(ticks))
        ax.xaxis.set_major_formatter(matplotlib.ticker.FixedFormatter([mu(t) for t in ticks]))
        ax.xaxis.set_minor_locator(matplotlib.ticker.NullLocator())
        ax.set_xlim(22, 1.4)
        ax.set_xlabel("Ngưỡng minUtil (% tổng utility, trục log, giảm dần sang phải)")
    axes[0].set_yscale("symlog", linthresh=1)
    axes[0].set_ylabel("Số HUI (trục symlog)")
    axes[0].set_title("(a) Số HUI bùng nổ khi hạ minUtil")
    axes[1].set_ylabel("Số WAP trong hợp các HUI")
    axes[1].set_title("(b) Số WAP dừng ở 79–91 với các lần chạy xong")
    axes[0].scatter([], [], s=46, facecolors="white", edgecolors=GREY,
                    label="Điểm rỗng: dừng do giới hạn")
    axes[0].legend(loc="upper left")
    fig.tight_layout(h_pad=1.6)
    return _save(fig, "fig30_minutil_hui_wap",
                 H[["utility_model", "rel_minutil_pct", "abs_minutil", "status",
                    "n_huis", "unique_waps_in_huis", "runtime_s",
                    "peak_memory_delta_mb"]])


def matched_budgets(D: Data) -> pd.DataFrame:
    rows = [r for r in D.O["observation_5_huim_vs_controls"]["matched_budget_comparison"]
            if "huim_A_signal" in r]
    return pd.DataFrame(rows)


def chart_topn(D: Data) -> Path:
    M = matched_budgets(D)
    fig, ax = plt.subplots(figsize=(7.4, 5.4))
    ax.fill_between(M["budget_n"], M["control_random_min"] * 100,
                    M["control_random_max"] * 100, color=LIGHT, alpha=0.45, lw=0,
                    label="Ngẫu nhiên: khoảng min–max của 3 seed")
    ax.plot(M["budget_n"], M["control_random_mean"] * 100, color=GREY, ls=":",
            lw=1.3, label="Ngẫu nhiên (trung bình 3 seed)")
    ax.plot(M["budget_n"], M["control_top_support"] * 100, color="#555555", ls="--",
            lw=1.3, marker="^", ms=4, label="Top-n theo số lần phát hiện")
    ax.plot(M["budget_n"], M["control_top_mi"] * 100, color="black", ls="--",
            lw=1.3, marker="s", ms=4, label="Top-n theo mutual information")
    for model in ("A_signal", "B_discriminative", "C_rare_discriminative"):
        p = D.primary[D.primary["utility_model"] == model].iloc[0]
        ax.plot(M["budget_n"], M[f"huim_{model}"] * 100, color=MODEL_COL[model],
                lw=1.8, marker="o", ms=4.5,
                label=f"HUIM top-n, mô hình {LETTER[model]} (minUtil {mu(p['rel_minutil_pct'])}%)")
    b = float(D.base["joint_building_floor_accuracy"]) * 100
    ax.axhline(b, color="black", lw=1, ls="-.")
    ax.text(M["budget_n"].max(), b + 1.5, f"Baseline 520 WAP: {vn(b)}%", ha="right",
            fontsize=10.5)
    ax.set_xlabel("Số WAP được chọn")
    ax.set_ylabel("Độ chính xác joint (%)")
    ax.set_ylim(0, 100)
    ax.set_title("HUIM top-n thua top-n theo MI ở mọi ngân sách từ 30 WAP")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.14), ncol=2, fontsize=10)
    return _save(fig, "fig31_huim_topn_vs_controls", M)


def union_vs_controls(D: Data) -> pd.DataFrame:
    rows = []
    for _, p in D.primary.iterrows():
        u = D.cond(f"huim_union__{p['run_tag']}")
        n = int(u["n_waps_selected"])
        rd = D.I[(D.I["family"] == "control_random") & (D.I["budget_n"] == n)]
        rows.append({"model": p["utility_model"], "minutil_pct": p["rel_minutil_pct"],
                     "n_waps": n,
                     "huim_union": u["joint_building_floor_accuracy"],
                     "top_mi": D.cond(f"control_top_mi__n{n}")["joint_building_floor_accuracy"],
                     "top_support": D.cond(f"control_top_support__n{n}")["joint_building_floor_accuracy"],
                     "random_mean": rd["joint_building_floor_accuracy"].mean()})
    return pd.DataFrame(rows)


def chart_union(D: Data) -> Path:
    U = union_vs_controls(D)
    fig, ax = plt.subplots(figsize=(7.4, 4.8))
    series = [("huim_union", "HUIM union", ACCENT, 1.0),
              ("top_mi", "Top-n theo MI", GREY, 1.0),
              ("top_support", "Top-n theo số lần phát hiện", GREY, 0.65),
              ("random_mean", "Ngẫu nhiên (TB 3 seed)", GREY, 0.35)]
    x = np.arange(len(U))
    w = 0.2
    for k, (col, lab, c, a) in enumerate(series):
        xs = x + (k - 1.5) * w
        ax.bar(xs, U[col] * 100, w, color=c, alpha=a, label=lab)
        for xi, v in zip(xs, U[col] * 100):
            ax.text(xi, v + 1, vn(v, 1), ha="center", fontsize=9.5,
                    fontweight="bold" if col == "huim_union" else "normal")
    b = float(D.base["joint_building_floor_accuracy"]) * 100
    ax.axhline(b, color="black", lw=1, ls="--")
    ax.text(len(U) - 0.5, b + 1.5, f"Baseline 520 WAP: {vn(b)}%", ha="right",
            fontsize=10.5)
    ax.set_xticks(x)
    ax.set_xticklabels([f"Mô hình {LETTER[m]} · {n} WAP\nminUtil {mu(p)}%"
                        for m, n, p in zip(U["model"], U["n_waps"], U["minutil_pct"])])
    ax.set_ylim(0, 100)
    ax.set_ylabel("Độ chính xác joint (%)")
    ax.set_title("HUIM union thua top-n theo MI ở cả ba mô hình")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.17), ncol=2, fontsize=10)
    ax.grid(axis="x", alpha=0)
    return _save(fig, "fig32_union_vs_controls", U)


def chart_cost(D: Data) -> Path:
    C = D.COST.merge(D.I[["condition", "utility_model", "rel_minutil_pct"]],
                     on="condition", how="left").sort_values("n_features")
    fig, axes = plt.subplots(2, 1, figsize=(7.4, 7.6))
    ax = axes[0]
    ax.plot(C["n_features"], C["predict_ms_per_fingerprint"], color=LIGHT, lw=1, zorder=1)
    for fam, lab, mk in (("baseline", "Baseline", "*"), ("control_top_mi", "Top-n theo MI", "s"),
                         ("control_top_support", "Top-n theo số lần phát hiện", "^")):
        g = C[C["family"] == fam]
        ax.scatter(g["n_features"], g["predict_ms_per_fingerprint"], marker=mk, s=34,
                   color="black" if fam == "baseline" else GREY, label=lab, zorder=3)
    g = C[C["family"].str.startswith("huim")]
    ax.scatter(g["n_features"], g["predict_ms_per_fingerprint"], s=40, color=ACCENT,
               label="Cấu hình HUIM (union / top-n)", zorder=4)
    ax.set_xlabel("Số WAP")
    ax.set_ylabel("ms / fingerprint (1 lõi)")
    ax.set_title("(a) Thời gian định vị tăng gần tuyến tính theo số WAP")
    ax.legend(fontsize=10, loc="upper left")
    ax = axes[1]
    ax.plot(C["n_features"], C["db_memory_float64_mb"], color="black", marker="o", ms=3,
            lw=1.2, label="float64 (bản cài đặt)")
    ax.plot(C["n_features"], C["db_memory_float32_mb"], color=GREY, marker="s", ms=3,
            lw=1.2, ls="--", label="float32")
    ax.plot(C["n_features"], C["db_memory_int8_mb"], color=LIGHT, marker="^", ms=3,
            lw=1.2, ls=":", label="int8")
    ax.set_xlabel("Số WAP")
    ax.set_ylabel("CSDL vân tay (MB)")
    ax.set_title("(b) Bộ nhớ tỉ lệ thuận với số WAP")
    ax.legend(fontsize=10, loc="upper left")
    fig.tight_layout(h_pad=1.6)
    return _save(fig, "fig33_ips_cost",
                 C[["condition", "family", "utility_model", "rel_minutil_pct",
                    "n_features", "predict_ms_per_fingerprint",
                    "db_memory_float64_mb", "db_memory_float32_mb",
                    "db_memory_int8_mb", "joint_building_floor_accuracy"]])


def chart_stages(D: Data) -> Path:
    A = D.AB
    fig, ax = plt.subplots(figsize=(7.4, 4.6))
    x = np.arange(len(A))
    w = 0.36
    ax.bar(x - w / 2, A["dev_mean_ground_m"], w, color=LIGHT, label="Dev (fold giữ lại người đo, dùng để CHỌN)")
    ax.bar(x + w / 2, A["val_mean_ground_m"], w, color=ACCENT, label="Kiểm định (validationData, chỉ để BÁO CÁO)")
    for xi, v in zip(x - w / 2, A["dev_mean_ground_m"]):
        ax.text(xi, v + 0.2, vn(v), ha="center", fontsize=10)
    for xi, v in zip(x + w / 2, A["val_mean_ground_m"]):
        ax.text(xi, v + 0.2, vn(v), ha="center", fontsize=10, fontweight="bold")
    ax.axhspan(7, 8, color="#2e7d32", alpha=0.12, lw=0)
    ax.text(len(A) - 0.45, 7.5, "mức tham chiếu 7–8 m", va="center", ha="right",
            fontsize=10, color="#2e7d32")
    ax.grid(axis="x", alpha=0)
    ax.set_xticks(x)
    ax.set_xticklabels(A["stage"] + ["\nk-NN V1", "\nkhoảng cách", "\n+ k, 1/d^p",
                                     "\n+ dò lệch", "\n+ ensemble"])
    ax.set_ylabel("Sai số trung bình (m thực địa)")
    ax.set_title("Phần lớn mức giảm sai số đến từ bước S1–S2")
    ax.legend(loc="upper right", fontsize=10)
    ax.set_ylim(0, float(A[["dev_mean_ground_m", "val_mean_ground_m"]].max().max()) * 1.2)
    return _save(fig, "fig34_v2_stages", A)


def cdf_table(D: Data) -> pd.DataFrame:
    rows = []
    for name, P in (("Version 1", D.Pold), ("Version 2", D.Pnew)):
        e = P["err_ground_m"].to_numpy()
        r = {"version": name, "mean": e.mean()}
        for q in (10, 25, 50, 75, 90, 95):
            r[f"p{q}"] = np.percentile(e, q)
        for t in (2, 5, 8, 10):
            r[f"le{t}"] = (e <= t).mean()
        r["max"] = e.max()
        rows.append(r)
    return pd.DataFrame(rows)


def chart_cdf(D: Data) -> Path:
    fig, ax = plt.subplots(figsize=(7.4, 4.6))
    for name, P, c in (("Version 1 (k-NN, k = 1)", D.Pold, GREY),
                       ("Version 2 (mô hình được chọn)", D.Pnew, ACCENT)):
        e = np.sort(P["err_ground_m"].to_numpy())
        ax.plot(e, np.arange(1, len(e) + 1) / len(e) * 100, color=c, lw=1.8,
                label=f"{name}: TB {vn(e.mean())} m, trung vị {vn(np.median(e))} m")
    ax.axvspan(7, 8, color="#2e7d32", alpha=0.12, lw=0)
    ax.text(7.5, 4, "7–8 m", ha="center", fontsize=10, color="#2e7d32")
    ax.set_xlim(0, 40)
    ax.set_ylim(0, 100)
    ax.set_xlabel("Sai số định vị (m thực địa)")
    ax.set_ylabel("Tỉ lệ fingerprint có sai số ≤ x (%)")
    ax.set_title("Version 2 dịch cả phân bố sai số sang trái")
    ax.legend(loc="lower right", fontsize=10.5)
    data = pd.DataFrame({"err_v1_sorted": np.sort(D.Pold["err_ground_m"]),
                         "err_v2_sorted": np.sort(D.Pnew["err_ground_m"])})
    return _save(fig, "fig35_v2_error_cdf", data)


SET_ORDER = ["baseline_all_520", "baseline_detected_465", "control_top_mi__n200",
             "control_top_support__n200", "control_top_mi__n91",
             "control_top_support__n91", "huim_union__C_rare_discriminative__1.7321pct",
             "huim_union__A_signal__1.9294pct", "huim_union__B_discriminative__2pct"]


def set_label(D: Data, r) -> str:
    c = r["condition"]
    if c == "baseline_all_520":
        return "Tất cả 520 WAP"
    if c == "baseline_detected_465":
        return "465 WAP từng phát hiện"
    if c.startswith("control"):
        kind = "MI" if "_mi_" in c else "số lần phát hiện"
        return f"Top-{int(r['n_waps'])} theo {kind}"
    kind = "union" if c.startswith("huim_union") else "top-n"
    return (f"HUIM {kind} {LETTER[r['utility_model']]} · {int(r['n_waps'])} WAP · "
            f"minUtil {mu(r['rel_minutil_pct'])}%")


def chart_sets(D: Data) -> Path:
    C = D.CN.set_index("condition").loc[SET_ORDER].reset_index()
    labels = [set_label(D, r) for _, r in C.iterrows()]
    fig, ax = plt.subplots(figsize=(7.6, 5.4))
    y = np.arange(len(C))[::-1]
    ax.axvspan(7, 8, color="#2e7d32", alpha=0.12, lw=0)
    for yi, a, b in zip(y, C["study_mean_ground_m"], C["opt_mean_ground_m"]):
        ax.plot([b, a], [yi, yi], color=LIGHT, lw=2, zorder=1)
        ax.scatter(a, yi, color=GREY, s=50, zorder=3)
        ax.scatter(b, yi, color=ACCENT, s=50, zorder=4)
        ax.text(a + 0.5, yi, vn(a), va="center", fontsize=10, color="#555555")
        ax.text(b - 0.5, yi, vn(b), va="center", ha="right", fontsize=10,
                fontweight="bold")
    ax.scatter([], [], color=GREY, s=50, label="Version 1 (k-NN, k = 1)")
    ax.scatter([], [], color=ACCENT, s=50, label="Version 2")
    ax.set_yticks(y)
    ax.set_yticklabels(labels, fontsize=10.5)
    ax.set_xlim(0, float(C["study_mean_ground_m"].max()) * 1.15)
    ax.set_xlabel("Sai số trung bình (m thực địa); dải xanh = 7–8 m")
    ax.set_title("Version 2 giảm sai số mọi tập WAP; tập HUIM vẫn trên 17 m")
    ax.legend(loc="upper right", fontsize=10.5)
    ax.grid(axis="y", alpha=0)
    return _save(fig, "fig36_v2_by_wap_set", C)


# ------------------------------------------------------------------- content
REFS = [
    'P. Bahl and V. N. Padmanabhan, "RADAR: An in-building RF-based user location and tracking system," in Proc. IEEE INFOCOM 2000, vol. 2, pp. 775–784. doi:10.1109/INFCOM.2000.832252',
    'M. Youssef and A. Agrawala, "The Horus WLAN location determination system," in Proc. 3rd Int. Conf. Mobile Systems, Applications, and Services (MobiSys), 2005, pp. 205–218. doi:10.1145/1067170.1067193',
    'A. Khalajmehrabadi, N. Gatsis, and D. Akopian, "Modern WLAN fingerprinting indoor positioning methods and deployment challenges," IEEE Commun. Surveys Tuts., vol. 19, no. 3, pp. 1974–2002, 2017. doi:10.1109/COMST.2017.2671454',
    'K. Kaemarungsi and P. Krishnamurthy, "Modeling of indoor positioning systems based on location fingerprinting," in Proc. IEEE INFOCOM 2004, vol. 2, pp. 1012–1022. doi:10.1109/INFCOM.2004.1356988',
    'Y. Chen, Q. Yang, J. Yin, and X. Chai, "Power-efficient access-point selection for indoor location estimation," IEEE Trans. Knowl. Data Eng., vol. 18, no. 7, pp. 877–888, 2006. doi:10.1109/TKDE.2006.112',
    'C. Song, J. Wang, and G. Yuan, "Hidden naive Bayes indoor fingerprinting localization based on best-discriminating AP selection," ISPRS Int. J. Geo-Inf., vol. 5, no. 10, art. 189, 2016. doi:10.3390/ijgi5100189',
    'J. Torres-Sospedra et al., "UJIIndoorLoc: A new multi-building and multi-floor database for WLAN fingerprint-based indoor localization problems," in Proc. Int. Conf. Indoor Positioning and Indoor Navigation (IPIN), 2014, pp. 261–270. doi:10.1109/IPIN.2014.7275492',
    'R. Agrawal, T. Imieliński, and A. Swami, "Mining association rules between sets of items in large databases," in Proc. ACM SIGMOD, 1993, pp. 207–216. doi:10.1145/170035.170072',
    'J. Han, J. Pei, and Y. Yin, "Mining frequent patterns without candidate generation," in Proc. ACM SIGMOD, 2000, pp. 1–12. doi:10.1145/342009.335372',
    'H. Yao, H. J. Hamilton, and C. J. Butz, "A foundational approach to mining itemset utilities from databases," in Proc. SIAM Int. Conf. Data Mining (SDM), 2004, pp. 482–486. doi:10.1137/1.9781611972740.51',
    'Y. Liu, W.-K. Liao, and A. Choudhary, "A two-phase algorithm for fast discovery of high utility itemsets," in Lecture Notes in Computer Science (PAKDD), 2005, pp. 689–695. doi:10.1007/11430919_79',
    'V. S. Tseng, C.-W. Wu, B.-E. Shie, and P. S. Yu, "UP-Growth: An efficient algorithm for high utility itemset mining," in Proc. 16th ACM SIGKDD, 2010, pp. 253–262. doi:10.1145/1835804.1835839',
    'V. S. Tseng, B.-E. Shie, C.-W. Wu, and P. S. Yu, "Efficient algorithms for mining high utility itemsets from transactional databases," IEEE Trans. Knowl. Data Eng., vol. 25, no. 8, pp. 1772–1786, 2013. doi:10.1109/TKDE.2012.59',
    'M. Liu and J. Qu, "Mining high utility itemsets without candidate generation," in Proc. 21st ACM CIKM, 2012, pp. 55–64. doi:10.1145/2396761.2396773',
    'P. Fournier-Viger, C.-W. Wu, S. Zida, and V. S. Tseng, "FHM: Faster high-utility itemset mining using estimated utility co-occurrence pruning," in Lecture Notes in Computer Science (ISMIS), 2014, pp. 83–92. doi:10.1007/978-3-319-08326-1_9',
    'S. Zida, P. Fournier-Viger, J. C.-W. Lin, C.-W. Wu, and V. S. Tseng, "EFIM: A fast and memory efficient algorithm for high-utility itemset mining," Knowl. Inf. Syst., vol. 51, no. 2, pp. 595–625, 2017. doi:10.1007/s10115-016-0986-0',
    'C. Zhang, G. Almpanidis, W. Wang, and C. Liu, "An empirical evaluation of high utility itemset mining algorithms," Expert Syst. Appl., vol. 101, pp. 91–115, 2018. doi:10.1016/j.eswa.2018.02.008',
    'W. Gan, J. C.-W. Lin, P. Fournier-Viger, H.-C. Chao, V. S. Tseng, and P. S. Yu, "A survey of utility-oriented pattern mining," IEEE Trans. Knowl. Data Eng., vol. 33, no. 4, pp. 1306–1327, 2021. doi:10.1109/TKDE.2019.2942594',
    'P. Fournier-Viger, J. C.-W. Lin, T. Truong-Chi, and R. Nkambou, "A survey of high utility itemset mining," in High-Utility Pattern Mining, Studies in Big Data, Springer, 2019, pp. 1–45. doi:10.1007/978-3-030-04921-8_1',
    'P. Fournier-Viger et al., "The SPMF open-source data mining library version 2," in Lecture Notes in Computer Science (ECML PKDD), 2016, pp. 36–40. doi:10.1007/978-3-319-46131-1_8',
    'J. Torres-Sospedra, R. Montoliu, S. Trilles, Ó. Belmonte, and J. Huerta, "Comprehensive analysis of distance and similarity measures for Wi-Fi fingerprinting indoor positioning systems," Expert Syst. Appl., vol. 42, no. 23, pp. 9263–9278, 2015. doi:10.1016/j.eswa.2015.08.013',
    'E. S. Lohan, J. Torres-Sospedra, H. Leppäkoski, P. Richter, Z. Peng, and J. Huerta, "Wi-Fi crowdsourced fingerprinting dataset for indoor positioning," Data, vol. 2, no. 4, art. 32, 2017. doi:10.3390/data2040032',
    'S. Adler, S. Schmitt, K. Wolter, and M. Kyas, "A survey of experimental evaluation in indoor localization research," in Proc. IPIN, 2015, pp. 1–10. doi:10.1109/IPIN.2015.7346749',
    'E. Jedari, Z. Wu, R. Rashidzadeh, and M. Saif, "Wi-Fi based indoor location positioning employing random forest classifier," in Proc. IPIN, 2015, pp. 1–5. doi:10.1109/IPIN.2015.7346754',
]


def bootstrap(D: Data, B: int = 10000, seed: int = 20251007) -> dict:
    eo = D.Pnew["err_ground_m"].to_numpy()
    es = D.Pold["err_ground_m"].to_numpy()
    rng = np.random.default_rng(seed)
    idx = rng.integers(0, len(eo), size=(B, len(eo)))
    ci = lambda a: np.percentile(a, [2.5, 97.5])
    return {"new": eo.mean(), "new_ci": ci(eo[idx].mean(1)),
            "old": es.mean(), "old_ci": ci(es[idx].mean(1)),
            "diff": (es - eo).mean(), "diff_ci": ci((es - eo)[idx].mean(1)),
            "better": (eo < es).mean(), "B": B}


def excel_checks() -> int | None:
    p = PROJECT_ROOT / "verification" / "kiem_chung_thu_cong.xlsx"
    if not p.exists():
        return None
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True, read_only=True)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            n += sum(v in ("KHỚP", "ĐẠT", "ĐẠT - là HUI") for v in row)
    return n or None


def build(D: Data) -> Report:
    R = Report()
    S, cs = D.S, D.S["cross_split"]
    tr = S["training"]
    ndp = tr["not_detected_probe"]
    b = D.base
    A = D.AB
    s0, sel = A[A["stage"] == "S0"].iloc[0], A.iloc[-1]
    best_val = A.loc[A["val_mean_ground_m"].idxmin()]
    BS = bootstrap(D)
    O5 = D.O["observation_5_huim_vs_controls"]["verdict_inputs"]
    O6 = D.O["observation_6_strong_vs_discriminative"]
    cc = D.coord["conclusion"]
    cand = {c["candidate"]: c.get("distance_to_uji_campus_m") for c in D.coord["candidates_tested"]}
    utm30 = next(v for k, v in cand.items() if "30N" in k)
    utm31 = next(v for k, v in cand.items() if "31N" in k)
    n_pos = len(D.va.drop_duplicates(["LONGITUDE", "LATITUDE", "FLOOR"]))
    va_days = (D.va["TIMESTAMP"].max() - D.va["TIMESTAMP"].min()) / 86400
    n_ph = len(cs["PHONEID_validation_values"])
    waps_tr = cs["waps_detected_in_train"]
    waps_tr = len(waps_tr) if isinstance(waps_tr, list) else int(waps_tr)
    cfg = D.PO["selected_config"]
    pr = {r["utility_model"]: r for _, r in D.primary.iterrows()}
    tot = D.totals
    nb = D.I[~D.I["family"].isin(["baseline"])]
    match = nb[nb["joint_building_floor_accuracy"] >= b["joint_building_floor_accuracy"]]
    smallest = match.loc[match["n_waps_selected"].idxmin()]
    cost = D.COST.set_index("condition")
    cb = cost.loc["baseline_all_520"]
    huim_cost = D.COST[D.COST["family"].str.startswith("huim")]
    hb = huim_cost.loc[huim_cost["joint_building_floor_accuracy"].idxmax()]
    cs200 = cost.loc[smallest["condition"]] if smallest["condition"] in cost.index else None
    U = union_vs_controls(D)
    rB2, rB15 = D.run("B_discriminative", 2.0), D.run("B_discriminative", 1.5)
    rA2 = D.run("A_signal", 2.0)
    n_runs = len(D.H)
    n_cond = len(D.I)
    tot_ratio = tot.max() / tot.min()

    # ------------------------------------------------------------ title page
    R.title_page([
        ("HỌC PHẦN CAO HỌC — BÀI TÌM HIỂU THUẬT TOÁN", 13, False),
        ("", 12, False), ("", 12, False),
        ("BÁO CÁO THUẬT TOÁN HUIM", 22, True),
        ("High Utility Itemset Mining và ứng dụng chọn điểm truy cập", 15, False),
        ("cho định vị trong nhà bằng vân tay Wi-Fi", 15, False),
        ("", 12, False),
        ("Thực nghiệm trên UJIIndoorLoc — Version 1 và Version 2", 13, False),
        ("", 12, False), ("", 12, False),
        ("Mọi số liệu trong báo cáo được sinh tự động từ các file kết quả "
         "của dự án (results/); không có số liệu nhập tay.", 11, False),
    ])
    R.toc_title("Mục lục")
    R.toc_field()
    R.page_break()

    # ------------------------------------------------------------ abstract
    R.h(1, "Tóm tắt")
    R.para(f"Báo cáo khảo sát khả năng dùng khai phá tập mục hữu ích cao (High "
           f"Utility Itemset Mining, HUIM) với thuật toán EFIM [16] để chọn nhóm "
           f"điểm truy cập (WAP) cho định vị trong nhà bằng vân tay Wi-Fi, thực "
           f"nghiệm trên UJIIndoorLoc [7]. Phiên bản 1 (Version 1) cho thấy tập WAP "
           f"do HUIM chọn ({U['n_waps'].min()}–{U['n_waps'].max()} WAP) không đạt "
           f"độ chính xác của mô hình dùng đủ {D.S['training']['n_wap_columns_detected']} "
           f"WAP và thua tiêu chí đối chứng top-n theo mutual information ở cả ba "
           f"mô hình utility. Phiên bản 2 (Version 2) giữ nguyên dữ liệu và giao "
           f"thức chống rò rỉ, chỉ tối ưu bộ ước lượng toạ độ; sai số trung bình "
           f"trên tập kiểm định giảm từ {vn(s0['val_mean_ground_m'])} m xuống "
           f"{vn(sel['val_mean_ground_m'])} m (khoảng tin cậy bootstrap 95%: "
           f"{vn(BS['new_ci'][0])}–{vn(BS['new_ci'][1])} m) và độ chính xác tòa "
           f"nhà–tầng tăng từ {pc(s0['val_joint_building_floor_accuracy'])} lên "
           f"{pc(sel['val_joint_building_floor_accuracy'])}.")
    R.para("Phương pháp gồm ba phần: ánh xạ mỗi fingerprint thành một giao dịch "
           "utility với q(i,T) là RSSI chuẩn hoá và ba định nghĩa trọng số ngoài "
           "p(i); cài đặt EFIM và kiểm chứng với một oracle vét cạn; so sánh có đối "
           "chứng ở cùng số WAP với ba tiêu chí chọn đơn giản. Toàn bộ siêu tham số "
           "và lựa chọn mô hình được xác định chỉ trên trainingData; validationData "
           "chỉ dùng để đánh giá.")
    R.para([("Từ khoá: ", "b"), ("high utility itemset mining; EFIM; định vị trong "
                                 "nhà; vân tay Wi-Fi; chọn điểm truy cập; k-NN; "
                                 "UJIIndoorLoc.", "")])

    # ------------------------------------------------------------ 1
    R.h(1, "1. Giới thiệu")
    R.para("Báo cáo trả lời câu hỏi: các nhóm WAP có utility cao do HUIM tìm ra có "
           "phải là tập đặc trưng tốt cho định vị hay không. Câu trả lời đo được "
           "trên UJIIndoorLoc là không, ở ngân sách tính toán đã dùng.")
    R.para([("Bối cảnh. ", "b"),
            ("Định vị bằng vân tay RF so khớp vector cường độ tín hiệu (RSSI) của "
             "thiết bị với cơ sở dữ liệu đo trước tại các vị trí biết trước; cách "
             "tiếp cận láng giềng gần nhất trong không gian tín hiệu bắt nguồn từ hệ "
             "RADAR [1], bên cạnh các bộ so khớp xác suất như Horus [2]. Các khảo "
             "sát nêu rõ hai giai đoạn offline/online và vấn đề số chiều khi số điểm "
             "truy cập lớn [3]. Độ biến thiên của RSSI khiến một số điểm truy cập "
             "mang nhiều thông tin vị trí hơn số khác [4].", "")])
    R.para([("Vấn đề chọn điểm truy cập. ", "b"),
            ("Chọn một tập con điểm truy cập để giảm chi phí tính toán đã được "
             "nghiên cứu theo tiêu chí lý thuyết thông tin [5] và theo hướng chọn "
             "điểm truy cập phân biệt tốt nhất [6]; vì vậy việc chọn điểm truy cập "
             "tự nó không phải đóng góp mới của báo cáo này.", "")])
    R.para([("Câu hỏi nghiên cứu.", "b")])
    R.numbered([
        "RQ1: Có thể mô hình hoá dữ liệu vân tay Wi-Fi thành bài toán HUIM một "
        "cách nhất quán không?",
        "RQ2: Tập WAP do HUIM chọn có cho độ chính xác cao hơn các tiêu chí chọn "
        "đơn giản ở cùng số WAP không?",
        "RQ3: Chi phí thời gian và bộ nhớ của khai phá và của hệ định vị thay đổi "
        "thế nào theo ngưỡng minUtil?",
        "RQ4 (Version 2): Sai số định vị theo mét có thể giảm bao nhiêu mà không "
        "dùng tập kiểm định để chọn mô hình?"])
    R.para([("Đóng góp và phạm vi. ", "b"),
            ("Đây là một bài tìm hiểu thuật toán trong học phần cao học, không phải "
             "luận văn. Đóng góp gồm: một ánh xạ HUIM–IPS có ba mô hình utility; một "
             "cài đặt EFIM được kiểm chứng độc lập; một đánh giá có đối chứng báo cáo "
             "cả kết quả âm; và một quy trình tối ưu sai số mét không rò rỉ dữ liệu. "
             "Báo cáo không khẳng định đây là nghiên cứu đầu tiên áp dụng HUIM cho "
             "IPS [CẦN KIỂM CHỨNG/BỔ SUNG DOI: chưa có tìm kiếm hệ thống trên "
             "Scopus/WoS về HUIM áp dụng cho Wi-Fi fingerprinting].", "")])

    # ------------------------------------------------------------ 2
    R.h(1, "2. Cơ sở lý thuyết HUIM")
    R.para("HUIM tổng quát hoá khai phá tập phổ biến bằng cách gán cho mỗi mục một "
           "lượng và một trọng số; vì utility không đơn điệu, mọi thuật toán đều phải "
           "cắt tỉa bằng cận trên.")
    R.para([("Từ tập phổ biến đến utility. ", "b"),
            ("Khai phá tập phổ biến định nghĩa mục, tập mục, giao dịch và độ hỗ trợ "
             "[8]; hướng tăng trưởng mẫu không sinh ứng viên được giới thiệu bởi "
             "FP-growth [9]. Độ hỗ trợ coi mọi mục ngang nhau; khung utility bổ sung "
             "utility nội (số lượng trong giao dịch) và utility ngoài (giá trị đơn "
             "vị) [10].", "")])
    R.para([("Định nghĩa. ", "b"),
            ("Cho cơ sở dữ liệu giao dịch D, mỗi giao dịch T gồm các mục i với lượng "
             "q(i,T) > 0 và trọng số ngoài p(i) > 0 [10], [18], [19]:", "")])
    R.eq(r"u(i,T)=q(i,T)\,p(i),\qquad u(X,T)=\sum_{i\in X}u(i,T),\qquad "
         r"u(X)=\sum_{T\in D,\;X\subseteq T}u(X,T)")
    R.para("Tập X là tập mục hữu ích cao (HUI) khi u(X) ≥ minUtil. Utility của giao "
           "dịch TU(T) là tổng u(i,T) trên mọi mục của T.")
    R.para([("Tính không đơn điệu. ", "b"),
            ("Utility không đơn điệu cũng không phản đơn điệu: thêm một mục vừa cộng "
             "utility của mục đó, vừa thu hẹp tập giao dịch chứa X [10], [18]. Do đó "
             "không thể cắt tỉa kiểu Apriori bằng chính u(X).", "")])
    R.para([("Cận trên TWU. ", "b"),
            ("Liu và cộng sự đưa ra transaction-weighted utility, phản đơn điệu và là "
             "cận trên của utility mọi tập chứa mục đó [11]:", "")])
    R.eq(r"\mathrm{TWU}(X)=\sum_{T\in D,\;X\subseteq T}\mathrm{TU}(T)\;\geq\;u(Y)"
         r"\qquad\forall\,Y\supseteq X")
    R.para("Các khảo sát phân loại thuật toán HUIM theo cấu trúc dữ liệu và chiến "
           "lược cắt tỉa [18], [19] (Bảng 1).")
    R.tbl(pd.DataFrame([
        ["Hai pha", "Two-Phase [11]", "Sinh ứng viên bằng TWU, sau đó quét lại để tính utility thật", "Nhiều ứng viên, nhiều lần quét"],
        ["Tăng trưởng mẫu", "UP-Growth, UP-Growth+ [12], [13]", "Cây UP-Tree, giảm utility ước lượng ở pha sinh ứng viên", "Vẫn cần pha kiểm tra ứng viên"],
        ["Một pha, utility-list", "HUI-Miner [14], FHM [15]", "Danh sách utility với iutil/rutil; FHM thêm cắt tỉa theo đồng xuất hiện", "Chi phí nối danh sách và bộ nhớ"],
        ["Một pha, chiếu CSDL", "EFIM [16]", "Hai cận trên mới, chiếu và gộp giao dịch, đếm utility tuyến tính", "—"],
    ], columns=["Họ", "Đại diện", "Ý tưởng chính", "Giới hạn"]),
        "Các họ thuật toán HUIM", [2.8, 3.4, 6.2, 3.6])
    R.para([("Lý do chọn EFIM. ", "b"),
            ("Căn cứ chính là đánh giá thực nghiệm độc lập của Zhang và cộng sự trên "
             "nhiều thuật toán HUIM [17], không do nhóm tác giả EFIM thực hiện. Thư "
             "viện tham chiếu SPMF [20] cài đặt EFIM bằng Java; môi trường thực nghiệm "
             "không có JVM nên EFIM được cài lại bằng Python và kiểm chứng độc lập "
             "(mục 5).", "")])

    # ------------------------------------------------------------ 3 EFIM
    sys.path.insert(0, str(PROJECT_ROOT / "tests"))
    from huim.brute_force import itemset_utility, transaction_utility, twu
    from huim.efim import mine_efim
    from test_efim import RUNNING_DB as DB, RUNNING_P as P
    mu_ex = 30
    R.h(1, "3. Thuật toán EFIM")
    R.para("EFIM duyệt sâu cây liệt kê tập mục theo thứ tự TWU tăng dần và cắt tỉa "
           "bằng hai cận trên chặt hơn TWU: local utility và subtree utility [16].")
    R.para([("Hai cận trên. ", "b"),
            ("Với tập tiền tố α, mục mở rộng z, và re(α,T) là utility còn lại của các "
             "mục đứng sau α trong T [16]:", "")])
    R.eq(r"\mathrm{lu}(\alpha,z)=\sum_{T\supseteq\alpha\cup\{z\}}"
         r"\left[u(\alpha,T)+\mathrm{re}(\alpha,T)\right]")
    R.eq(r"\mathrm{su}(\alpha,z)=\sum_{T\supseteq\alpha\cup\{z\}}"
         r"\left[u(\alpha,T)+u(z,T)+\sum_{i\in T,\;i\succ z}u(i,T)\right]")
    R.para("Mục có lu < minUtil bị loại khỏi toàn bộ cây con của α (tập Secondary). "
           "Nhánh mở rộng α thêm z chỉ được mở khi su(α, z) ≥ minUtil (tập Primary). "
           "Tại gốc (α rỗng), lu trùng với TWU(z), nên cắt tỉa TWU là trường hợp "
           "riêng [16].")
    R.para([("Các bước của thuật toán [16]:", "b")])
    R.numbered([
        "Tính TU(T) và TWU(i); loại mục có TWU < minUtil.",
        "Sắp mục theo TWU tăng dần; thứ tự này quyết định hình dạng cây tìm kiếm.",
        "Chiếu cơ sở dữ liệu (high-utility database projection): khi xuống nút "
        "α thêm z, mỗi giao dịch chỉ giữ phần sau z kèm utility tiền tố.",
        "Gộp giao dịch (high-utility transaction merging): các giao dịch chiếu giống "
        "hệt nhau được gộp, cộng utility theo từng phần tử.",
        "Fast Utility Counting: lu và su của mọi mục được tính trong một lượt quét "
        "bằng hai mảng utility-bin, tuyến tính theo kích thước cơ sở dữ liệu chiếu.",
        "Đệ quy trên các mục Primary; xuất X khi u(X) ≥ minUtil."])
    R.para("Bản cài đặt của báo cáo bổ sung hai điểm không có trong [16]: đếm số lần "
           "xuất hiện khi gộp để độ hỗ trợ vẫn chính xác, và dừng có kiểm soát khi "
           "vượt 600 giây hoặc 1.000.000 HUI (trạng thái ABORTED_RESOURCE_LIMIT).")
    R.para([("Ví dụ minh hoạ. ", "b"),
            (f"Cơ sở dữ liệu {len(DB)} giao dịch, {len(P)} mục a–g là ví dụ chạy chuẩn "
             f"trong [16]. Các giá trị trong Bảng 2–4 do bản cài đặt của báo cáo tính, "
             f"với minUtil = {mu_ex}.", "")])
    R.tbl(pd.DataFrame([[f"T{k}", ", ".join(f"{i}({q})" for i, q in t.items()),
                         vn(transaction_utility(t, P), 0)]
                        for k, t in enumerate(DB, start=1)],
                       columns=["Giao dịch", "Mục (lượng q)", "TU(T)"]),
          "Cơ sở dữ liệu ví dụ; trọng số " + ", ".join(f"p({i}) = {v}" for i, v in P.items()),
          [2.5, 8.5, 2.5])
    tw = {i: twu(i, DB, P) for i in sorted(P)}
    order = sorted((i for i in tw if tw[i] >= mu_ex), key=lambda i: (tw[i], i))
    su0 = {}
    for z, item in enumerate(order):
        su0[item] = sum(sum(q * P[i] for i, q in t.items()
                            if i in order and order.index(i) >= z)
                        for t in DB if item in t)
    R.tbl(pd.DataFrame([[i, vn(tw[i], 0), order.index(i) + 1, vn(su0[i], 0),
                         "Primary" if su0[i] >= mu_ex else "Chỉ Secondary"]
                        for i in order],
                       columns=["Mục", "TWU", "Thứ tự duyệt", "su tại gốc", "Vai trò tại gốc"]),
          "TWU, thứ tự duyệt và subtree utility tại gốc", [2.0, 2.4, 3.0, 3.0, 4.0])
    huis, st = mine_efim(DB, P, mu_ex)
    R.tbl(pd.DataFrame([["{" + ", ".join(sorted(x)) + "}", vn(u, 0), s]
                        for x, u, s in sorted(huis, key=lambda r: (len(r[0]), sorted(r[0])))],
                       columns=["HUI", "u(X)", "Độ hỗ trợ"]),
          f"Các HUI tìm được với minUtil = {mu_ex}", [6.0, 3.0, 3.0])
    R.para(f"EFIM trả về {len(huis)} HUI nhưng chỉ duyệt {st.n_nodes_visited} nút "
           f"trên {2 ** len(P) - 1} tập mục khả dĩ, loại {st.n_candidates_pruned_by_su} "
           f"nhánh nhờ subtree utility và gộp {st.n_transactions_saved_by_merging} giao "
           f"dịch. Ví dụ cũng cho thấy tính không đơn điệu: u({{b}}) = "
           f"{vn(itemset_utility(('b',), DB, P), 0)} nhưng u({{b,d}}) = "
           f"{vn(itemset_utility(('b', 'd'), DB, P), 0)}.")

    # ------------------------------------------------------------ 4 model
    R.h(1, "4. Mô hình hoá HUIM cho định vị trong nhà")
    R.para("Mỗi fingerprint trở thành một giao dịch utility: mục là WAP được phát "
           "hiện, lượng là RSSI chuẩn hoá, trọng số là mức giá trị của WAP. Dữ liệu "
           "Wi-Fi không có sẵn q và p, nên ánh xạ dưới đây là lựa chọn thiết kế của "
           "báo cáo, không phải chuẩn trong tài liệu.")
    R.tbl(pd.DataFrame([
        ["Mục i", f"Một WAP ({tr['n_wap_columns_detected']} cột trong UJIIndoorLoc [7])"],
        ["Giao dịch T", "Một fingerprint; WAP có giá trị 100 (không phát hiện) không thuộc T"],
        ["Utility nội q(i,T)", "RSSI chuẩn hoá về số nguyên 1…100"],
        ["Utility ngoài p(i)", "Một trong ba định nghĩa A, B, C (Bảng 6)"],
        ["HUI", "Nhóm WAP có tổng utility ≥ minUtil"],
    ], columns=["Khái niệm HUIM", "Ánh xạ trong IPS"]), "Ánh xạ HUIM – IPS", [4.5, 11.5])
    R.para([("Giá trị 100. ", "b"),
            (f"Tài liệu bộ dữ liệu dùng giá trị 100 để chỉ WAP không được phát hiện "
             f"[7]. Audit của báo cáo xác nhận: {vi(ndp['count_in_1_to_99'])} ô trong "
             f"{vi(ndp['total_wap_cells'])} ô WAP của trainingData mang giá trị 1…99, "
             f"và {vn(ndp['share_eq_100_pct'])}% số ô bằng 100; trung bình mỗi "
             f"fingerprint phát hiện {vn(ndp['detections_per_fingerprint']['mean'], 1)} "
             f"WAP.", "")])
    R.para([("Utility nội. ", "b"),
            ("Theo hướng biểu diễn \"positive\" trong các nghiên cứu trên UJIIndoorLoc "
             "[21], [22], mỗi RSSI phát hiện được ánh xạ tuyến tính về thang dương; "
             "RSSI_min và RSSI_max được ước lượng chỉ trên trainingData:", "")])
    R.eq(r"q(i,T)=\mathrm{round}\left(1+\dfrac{\mathrm{RSSI}_i-\mathrm{RSSI}_{\min}}"
         r"{\mathrm{RSSI}_{\max}-\mathrm{RSSI}_{\min}}\times 99\right)")
    db = D.hsum["database"]
    R.para(f"Kết quả là {vi(db['n_non_empty_transactions'])} giao dịch không rỗng từ "
           f"{vi(db['n_rows'])} fingerprint ({db['n_empty_fingerprints_dropped']} "
           f"fingerprint không phát hiện WAP nào).")
    R.para([("Ba mô hình utility ngoài. ", "b"),
            (f"Cả ba dùng chung q(i,T), chỉ khác p(i), để phép so sánh cô lập đúng một "
             f"biến. Y là nhãn tòa nhà–tầng ({tr['joint_building_floor_n_classes']} lớp); "
             f"I(A_i;Y) là mutual information giữa trạng thái WAP (không phát hiện và "
             f"5 khoảng dBm với biên −104, −90, −80, −70, −60, 0) và Y, tính từ "
             f"histogram đồng thời.", "")])
    R.tbl(pd.DataFrame([
        ["A – Signal", "1", vi(tot["A_signal"]), "Đối chứng: chỉ cường độ tín hiệu và đồng xuất hiện"],
        ["B – Discriminative", "I(A_i;Y) / max_j I(A_j;Y)", vi(tot["B_discriminative"]),
         "Mô hình chính; tương ứng tiêu chí information gain trong chọn điểm truy cập [5], [6]"],
        ["C – Rare-discriminative", "tỉ lệ với I(A_i;Y) × log(N/(1+support_i))",
         vi(tot["C_rare_discriminative"]), "Giảm trọng số WAP xuất hiện khắp nơi"],
    ], columns=["Mô hình", "p(i)", "Tổng utility CSDL", "Vai trò"]),
        "Ba mô hình utility ngoài (tổng utility đọc từ results/metrics/huim_runs.csv)",
        [3.4, 4.4, 2.8, 5.4])
    R.para(f"Trọng số p(i) có sàn 10⁻⁶ để bảo đảm utility dương theo yêu cầu của các "
           f"cận trên EFIM [16]. Hàm ước lượng p(i) báo lỗi nếu được gọi trên dữ liệu "
           f"khác trainingData. Vì tổng utility chênh {vn(tot_ratio)} lần giữa ba mô "
           f"hình, minUtil được biểu diễn theo phần trăm tổng utility; ngưỡng tuyệt "
           f"đối bằng phần trăm nhân tổng utility.")
    R.para([("Từ HUI đến tập WAP. ", "b"),
            ("Hai quy tắc được cố định trước khi có bất kỳ kết quả độ chính xác nào: "
             "union lấy mọi WAP xuất hiện trong ít nhất một HUI (minUtil là tham số "
             "duy nhất); top-n lấy n WAP có điểm đóng góp cao nhất, với mỗi HUI chia "
             "đều utility cho các thành viên:", "")])
    R.eq(r"\mathrm{contrib}(i)=\sum_{X\in\mathrm{HUIs},\;i\in X}\;\dfrac{u(X)}{|X|}")

    # ------------------------------------------------------------ 5 design
    R.h(1, "5. Thiết kế thực nghiệm (Version 1)")
    R.para("Mọi tham số được học hoặc chọn trên trainingData; validationData được "
           "đánh giá đúng một lần cho mỗi điều kiện và đóng vai tập kiểm tra.")
    R.para([("Dữ liệu. ", "b"),
            (f"UJIIndoorLoc gồm {vi(cs['train_rows'])} fingerprint huấn luyện và "
             f"{vi(cs['validation_rows'])} fingerprint kiểm định, mỗi fingerprint có "
             f"{tr['n_wap_columns_detected']} cột WAP cùng toạ độ, tòa nhà và tầng [7]. "
             f"trainingData do {len(cs['USERID_train_values'])} người đo (USERID "
             f"{min(cs['USERID_train_values'])}–{max(cs['USERID_train_values'])}); "
             f"validationData không ghi mã người đo (USERID = "
             f"{cs['USERID_validation_values'][0]} ở mọi dòng), được thu bằng "
             f"{n_ph} mã thiết bị, trong đó chỉ {len(cs['PHONEID_overlap'])} mã "
             f"xuất hiện trong trainingData, tại {vi(n_pos)} vị trí phân biệt. Theo "
             f"TIMESTAMP, trainingData được thu từ {day(tr['TIMESTAMP_stats'])}, còn "
             f"validationData từ {day(S['validation']['TIMESTAMP_stats'])}. Chỉ {waps_tr}/"
             f"{tr['n_wap_columns_detected']} WAP từng xuất hiện trong trainingData. "
             f"Kết quả vì vậy là kết quả trên thiết bị và thời điểm khác với dữ liệu "
             f"huấn luyện.", "")])
    R.tbl(pd.DataFrame([
        ["Chuẩn hoá RSSI, trọng số p(i), khai phá HUIM, chọn WAP", "trainingData"],
        ["Chọn siêu tham số k-NN (Version 1)", "Phép chia phân tầng 75/25 của trainingData"],
        ["Chọn mô hình định vị (Version 2)", "6 fold giữ lại người đo của trainingData"],
        ["Huấn luyện bộ định vị cuối", "Toàn bộ trainingData"],
        ["Đánh giá", "validationData, một lần mỗi điều kiện"],
    ], columns=["Bước", "Dữ liệu dùng"]), "Giao thức chống rò rỉ dữ liệu", [9.5, 6.5])
    R.para([("Kiểm chứng cài đặt EFIM. ", "b"),
            ("EFIM được so với một oracle vét cạn tính trực tiếp từ định nghĩa: ví dụ "
             "chuẩn ở 16 ngưỡng, 40 cơ sở dữ liệu ngẫu nhiên × 5 ngưỡng, 15 cơ sở dữ "
             "liệu thưa–rộng giống Wi-Fi, trọng số thực, và hai utility tính tay. Toàn "
             "bộ kiểm thử đạt trước khi chạy trên UJIIndoorLoc.", "")])
    R.para([("Khai phá. ", "b"),
            (f"Lưới minUtil bắt đầu từ 20% đến 1% tổng utility, hạ dần cho đến khi chạm "
             f"ngân sách (600 giây hoặc 1.000.000 HUI), sau đó chia đôi theo cấp số nhân "
             f"giữa ngưỡng chạy xong thấp nhất và ngưỡng bị dừng cao nhất; tổng cộng "
             f"{n_runs} lần khai phá. Tập HUI bị dừng giữa chừng thiên lệch theo thứ tự "
             f"duyệt sâu nên không được dùng để xếp hạng WAP.", "")])
    kcfg = D.PO.get("stages", [{}])[0].get("config", {})
    R.para([("Bộ định vị. ", "b"),
            ("k-NN là bộ so khớp tham chiếu cho vân tay Wi-Fi [1], [21]. Lưới "
             "k thuộc {1, 3, 5, 7, 9, 11}, khoảng cách Euclid hoặc Manhattan, trọng số đều "
             "hoặc theo khoảng cách được xét một lần với đủ 520 WAP; cấu hình thắng "
             f"(k = {kcfg.get('k', 1)}, {kcfg.get('distance', 'manhattan').capitalize()}, "
             "trọng số đều) được đóng băng cho mọi điều kiện. Random Forest 200 cây là "
             "bộ phân loại kiểm tra phụ theo tiền lệ [24]. Toạ độ dự đoán là trung bình "
             "toạ độ của các láng giềng.", "")])
    R.para([("Đối chứng ở cùng số WAP. ", "b"),
            (f"Ba tiêu chí: ngẫu nhiên (3 seed), top-n theo số lần phát hiện, và top-n "
             f"theo mutual information. Tiêu chí thứ ba dùng chính thước đo của Mô hình "
             f"B nhưng không khai phá tập mục, nên là đối chứng khó nhất. Tổng cộng "
             f"{n_cond} điều kiện được đánh giá.", "")])
    R.para([("Độ đo. ", "b"),
            (f"Độ chính xác tòa nhà, tầng và joint (đúng cả hai), macro F1, và sai số "
             f"định vị Euclid. Bộ dữ liệu không ghi hệ toạ độ; báo cáo xác định bằng "
             f"thực nghiệm đó là Web Mercator (EPSG:3857): nghịch đảo phép chiếu đặt tâm "
             f"dữ liệu cách khuôn viên Đại học Jaume I {vn(cc['distance_to_uji_campus_m'], 1)} m, "
             f"trong khi UTM 30N/31N lệch {vi(utm30 / 1000)} km và {vi(utm31 / 1000)} km. "
             f"Vì Web Mercator phóng đại khoảng cách 1/cos(vĩ độ) lần, sai số được báo "
             f"cáo theo mét thực địa (nhân {vn(cc['ground_distance_correction_factor'], 6)}) "
             f"và theo mét toạ độ chiếu, gồm trung vị, trung bình và các phân vị theo "
             f"khuyến nghị đánh giá thực nghiệm IPS [23].", "")])
    env = D.S.get("environment", {})
    R.para([("Môi trường. ", "b"),
            (f"{env.get('platform', '')}, Python {env.get('python_version', '')}, CPU "
             f"{env.get('cpu_model', '')}, RAM {vn(env.get('ram_total_gb', 0))} GB, "
             f"seed 20251007.", "")])

    # ------------------------------------------------------------ 6 results v1
    R.h(1, "6. Kết quả Version 1")
    pc_ = pr["C_rare_discriminative"]
    uC = D.cond(f"huim_union__{pc_['run_tag']}")
    R.para(f"HUIM chạy được trên toàn bộ dữ liệu, nhưng tập WAP lớn nhất nó cung cấp "
           f"chỉ có {int(uC['n_waps_selected'])} WAP và đạt "
           f"{pc(uC['joint_building_floor_accuracy'])} độ chính xác joint, so với "
           f"{pc(b['joint_building_floor_accuracy'])} của mô hình dùng đủ 520 WAP.")
    R.h(2, "6.1. Khai phá HUIM (RQ1, RQ3)")
    R.para(f"Hình 1 và Bảng 8 trình bày cả {n_runs} lần khai phá. Số HUI rất nhạy với "
           f"ngưỡng trong khi số WAP gần như không đổi. Với mô hình B, hạ minUtil từ "
           f"{mu(rB2['rel_minutil_pct'])}% xuống {mu(rB15['rel_minutil_pct'])}% làm số HUI "
           f"vượt giới hạn {vi(1_000_000)} và bộ nhớ tăng lên "
           f"{vn(rB15['peak_memory_delta_mb'], 1)} MB, trong khi số mục còn lại sau cắt "
           f"tỉa TWU chỉ đổi từ {int(rB2['items_after_twu_pruning'])} lên "
           f"{int(rB15['items_after_twu_pruning'])}. Ở cùng ngưỡng 2%, mô hình A (p = 1) "
           f"mất {vn(rA2['runtime_s'], 1)} giây cho {vi(rA2['n_huis'])} HUI, mô hình B "
           f"mất {vn(rB2['runtime_s'], 1)} giây cho {vi(rB2['n_huis'])} HUI: khi p đều "
           f"bằng 1, các cận trên ít phân hoá mục nên cắt tỉa kém.")
    R.fig(chart_mining(D), "Số HUI (a) và số WAP trong hợp các HUI (b) theo ngưỡng "
                           "minUtil; điểm rỗng là lần chạy dừng do giới hạn",
          width_cm=12.5)
    HT = D.H.sort_values(["utility_model", "rel_minutil_pct"], ascending=[True, False])
    R.tbl(pd.DataFrame([[LETTER[r["utility_model"]], mu(r["rel_minutil_pct"]),
                         vi(r["abs_minutil"]),
                         "Xong" if r["status"] == "COMPLETED" else "Dừng",
                         vi(r["n_huis"]), int(r["unique_waps_in_huis"]),
                         int(r["items_after_twu_pruning"]), vn(r["runtime_s"], 1),
                         vn(r["peak_memory_delta_mb"], 1)]
                        for _, r in HT.iterrows()],
                       columns=["Mô hình", "minUtil (%)", "minUtil tuyệt đối", "Trạng thái",
                                "Số HUI", "Số WAP", "Mục sau TWU", "Thời gian (s)",
                                "Bộ nhớ tăng (MB)"]),
          "Số liệu đầy đủ của Hình 1: mọi lần khai phá; \"Dừng\" = chạm giới hạn 600 "
          "giây hoặc 1.000.000 HUI (results/metrics/huim_runs.csv)",
          [1.4, 1.6, 2.3, 1.6, 2.1, 1.4, 1.7, 1.9, 2.0], size=10)
    R.h(2, "6.2. Baseline định vị")
    R.tbl(pd.DataFrame([
        ["Độ chính xác tòa nhà", pc(b["building_accuracy"])],
        ["Độ chính xác tầng", pc(b["floor_accuracy"])],
        ["Độ chính xác joint (tòa nhà + tầng)", pc(b["joint_building_floor_accuracy"])],
        ["Macro F1 (joint)", vn(b["joint_macro_f1"], 4)],
        ["Sai số trung vị (m thực địa / m toạ độ chiếu)",
         f"{vn(b['pos_median_error_ground_m'])} / {vn(b['pos_median_error'])}"],
        ["Sai số trung bình (m thực địa / m toạ độ chiếu)",
         f"{vn(b['pos_mean_error_ground_m'])} / {vn(b['pos_mean_error'])}"],
        ["Phân vị 75 / 90 / 95 (m thực địa)",
         f"{vn(b['pos_p75_error_ground_m'])} / {vn(b['pos_p90_error_ground_m'])} / "
         f"{vn(b['pos_p95_error_ground_m'])}"],
    ], columns=["Chỉ số", "Giá trị"]),
        "Baseline k-NN với đủ 520 WAP trên validationData", [10.0, 6.0])
    R.para("Sai số trung bình lớn gấp đôi trung vị do đuôi phân bố dài: một số "
           "fingerprint bị gán sai tầng hoặc sai tòa nhà.")
    R.h(2, "6.3. So với đối chứng ở cùng số WAP (RQ2)")
    R.para(f"Trên {O5['n_comparisons']} ngân sách WAP khớp nhau (Hình 2, Bảng 10), lựa "
           f"chọn HUIM tốt nhất vượt chọn ngẫu nhiên ở {O5['huim_beats_random']}/"
           f"{O5['n_comparisons']}, vượt top-n theo mutual information ở "
           f"{O5['huim_beats_top_mi']}/{O5['n_comparisons']} và vượt top-n theo số lần "
           f"phát hiện ở {O5['huim_beats_top_support']}/{O5['n_comparisons']}.")
    R.fig(chart_topn(D), "Độ chính xác joint theo số WAP: HUIM top-n so với ba đối "
                         "chứng (k-NN Version 1)")
    M = matched_budgets(D)
    R.tbl(pd.DataFrame([[int(r["budget_n"]), pc(r["huim_A_signal"]),
                         pc(r["huim_B_discriminative"]), pc(r["huim_C_rare_discriminative"]),
                         pc(r["control_top_mi"]), pc(r["control_top_support"]),
                         pc(r["control_random_mean"]),
                         f"{pc(r['control_random_min'])}–{pc(r['control_random_max'])}"]
                        for _, r in M.iterrows()],
                       columns=["Số WAP", "HUIM A", "HUIM B", "HUIM C", "Top-n MI",
                                "Top-n tần suất", "Ngẫu nhiên TB", "Ngẫu nhiên min–max"]),
          "Số liệu đầy đủ của Hình 2: độ chính xác joint ở từng ngân sách WAP "
          "(results/reverse_thinking_observations.json)",
          [1.4, 1.8, 1.8, 1.8, 1.9, 2.1, 2.0, 3.2], size=10)
    R.para("Ở chế độ union, mỗi mô hình cho một tập WAP cố định tại ngưỡng thấp nhất "
           "chạy xong (Hình 3, Bảng 11).")
    R.fig(chart_union(D), "HUIM union so với ba đối chứng ở đúng số WAP của nó")
    R.tbl(pd.DataFrame([[f"Mô hình {LETTER[r['model']]}", mu(r["minutil_pct"]),
                         int(r["n_waps"]), pc(r["huim_union"]), pc(r["top_mi"]),
                         pc(r["top_support"]), pc(r["random_mean"])]
                        for _, r in U.iterrows()],
                       columns=["Mô hình", "minUtil (%)", "Số WAP", "HUIM union",
                                "Top-n MI", "Top-n tần suất", "Ngẫu nhiên TB"]),
          "Số liệu đầy đủ của Hình 3 (results/metrics/ips_results.csv)",
          [2.6, 2.0, 1.6, 2.4, 2.4, 2.6, 2.4])
    R.para(f"Tập con nhỏ nhất chạm mức baseline có {int(smallest['n_waps_selected'])} WAP "
           f"({smallest['condition']}, joint {pc(smallest['joint_building_floor_accuracy'])}), "
           f"vượt xa trần {int(uC['n_waps_selected'])} WAP mà HUIM cung cấp. Trên "
           f"{O6['n_detected_waps_compared']} WAP được phát hiện, hạng theo số lần phát "
           f"hiện và hạng theo mutual information tương quan Spearman ρ = "
           f"{vn(O6['spearman_support_vs_mi_over_all_detected_waps'], 4)}, nên Mô hình B "
           f"trên thực tế gần như tham số hoá lại tần suất phát hiện.")
    UA = D.I[D.I["family"] == "huim_union"].sort_values(["utility_model", "rel_minutil_pct"],
                                                          ascending=[True, False])
    R.tbl(pd.DataFrame([[LETTER[r["utility_model"]], mu(r["rel_minutil_pct"]),
                         "Xong" if r["huim_status"] == "COMPLETED" else "Dừng",
                         vi(r["n_huis"]), int(r["n_waps_selected"]),
                         pc(r["building_accuracy"]), pc(r["floor_accuracy"]),
                         pc(r["joint_building_floor_accuracy"]),
                         vn(r["pos_mean_error_ground_m"])]
                        for _, r in UA.iterrows()],
                       columns=["Mô hình", "minUtil (%)", "Khai phá", "Số HUI", "Số WAP",
                                "Tòa nhà", "Tầng", "Joint", "Sai số TB (m)"]),
          "Mọi điều kiện HUIM union (k-NN Version 1; \"Dừng\" = HUI lấy từ lần khai "
          "phá chạm giới hạn)", [1.4, 1.7, 1.6, 2.1, 1.4, 1.95, 1.95, 1.95, 1.95], size=10)
    R.h(2, "6.4. Thời gian và bộ nhớ của hệ định vị (RQ3)")
    R.para(f"Đo trên một lõi, lấy trung vị 5 lần lặp (Hình 4, Bảng 13). Baseline 520 WAP "
           f"cần {vn(cb['predict_ms_per_fingerprint'], 3)} ms cho mỗi fingerprint và "
           f"{vn(cb['db_memory_float64_mb'], 1)} MB cơ sở dữ liệu vân tay (float64). Cấu "
           f"hình HUIM tốt nhất ({int(hb['n_features'])} WAP) nhanh hơn "
           f"{vn(cb['predict_ms_per_fingerprint'] / hb['predict_ms_per_fingerprint'], 1)} "
           f"lần và dùng {vn(hb['db_memory_float64_mb'] / cb['db_memory_float64_mb'] * 100, 0)}% "
           f"bộ nhớ, nhưng độ chính xác joint chỉ {pc(hb['joint_building_floor_accuracy'])}."
           + (f" Tập {int(cs200['n_features'])} WAP chạm baseline nhanh hơn "
              f"{vn(cb['predict_ms_per_fingerprint'] / cs200['predict_ms_per_fingerprint'], 1)} "
              f"lần và dùng {vn(cs200['db_memory_float64_mb'] / cb['db_memory_float64_mb'] * 100, 0)}% "
              f"bộ nhớ." if cs200 is not None else ""))
    R.fig(chart_cost(D), "Thời gian định vị (a) và kích thước cơ sở dữ liệu vân tay "
                         "(b) theo số WAP", width_cm=12.5)
    CT = D.COST.merge(D.I[["condition", "utility_model", "rel_minutil_pct"]],
                      on="condition", how="left").sort_values("n_features", ascending=False)
    R.tbl(pd.DataFrame([[r["condition"], int(r["n_features"]),
                         vn(r["predict_ms_per_fingerprint"], 4), vn(r["db_memory_float64_mb"], 2),
                         vn(r["db_memory_float32_mb"], 2), vn(r["db_memory_int8_mb"], 2),
                         pc(r["joint_building_floor_accuracy"])]
                        for _, r in CT.iterrows()],
                       columns=["Điều kiện", "Số WAP", "ms / fingerprint", "float64 (MB)",
                                "float32 (MB)", "int8 (MB)", "Joint"]),
          "Số liệu đầy đủ của Hình 4 (results/metrics/ips_cost.csv)",
          [5.8, 1.4, 2.0, 1.7, 1.7, 1.5, 1.9], size=9.5)

    # ------------------------------------------------------------ 7 version 2
    R.h(1, f"7. Version 2: các cải tiến rút sai số từ {vn(s0['val_mean_ground_m'])} m "
           f"xuống {vn(sel['val_mean_ground_m'])} m")
    R.para(f"Version 2 giữ nguyên dữ liệu, phần HUIM và giao thức chống rò rỉ, chỉ thay "
           f"bộ ước lượng toạ độ; sai số trung bình thực địa trên tập kiểm định giảm từ "
           f"{vn(s0['val_mean_ground_m'])} m xuống {vn(sel['val_mean_ground_m'])} m (giảm "
           f"{vn((1 - sel['val_mean_ground_m'] / s0['val_mean_ground_m']) * 100, 0)}%), "
           f"dưới mức tham chiếu 7–8 m mà đề tài đặt ra.")
    R.h(2, "7.1. Nguyên nhân sai số lớn của Version 1")
    R.numbered([
        [("Chọn tham số sai mục tiêu. ", "b"), ("Cấu hình k = 1, Manhattan được chọn "
          "cho bài toán phân loại tầng, không phải ước lượng toạ độ.", "")],
        [("Phép chia dev không giống tập kiểm định. ", "b"), (f"Phép chia ngẫu nhiên "
          f"75/25 đặt các fingerprint đo tại cùng một điểm tham chiếu vào cả hai phía, "
          f"nên láng giềng gần nhất gần như luôn trùng vị trí và k = 1 thắng. Tập kiểm "
          f"định được thu ở thời điểm khác, chủ yếu bằng thiết bị không có trong "
          f"trainingData ({n_ph - len(cs['PHONEID_overlap'])}/{n_ph} mã thiết bị), tại "
          f"{vi(n_pos)} vị trí phân biệt.", "")],
        [("Phương sai cao của k = 1. ", "b"), ("Toạ độ dự đoán chính là một vị trí "
          "huấn luyện, nên mỗi láng giềng sai thành một sai số trọn vẹn.", "")],
        [("Hoà khoảng cách. ", "b"), ("RSSI là số nguyên nên khoảng cách Manhattan "
          "thường bằng nhau tuyệt đối; với k = 1, việc chọn ai trong nhóm hoà quyết "
          "định trực tiếp toạ độ dự đoán.", "")],
        [("Lệch thiết bị. ", "b"), ("Khác biệt mức RSSI giữa các điện thoại làm sai "
          "khoảng cách tính trên giá trị RSSI thô [3], [4].", "")],
    ])
    R.h(2, "7.2. Giao thức chọn mô hình")
    R.para(f"trainingData được chia thành 6 fold theo mã người đo (GroupKFold(6) trên "
           f"USERID): mỗi fold dev chỉ gồm người đo mà mô hình chưa thấy. Tiêu chí chọn là sai số trung bình thực địa trên các fold "
           f"dev. Năm bước S0–S4 nối tiếp nhau; mỗi bước giữ cấu hình tốt nhất trên dev "
           f"và chỉ được nhận khi giảm sai số dev ít nhất {vn(D.PO['protocol']['min_gain_m'])} m. "
           f"validationData chỉ chấm điểm mỗi bước một lần, sau khi lựa chọn đã xong.")
    R.h(2, "7.3. Các cải tiến")
    g = D.PO["protocol"]["grids"]
    R.numbered([
        [("Biểu diễn tín hiệu và hàm khoảng cách. ", "b"),
         (f"Lưới {len(g['representations'])} biểu diễn ({', '.join(g['representations'])}) "
          f"× {len(g['distances'])} khoảng cách ({', '.join(g['distances'])}). Các biểu diễn "
          f"powed, exponential và vai trò của việc chọn khoảng cách cho k-NN trên "
          f"UJIIndoorLoc lấy từ [21], [22]. Khoảng cách Bray-Curtis chuẩn hoá theo tổng "
          f"cường độ của hai vector:", "")],
    ])
    R.eq(r"d_{\mathrm{BC}}(a,b)=\dfrac{\sum_{j}\,|a_j-b_j|}{\sum_{j}\,(a_j+b_j)}")
    R.numbered([
        [("Số láng giềng và trọng số. ", "b"),
         (f"k từ {min(g['k'])} đến {max(g['k'])}; toạ độ dự đoán là trung bình có trọng "
          f"số của k láng giềng, trọng số tỉ lệ nghịch với khoảng cách mũ p, p từ "
          f"{min(g['powers'])} đến {max(g['powers'])} (p = 0 là trung bình đều):", "")],
    ], start=2)
    R.eq(r"\hat{x}=\dfrac{\sum_{i=1}^{k}\,w_i\,x_i}{\sum_{i=1}^{k}\,w_i},\qquad "
         r"w_i=\dfrac{1}{d_i^{\,p}}")
    R.numbered(start=3, items=[
        [("Dò độ lệch thiết bị lúc truy vấn. ", "b"),
         (f"Mọi RSSI phát hiện của truy vấn được cộng δ dB, δ trong khoảng ±"
          f"{cfg['offset_range']} dB với bước {g['offset_step_db']} dB; giữ δ cho khoảng "
          f"cách tới láng giềng gần nhất nhỏ nhất. Bước này chỉ dùng tín hiệu của chính "
          f"truy vấn, không dùng nhãn hay toạ độ [CẦN KIỂM CHỨNG/BỔ SUNG DOI: tài liệu "
          f"về hiệu chỉnh lệch thiết bị cho Wi-Fi fingerprinting để đối chiếu kỹ thuật "
          f"này].", "")],
        [("Ensemble. ", "b"), (f"Trung bình toạ độ của hai mô hình dùng hai biểu diễn "
                               f"khác nhau ({cfg['representation']} và "
                               f"{cfg.get('ensemble_with') or '—'}).", "")],
        [("Phá hoà tất định. ", "b"), ("Láng giềng được sắp theo (khoảng cách, chỉ số "
                                       "dòng), nên kết quả không phụ thuộc cách thư "
                                       "viện xử lý hoà.", "")],
    ])
    R.para("Sai số định vị của từng fingerprint được tính theo mét thực địa:")
    R.eq(r"e=\sqrt{(x-\hat{x})^2+(y-\hat{y})^2}\;\times\;\cos\varphi")
    R.h(2, "7.4. Kết quả từng bước")
    R.para(f"Hình 5 và Bảng 14 trình bày sai số của từng bước trên dev (dùng để chọn) và "
           f"trên tập kiểm định (chỉ để báo cáo). Mô hình được chọn đạt độ chính xác tòa "
           f"nhà {pc(sel['val_building_accuracy'])} và tầng {pc(sel['val_floor_accuracy'])}; "
           f"theo quy ước toạ độ chiếu, sai số trung bình là "
           f"{vn(sel['val_mean_projected_m'])} m.")
    R.fig(chart_stages(D), "Sai số trung bình của từng bước S0–S4 trên dev và trên "
                           "tập kiểm định")
    what = {"S0": "k-NN Version 1", "S1": "biểu diễn × khoảng cách",
            "S2": "+ k và trọng số theo khoảng cách", "S3": "+ dò lệch thiết bị",
            "S4": "+ ensemble hai biểu diễn"}
    rows14 = [("Thay đổi", lambda r: what.get(r["stage"], r["what"])),
              ("Cấu hình", lambda r: re.sub(r"(?<=[a-z])\+(?=[a-z])", " + ", r["config"])
                                       .replace("+-", "±").replace(", ", ",\n"))]
    for col, lab in (("mean", "TB"), ("median", "trung vị"), ("p75", "P75"),
                     ("p90", "P90"), ("p95", "P95"), ("max", "lớn nhất")):
        rows14.append((f"Dev – sai số {lab} (m)",
                       lambda r, c=col: vn(r[f"dev_{c}_ground_m"])))
    rows14.append(("Dev – tỉ lệ ≤ 8 m", lambda r: pc(r["dev_share_within_8m"], 1)))
    for col, lab in (("mean", "TB"), ("median", "trung vị"), ("p75", "P75"),
                     ("p90", "P90"), ("p95", "P95"), ("max", "lớn nhất")):
        rows14.append((f"KĐ – sai số {lab} (m)",
                       lambda r, c=col: vn(r[f"val_{c}_ground_m"])))
    rows14 += [("KĐ – tỉ lệ ≤ 8 m", lambda r: pc(r["val_share_within_8m"], 1)),
               ("KĐ – sai số TB, m toạ độ chiếu", lambda r: vn(r["val_mean_projected_m"])),
               ("KĐ – độ chính xác tòa nhà", lambda r: pc(r["val_building_accuracy"])),
               ("KĐ – độ chính xác tầng", lambda r: pc(r["val_floor_accuracy"])),
               ("KĐ – độ chính xác joint",
                lambda r: pc(r["val_joint_building_floor_accuracy"]))]
    R.tbl(pd.DataFrame([[lab] + [f(r) for _, r in A.iterrows()] for lab, f in rows14],
                       columns=["Chỉ số"] + list(A["stage"])),
          "Số liệu đầy đủ của Hình 5 (m thực địa trừ khi ghi khác; Dev = 6 fold giữ "
          "lại người đo; KĐ = tập kiểm định; results/metrics/positioning_ablation.csv)",
          [4.6] + [2.28] * len(A), size=10)
    R.h(2, "7.5. Phân bố sai số và độ bất định")
    CDF = cdf_table(D)
    R.para(f"Version 2 dịch toàn bộ phân bố sai số về phía nhỏ (Hình 6, Bảng 15). "
           f"Bootstrap {vi(BS['B'])} lần trên {vi(len(D.Pnew))} fingerprint cho khoảng tin "
           f"cậy 95% của sai số trung bình là {vn(BS['new_ci'][0])}–{vn(BS['new_ci'][1])} m "
           f"(Version 2) và {vn(BS['old_ci'][0])}–{vn(BS['old_ci'][1])} m (Version 1); mức "
           f"giảm ghép cặp là {vn(BS['diff'])} m ({vn(BS['diff_ci'][0])}–"
           f"{vn(BS['diff_ci'][1])} m), và Version 2 sai ít hơn trên {pc(BS['better'], 1)} "
           f"fingerprint (Bảng 16).")
    R.fig(chart_cdf(D), "Hàm phân bố tích luỹ của sai số định vị, Version 1 và "
                        "Version 2")
    v1, v2 = CDF.iloc[0], CDF.iloc[1]
    cdf_rows = [("Sai số trung bình (m)", "mean", vn)]
    cdf_rows += [(f"Phân vị {q} (m)" + (" – trung vị" if q == 50 else ""), f"p{q}", vn)
                 for q in (10, 25, 50, 75, 90, 95)]
    cdf_rows += [("Sai số lớn nhất (m)", "max", vn)]
    cdf_rows += [(f"Tỉ lệ fingerprint sai số ≤ {t} m", f"le{t}", lambda v: pc(v, 1))
                 for t in (2, 5, 8, 10)]
    R.tbl(pd.DataFrame([[lab, f(v1[k]), f(v2[k])] for lab, k, f in cdf_rows],
                       columns=["Chỉ số (m thực địa)", "Version 1", "Version 2"]),
          "Số liệu của Hình 6: phân vị sai số và tỉ lệ fingerprint dưới các ngưỡng "
          "(results/predictions/val_predictions_*.csv)", [8.0, 4.0, 4.0])
    R.tbl(pd.DataFrame([
        ["Sai số TB Version 1 (m)", vn(BS["old"]), f"{vn(BS['old_ci'][0])} – {vn(BS['old_ci'][1])}"],
        ["Sai số TB Version 2 (m)", vn(BS["new"]), f"{vn(BS['new_ci'][0])} – {vn(BS['new_ci'][1])}"],
        ["Mức giảm ghép cặp (m)", vn(BS["diff"]), f"{vn(BS['diff_ci'][0])} – {vn(BS['diff_ci'][1])}"],
        ["Tỉ lệ fingerprint Version 2 sai ít hơn", pc(BS["better"], 1), "—"],
    ], columns=["Đại lượng", "Giá trị", "Khoảng tin cậy 95%"]),
        f"Bootstrap {vi(BS['B'])} lần (seed 20251007) trên tập kiểm định", [7.5, 3.0, 5.5])
    R.h(2, "7.6. Đọc kết quả một cách thận trọng")
    R.para(f"Phần lớn cải thiện đến từ bước S1–S2. Dò độ lệch và ensemble giảm sai số "
           f"trên dev nhưng không trên tập kiểm định: {best_val['stage']} có sai số kiểm "
           f"định {vn(best_val['val_mean_ground_m'])} m, thấp hơn {sel['stage']} "
           f"({vn(sel['val_mean_ground_m'])} m). {sel['stage']} được giữ vì giao thức quyết "
           f"định trước; chọn lại theo số kiểm định sẽ là dùng tập kiểm tra để chọn mô "
           f"hình.")
    R.h(2, "7.7. Áp Version 2 lên các tập WAP của HUIM")
    CN = D.CN.set_index("condition").loc[SET_ORDER].reset_index()
    tm91 = D.CN[D.CN["condition"] == "control_top_mi__n91"].iloc[0]
    R.para(f"Các tập WAP của HUIM giảm sai số nhưng vẫn cách xa dải 7–8 m, và đối chứng "
           f"top-n theo mutual information ở cùng 91 WAP vẫn tốt hơn "
           f"({vn(tm91['study_mean_ground_m'])} m → {vn(tm91['opt_mean_ground_m'])} m). "
           f"Kết luận RQ2 của Version 1 không đổi (Hình 7, Bảng 17).")
    R.fig(chart_sets(D), "Sai số trung bình của từng tập WAP trước và sau Version 2")
    R.tbl(pd.DataFrame([[set_label(D, r), int(r["n_waps"]),
                         vn(r["study_mean_ground_m"]), vn(r["opt_mean_ground_m"]),
                         vn(r["study_median_ground_m"]), vn(r["opt_median_ground_m"]),
                         pc(r["study_joint_accuracy"]), pc(r["opt_joint_accuracy"])]
                        for _, r in D.CN.iterrows()],
                       columns=["Tập WAP", "Số WAP", "TB V1 (m)", "TB V2 (m)",
                                "Trung vị V1", "Trung vị V2", "Joint V1", "Joint V2"]),
          "Số liệu đầy đủ của Hình 7 và mọi tập WAP đã đánh giá "
          "(results/metrics/positioning_conditions.csv)",
          [5.0, 1.3, 1.55, 1.55, 1.65, 1.65, 1.65, 1.65], size=9.5)

    # ------------------------------------------------------------ 8 discussion
    R.h(1, "8. Thảo luận, hạn chế và kiểm chứng")
    R.para("Kết quả âm của HUIM chủ yếu do trần kích thước tập WAP, không chỉ do cách "
           "xếp hạng; cải thiện của Version 2 đến từ bộ ước lượng toạ độ, không từ HUIM.")
    R.para([("Vì sao HUIM chưa thắng đối chứng. ", "b"),
            (f"HUIM là bài toán không giám sát; không có định lý nào bảo đảm một HUI là "
             f"nhóm đặc trưng tốt cho phân loại [18]. Muốn mở rộng tập WAP phải hạ "
             f"minUtil, nhưng số tập mục tăng nhanh hơn số WAP rất nhiều và chạm ngân sách "
             f"trước khi tập WAP đủ lớn (mục 6.1). Thêm vào đó, mutual information gần "
             f"trùng tần suất phát hiện (ρ = "
             f"{vn(O6['spearman_support_vs_mi_over_all_detected_waps'], 4)}), nên Mô hình B "
             f"không tách được \"phổ biến\" với \"hữu ích\" ở mức từng WAP.", "")])
    R.para([("Giới hạn về tính hiệu lực.", "b")])
    R.bullets([
        f"Một bộ dữ liệu, một tập kiểm tra: toàn bộ kết quả đo trên UJIIndoorLoc [7]; "
        f"{vi(len(D.Pnew))} fingerprint kiểm định, thu bằng {n_ph} mã thiết bị trong "
        f"khoảng {int(round(va_days))} ngày. Không có tập kiểm tra thứ hai chưa từng "
        f"dùng.",
        "Tập kiểm định đã được nhìn: sai số kiểm định của cả năm bước S0–S4 được báo "
        "cáo; việc chọn S4 dựa trên dev, nhưng người đọc nên coi đây là năm lần đánh giá "
        "trên cùng một tập.",
        "Ánh xạ utility là lựa chọn thiết kế: định nghĩa q và p khác sẽ cho tập HUI và "
        "tập WAP khác.",
        "Cài đặt Python: thời gian khai phá tuyệt đối không so trực tiếp được với bản "
        "Java trong SPMF [20]; có ý nghĩa là hình dạng đường cong theo minUtil.",
        "Hệ toạ độ: EPSG:3857 được xác định bằng thực nghiệm, không có trong tài liệu bộ "
        "dữ liệu. So sánh với công trình khác phải dùng số theo toạ độ chiếu [CẦN KIỂM "
        "CHỨNG/BỔ SUNG DOI: số liệu sai số của các công trình khác trên validationData "
        "của UJIIndoorLoc, nêu rõ có hay không phạt sai tầng/tòa nhà].",
        "Xếp hạng tạp chí: chỉ số Scopus/WoS và hạng Q1/Q2 của các nguồn chưa được xác "
        "minh trực tiếp [CẦN KIỂM CHỨNG/BỔ SUNG DOI: tra SJR/JCR cho từng tạp chí trong "
        "danh mục tài liệu]; bài hội nghị không có hạng quartile.",
    ])
    R.para([("Kiểm chứng độc lập. ", "b"),
            ("Ba lớp kiểm chứng không dựa vào mã thí nghiệm (Bảng 18).", "")])
    vrows = []
    for name, label, what in (
            ("huim_checks.json", "HUIM",
             "Tính lại q, p, u(X) từ CSV gốc; kiểm tra mọi HUI của ba lần khai phá; vét "
             "cạn mọi tổ hợp 1–3 WAP"),
            ("positioning_checks.json", "Định vị",
             "Tính lại độ chính xác và sai số chỉ bằng thư viện chuẩn; tìm lại láng giềng "
             "gần nhất trên toàn bộ trainingData")):
        p = RESULTS / "verification" / name
        if p.exists():
            v = load_json(p)
            vrows.append([label, what, f"{v['n_checks'] - v['n_failed']}/{v['n_checks']} "
                                       f"phép kiểm tra đạt"])
    nx = excel_checks()
    if nx:
        vrows.append(["Excel", "Công thức tính lại minUtil, utility HUI và mọi dự đoán của "
                               "hai phiên bản", f"{vi(nx)} ô kiểm tra khớp"])
    R.tbl(pd.DataFrame(vrows, columns=["Lớp", "Nội dung", "Kết quả"]),
          "Kết quả kiểm chứng độc lập (results/verification/, "
          "verification/kiem_chung_thu_cong.xlsx)", [2.2, 10.0, 3.8])

    # ------------------------------------------------------------ 9 conclusion
    R.h(1, "9. Kết luận và hướng phát triển")
    R.para(f"HUIM có thể mô hình hoá nhất quán cho dữ liệu vân tay Wi-Fi nhưng chưa chọn "
           f"WAP tốt hơn các tiêu chí đơn giản; sai số định vị giảm từ "
           f"{vn(s0['val_mean_ground_m'])} m xuống {vn(sel['val_mean_ground_m'])} m nhờ tối "
           f"ưu bộ ước lượng toạ độ.")
    R.bullets([
        f"RQ1: Có. {vi(db['n_rows'])} fingerprint được chuyển thành "
        f"{vi(db['n_non_empty_transactions'])} giao dịch utility; EFIM chạy trên {n_runs} "
        f"cấu hình ngưỡng và khớp oracle vét cạn.",
        "RQ2: Không, ở ngân sách đã dùng. HUIM union thua top-n theo mutual information "
        "ở cả ba mô hình utility (" + " / ".join(pc(v) for v in U["huim_union"]) + " so với "
        + " / ".join(pc(v) for v in U["top_mi"]) + f") và bị chặn ở tối đa "
        f"{int(U['n_waps'].max())} WAP.",
        "RQ3: Số HUI và bộ nhớ khai phá tăng vọt khi hạ minUtil trong khi số WAP gần như "
        "đứng yên; phía hệ định vị, thời gian và bộ nhớ giảm gần tuyến tính theo số WAP.",
        f"RQ4: Sai số trung bình giảm {vn(BS['diff'])} m (khoảng tin cậy 95%: "
        f"{vn(BS['diff_ci'][0])}–{vn(BS['diff_ci'][1])} m) mà không dùng tập kiểm định để "
        f"chọn mô hình.",
    ])
    R.para([("Hướng phát triển.", "b")])
    R.numbered([
        "Thiết kế utility ngoài độc lập với tần suất phát hiện, vì mutual information "
        "hiện tại gần trùng tần suất.",
        "Thay HUIM đầy đủ bằng top-k HUIM hoặc tập HUI đóng để lấy được nhiều WAP hơn "
        "với cùng ngân sách [CẦN KIỂM CHỨNG/BỔ SUNG DOI: công trình gốc về top-k HUIM và "
        "closed HUIM].",
        "Kiểm tra lại trên một bộ dữ liệu thứ hai và một tập kiểm tra chưa từng dùng, ví "
        "dụ bộ dữ liệu crowdsourced của Lohan và cộng sự [22].",
    ])

    # ------------------------------------------------------------ references
    R.h(1, "Tài liệu tham khảo")
    R.para("Mọi DOI dưới đây đã được đối chiếu tiêu đề, tác giả và nơi đăng qua Crossref "
           "(research/doi_verification.json của dự án). Chỉ số Scopus/WoS và hạng "
           "quartile chưa được xác minh trực tiếp.")
    for k, ref in enumerate(REFS, start=1):
        par = R.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.left_indent = Cm(1.0)
        par.paragraph_format.first_line_indent = Cm(-1.0)
        r = par.add_run(f"[{k}]\t{ref}")
        r.font.size = Pt(12)
        _rfonts(r)
    return R


def word_finish(path: Path) -> None:
    """If Word is available: fill the table of contents and save once more."""
    if os.name != "nt":
        return
    ps = ("$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          "try { $d = $w.Documents.Open('%s'); "
          "foreach ($t in $d.TablesOfContents) { $t.Update() }; $d.Save(); "
          "$d.Close() } finally { $w.Quit() }" % str(path).replace("'", "''"))
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                       capture_output=True, text=True, timeout=600)
    LOG.info("Word TOC update: %s", "ok" if r.returncode == 0 else
             (r.stderr or "").strip()[-200:])


def main() -> None:
    D = Data()
    R = build(D)
    info = R.save(OUT)
    word_finish(OUT)
    LOG.info("saved %s: %d paragraphs, %d tables, %d figures", info["path"],
             info["n_paragraphs"], info["n_tables"], info["n_figures"])
    print("OK ", info["path"], f"({info['n_tables']} tables, {info['n_figures']} figures)")


if __name__ == "__main__":
    main()
