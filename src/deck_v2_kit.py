"""Drawing kit for the Version 2 deck (src/make_deck_v2.py).

The Version 1 deck "HUIM và IPS.pdf" was designed in Canva: navy backgrounds,
gold accents, Public Sans. This module rebuilds that look in python-pptx:

  * templates()  - every background is a page of the Version 1 PDF with its
                   TEXT removed (images, frames, arrows and effects are kept),
                   so the new text sits on exactly the original design;
  * tb() / card() / table() / formula() - text boxes with a tiny markup,
                   rounded cards, native tables, formulas drawn by matplotlib
                   mathtext (STIX);
  * chart_style() - matplotlib defaults for the white chart cards.

Coordinates are in points on a 1440 x 810 pt slide, the page size of the
Canva export, so positions measured on the Version 1 PDF can be reused as is.
"""
from __future__ import annotations

import os
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from lxml import etree
from matplotlib import font_manager
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

FONT = "Public Sans"
SYMBOL_FONT = "Segoe UI Symbol"        # Greek letters and math signs
W, H = 1440, 810                       # slide size in points

GOLD = RGBColor(0xF2, 0xBE, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xDC, 0xE3, 0xEE)
MUTED = RGBColor(0xA9, 0xB8, 0xCC)
INK = RGBColor(0x14, 0x22, 0x3A)       # dark text on white / gold
CARD_FILL = RGBColor(0x0D, 0x26, 0x45)
CARD_LINE = RGBColor(0x3A, 0x55, 0x7C)
ROW_A = RGBColor(0x10, 0x2C, 0x4F)
ROW_B = RGBColor(0x0B, 0x22, 0x40)
HEAD_FILL = RGBColor(0x17, 0x3A, 0x66)
GREEN = RGBColor(0x5F, 0xC9, 0x8B)
RED = RGBColor(0xF0, 0x7A, 0x6A)

# model colours of the Version 1 "Ba cách định nghĩa p(i)" cards
MODEL_HEX = {"A_signal": "#4E5A66", "B_discriminative": "#17797F",
             "C_rare_discriminative": "#E2672B"}

_COLORS = {"g": GOLD, "m": MUTED, "w": WHITE, "k": INK, "s": SOFT,
           "G": GREEN, "R": RED}


# ------------------------------------------------------------------ fonts
def _font_file(name: str) -> Path | None:
    for d in (Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft/Windows/Fonts",
              Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"):
        p = d / name
        if p.exists():
            return p
    return None


def _public_sans_cmap() -> set[int] | None:
    p = _font_file("PublicSans-Regular.ttf")
    if p is None:
        return None
    from fontTools.ttLib import TTFont
    return set(TTFont(str(p)).getBestCmap())


_CMAP = _public_sans_cmap()


def font_available() -> bool:
    return _CMAP is not None


# ------------------------------------------------------------- templates
def templates(v1_pdf: Path, out_dir: Path, spec: dict[str, tuple[int, list]],
              dpi: int = 144) -> dict[str, Path]:
    """Render pages of the Version 1 PDF without their text.

    spec: name -> (1-based page, [bbox of images to drop]); the bboxes are the
    Version 1 charts/screenshots that the new deck replaces.
    """
    import pymupdf
    out_dir.mkdir(parents=True, exist_ok=True)
    out = {}
    for name, (pno, drop) in spec.items():
        path = out_dir / f"bg_{name}.jpg"
        out[name] = path
        if path.exists() and path.stat().st_mtime > v1_pdf.stat().st_mtime:
            continue
        doc = pymupdf.open(v1_pdf)
        page = doc[pno - 1]
        for info in page.get_image_info(xrefs=True):
            bb = info["bbox"]
            if any(max(abs(a - b) for a, b in zip(bb, d)) < 3 for d in drop):
                page.delete_image(info["xref"])
        for b in page.get_text("dict")["blocks"]:
            for line in b.get("lines", []):
                for s in line["spans"]:
                    if s["text"].strip():
                        page.add_redact_annot(pymupdf.Rect(s["bbox"]))
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                              text=pymupdf.PDF_REDACT_TEXT_REMOVE)
        pix = page.get_pixmap(dpi=dpi)
        from PIL import Image
        Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(
            path, quality=90, optimize=True)
        doc.close()
    return out


# ------------------------------------------------------------- slide basics
def new_slide(prs, bg: Path):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(bg), 0, 0, Pt(W), Pt(H))
    return s


def _rgb(c) -> RGBColor:
    if isinstance(c, RGBColor):
        return c
    c = c.lstrip("#")
    return RGBColor(int(c[:2], 16), int(c[2:4], 16), int(c[4:], 16))


def _segments(text: str):
    """Split text into runs Public Sans can draw and runs it cannot."""
    if _CMAP is None:
        yield text, False
        return
    buf, sym = "", None
    for ch in text:
        s = ord(ch) not in _CMAP and not ch.isspace()
        if sym is None or s == sym:
            buf += ch
        else:
            yield buf, sym
            buf = ch
        sym = s
    if buf:
        yield buf, bool(sym)


_MARK = re.compile(r"\{([gbimwksGRu]+):(.*?)\}", re.S)


def add_rich(p, text: str, size: float, color, bold=False, italic=False,
             font: str = FONT):
    """Markup: {g:gold bold} {b:bold} {i:italic} {m:muted} {k:ink} {w:white}
    {s:soft} {G:green} {R:red} {u:underline}; flags combine, e.g. {gi:...}."""
    pos = 0
    parts = []
    for m in _MARK.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], ""))
        parts.append((m.group(2), m.group(1)))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], ""))
    for chunk, flags in parts:
        col = color
        for f in flags:
            if f in _COLORS:
                col = _COLORS[f]
        b = bold or "b" in flags or "g" in flags
        it = italic or "i" in flags
        for seg, is_sym in _segments(chunk):
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = SYMBOL_FONT if is_sym else font
            f.size = Pt(size)
            f.bold = b
            f.italic = it
            f.underline = "u" in flags
            f.color.rgb = _rgb(col)


def _bullet(p, indent: float, char: str, color):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(indent))))
    pPr.set("indent", str(-int(Pt(indent))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", str(_rgb(color)))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", FONT)
    if char == "#":                                   # 1. 2. 3.
        etree.SubElement(pPr, qn("a:buAutoNum")).set("type", "arabicPeriod")
    else:
        etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def tb(slide, x, y, w, h, paras, size=22, color=WHITE, bold=False,
       italic=False, align="l", anchor="t", spacing=1.12, gap=0.0,
       bullet=None, bullet_color=GOLD, indent=None, font=FONT, shadow=False):
    """Text box. paras: str (newline = paragraph) or list of str / dict.
    A dict paragraph may override size, color, bold, italic, align, bullet,
    gap (space before, pt)."""
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = box.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = paras.split("\n")
    for k, para in enumerate(paras):
        d = para if isinstance(para, dict) else {"text": para}
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                       "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[d.get("align", align)]
        p.line_spacing = d.get("spacing", spacing)
        g = d.get("gap", gap if k else 0)
        if g:
            p.space_before = Pt(g)
        sz = d.get("size", size)
        bu = d.get("bullet", bullet)
        if bu:
            _bullet(p, d.get("indent", indent or sz * 0.9), bu,
                    d.get("bullet_color", bullet_color))
        add_rich(p, d["text"], sz, d.get("color", color), d.get("bold", bold),
                 d.get("italic", italic), font)
    if shadow:
        text_shadow(box)
    return box


def text_shadow(box, blur=6, dist=3, alpha=60):
    """Soft drop shadow on every run, like the Canva titles."""
    for p in box.text_frame.paragraphs:
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            eff = etree.SubElement(rPr, qn("a:effectLst"))
            sh = etree.SubElement(eff, qn("a:outerShdw"))
            sh.set("blurRad", str(int(Pt(blur))))
            sh.set("dist", str(int(Pt(dist))))
            sh.set("dir", "2700000")
            sh.set("algn", "tl")
            c = etree.SubElement(sh, qn("a:srgbClr"))
            c.set("val", "000000")
            etree.SubElement(c, qn("a:alpha")).set("val", str(alpha * 1000))
            # effectLst must precede latin/ea/cs fonts in rPr
            for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
                el = rPr.find(qn(tag))
                if el is not None:
                    rPr.remove(el)
                    rPr.append(el)


def title(slide, text, x=46, y=36, w=1350, size=60, color=WHITE, h=None):
    """Slide title in the Version 1 style: bold, white, gold key words."""
    return tb(slide, x, y, w, h or size * 1.35, text, size=size, color=color,
              bold=True, spacing=0.95, shadow=True)


def card(slide, x, y, w, h, fill=CARD_FILL, line=CARD_LINE, lw=1.25,
         radius=0.06, alpha=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        if alpha is not None:
            clr = shp.fill._xPr.find(qn("a:solidFill"))[0]
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def rule(slide, x, y, w, color=GOLD, weight=3):
    """Short gold rule, as under the Version 1 titles."""
    return card(slide, x, y, w, weight, fill=color, line=None,
                shape=MSO_SHAPE.RECTANGLE)


def arrow(slide, x, y, w=34, h=22, color=GOLD):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Pt(x), Pt(y), Pt(w), Pt(h))
    a.fill.solid()
    a.fill.fore_color.rgb = _rgb(color)
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), Pt(x), Pt(y),
                                    Pt(w) if w else None, Pt(h) if h else None)


def chart_card(slide, path, x, y, w, h, pad=10):
    """White rounded card holding a chart image (the Version 1 look)."""
    card(slide, x, y, w, h, fill=WHITE, line=None, radius=0.025)
    from PIL import Image
    iw, ih = Image.open(path).size
    bw, bh = w - 2 * pad, h - 2 * pad
    sc = min(bw / iw, bh / ih)
    pw, ph = iw * sc, ih * sc
    return picture(slide, path, x + (w - pw) / 2, y + (h - ph) / 2, pw, ph)


# ------------------------------------------------------------------ tables
_NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _cell_border(cell, side, color, width):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"l": "a:lnL", "r": "a:lnR", "t": "a:lnT", "b": "a:lnB"}[side]
    old = tcPr.find(qn(tag))
    if old is not None:
        tcPr.remove(old)
    ln = etree.SubElement(tcPr, qn(tag))
    if width <= 0:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    else:
        ln.set("w", str(int(Pt(width))))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", str(_rgb(color)))
    # border elements must come before the cell fill
    fill = tcPr.find(qn("a:solidFill"))
    if fill is not None:
        tcPr.remove(fill)
        tcPr.append(fill)


def table(slide, x, y, widths, rows, size=17, head_size=None, row_h=30,
          head=True, head_fill=HEAD_FILL, head_color=GOLD,
          fills=(ROW_A, ROW_B), color=WHITE, align=None, line=CARD_LINE,
          highlight: dict | None = None, pad=7):
    """Native table. rows: list of lists of markup strings; widths in pt.
    highlight: {row_index: RGBColor} fill override for whole rows."""
    n, m = len(rows), len(widths)
    gf = slide.shapes.add_table(n, m, Pt(x), Pt(y), Pt(sum(widths)),
                                Pt(row_h * n))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = _NO_STYLE
    tbl.first_row = False
    tbl.horz_banding = False
    for j, wd in enumerate(widths):
        tbl.columns[j].width = Pt(wd)
    for i in range(n):
        tbl.rows[i].height = Pt(row_h)
        is_head = head and i == 0
        for j in range(m):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Pt(pad)
            c.margin_top = c.margin_bottom = Pt(pad * 0.55)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = head_fill if is_head else (highlight or {}).get(i, fills[(i - head) % 2])
            c.fill.solid()
            c.fill.fore_color.rgb = _rgb(fill)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            a = (align[j] if align else ("l" if j == 0 else "c"))
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                           "r": PP_ALIGN.RIGHT}[a]
            lines = str(rows[i][j]).split("\n")
            for k, ln in enumerate(lines):
                pp = p if k == 0 else tf.add_paragraph()
                pp.alignment = p.alignment
                add_rich(pp, ln, head_size or size if is_head else size,
                         head_color if is_head else color, bold=is_head)
            for side in "lrt":
                _cell_border(c, side, line, 0)
            _cell_border(c, "b", line, 0.75)
    return gf


# ---------------------------------------------------------------- formulas
def tex_png(tex: str, path: Path, size: float = 28, color: str = "#FFFFFF",
            dpi: int = 300) -> tuple[float, float]:
    """Draw $tex$ with mathtext (STIX) on a transparent PNG; returns its
    natural size in points."""
    with plt.rc_context({"mathtext.fontset": "stix"}):
        f = plt.figure(figsize=(0.01, 0.01))
        f.text(0, 0, f"${tex}$", fontsize=size, color=color)
        f.savefig(path, dpi=dpi, transparent=True, bbox_inches="tight",
                  pad_inches=0.03)
        plt.close(f)
    from PIL import Image
    w, h = Image.open(path).size
    return w * 72 / dpi, h * 72 / dpi


class Formulas:
    """Numbered formula images in one folder."""

    def __init__(self, out_dir: Path) -> None:
        self.dir = out_dir
        self.dir.mkdir(parents=True, exist_ok=True)
        self.n = 0

    def put(self, slide, tex, x, y, size=28, color="#FFFFFF", max_w=None,
            align="l", box_w=None):
        self.n += 1
        p = self.dir / f"formula_{self.n:03d}.png"
        w, h = tex_png(tex, p, size, color)
        if max_w and w > max_w:
            h, w = h * max_w / w, max_w
        if align == "c" and box_w:
            x = x + (box_w - w) / 2
        picture(slide, p, x, y, w, h)
        return w, h


# ------------------------------------------------------------------ charts
def chart_style() -> None:
    """matplotlib defaults for charts printed on white cards. Figures are
    created at their size on the slide (pt / 72 inches), so font sizes are
    the sizes the audience sees."""
    for f in ("PublicSans-Regular.ttf", "PublicSans-Bold.ttf",
              "PublicSans-Italic.ttf", "PublicSans-SemiBold.ttf"):
        p = _font_file(f)
        if p:
            font_manager.fontManager.addfont(str(p))
    fam = [FONT, "Segoe UI", "DejaVu Sans"] if _CMAP else ["Segoe UI", "DejaVu Sans"]
    plt.rcParams.update({
        "font.family": "sans-serif", "font.sans-serif": fam,
        "mathtext.fontset": "stixsans", "font.size": 14,
        "axes.titlesize": 16, "axes.titleweight": "bold",
        "axes.labelsize": 14, "xtick.labelsize": 13, "ytick.labelsize": 13,
        "legend.fontsize": 12.5, "legend.frameon": False,
        "axes.grid": True, "grid.alpha": 0.28, "grid.color": "#8A94A0",
        "axes.spines.top": False, "axes.spines.right": False,
        "axes.edgecolor": "#5B6676", "axes.labelcolor": "#14223A",
        "xtick.color": "#3B4656", "ytick.color": "#3B4656",
        "text.color": "#14223A", "axes.titlecolor": "#14223A",
        "savefig.dpi": 200, "figure.dpi": 100,
        "axes.unicode_minus": True,
    })


def fig_pt(w_pt: float, h_pt: float):
    return plt.figure(figsize=(w_pt / 72, h_pt / 72))
