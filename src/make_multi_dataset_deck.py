"""Build the deck "HUIM-IPS trên nhiều bộ dữ liệu" in the style of
"HUIM và IPS - V2.pdf": the same Canva backgrounds (pages of the Version 1 PDF
with the text removed), the same kit (src/deck_v2_kit.py), the same rules -
every number is read from results/ at build time, every formula carries its
source, every chart is saved with the rows it plots.

The content is that of report/Bao_Cao_HUIM_IPS_Da_Dataset.docx
(src/make_multi_dataset_report_docx.py), whose loaders are reused so the two
documents cannot disagree.

Inputs  HUIM và IPS.pdf (backgrounds), results/multi_dataset/**,
        research/doi_verification_*.json
Outputs slides/HUIM_IPS_Da_Dataset.pptx, "HUIM và IPS - Đa dữ liệu.pdf"
        (via PowerPoint), figures/deck_multi/*.png + .csv

Run: python src/make_multi_dataset_deck.py [--no-pdf] [--png DIR]
"""
from __future__ import annotations

import argparse
import datetime as dt
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import make_multi_dataset_report_docx as rd                     # noqa: E402
import make_multi_dataset_summary as ms                         # noqa: E402
from deck_v2_kit import (CARD_FILL, GOLD, HEAD_FILL, INK, MODEL_HEX,  # noqa: E402
                         MUTED, SOFT, WHITE, Formulas, arrow, card, chart_card,
                         chart_style, fig_pt, font_available, new_slide, table,
                         tb, templates, title)
from make_deck_v2 import TEMPLATES, V1_PDF, export_pdf          # noqa: E402
from pptx import Presentation                                   # noqa: E402
from pptx.dml.color import RGBColor                             # noqa: E402
from pptx.util import Pt                                        # noqa: E402
from utils import FIGURES, PROJECT_ROOT, SLIDES, get_logger, load_json  # noqa: E402

import matplotlib.pyplot as plt                                 # noqa: E402

LOG = get_logger("deck_multi")
OUT_PPTX = SLIDES / "HUIM_IPS_Da_Dataset.pptx"
OUT_PDF = PROJECT_ROOT / "HUIM và IPS - Đa dữ liệu.pdf"
ASSETS = SLIDES / "v2_assets"                 # backgrounds shared with V2
CHARTS = FIGURES / "deck_multi"
vn, pc, vi, mu = rd.vn, rd.pc, rd.vi, rd.mu
MODELS, LET = rd.MODELS, rd.LET
DARK_GREEN = RGBColor(0x2A, 0x4A, 0x2A)
C_S0, C_TR, C_OPT = "#8A94A0", "#E2672B", "#2A78D6"


def _save(fig, name: str, data: pd.DataFrame) -> Path:
    CHARTS.mkdir(parents=True, exist_ok=True)
    p = CHARTS / f"{name}.png"
    fig.savefig(p, bbox_inches="tight")
    plt.close(fig)
    data.to_csv(CHARTS / f"{name}.csv", index=False)
    return p


# ------------------------------------------------------------------ charts
def ch_dlulib(M, T: pd.DataFrame, w, h) -> Path:
    fl = [c for c in T.columns if c.endswith("_detect_pct")]
    Z = T[fl].to_numpy().T
    fig = fig_pt(w, h)
    ax = fig.add_axes([0.08, 0.2, 0.86, 0.62])
    ax.imshow(Z, cmap="Blues", vmin=0, vmax=100, aspect="auto")
    for i in range(Z.shape[0]):
        for j in range(Z.shape[1]):
            ax.text(j, i, f"{Z[i, j]:.0f}", ha="center", va="center", fontsize=12,
                    color="white" if Z[i, j] > 60 else "#14223A")
    ax.set_xticks(range(len(T)))
    ax.set_xticklabels(T["AP"], fontsize=11.5, rotation=0)
    ax.set_yticks(range(len(fl)))
    ax.set_yticklabels([f"Tầng {c.split('_')[1]}" for c in fl])
    ax.grid(False)
    ax.set_title("% vân tay huấn luyện của mỗi tầng phát hiện từng AP", loc="left")
    return _save(fig, "dlulib_ap_by_floor", T)


def ch_mining(M, w, h) -> Path:
    keys = [k for k in M.keys if k != "hus"]
    fig = fig_pt(w, h)
    n = len(keys)
    ncol = min(3, n)
    nrow = int(np.ceil(n / ncol))
    axs = fig.subplots(nrow, ncol, squeeze=False,
                       gridspec_kw={"hspace": 0.62, "wspace": 0.32})
    rows = []
    for a, k in zip(axs.flat, keys):
        H = M.runs[k]
        for m in MODELS:
            g = H[H["utility_model"] == m]
            ok = g[(g["status"] == "COMPLETED") & (g["n_huis"] > 0)]
            a.plot(ok["rel_minutil_pct"], ok["n_huis"], marker="o", ms=4, lw=2,
                   color=MODEL_HEX[m], label=f"Model {LET[m]}")
            ab = g[g["status"] != "COMPLETED"]
            a.scatter(ab["rel_minutil_pct"], ab["n_huis"], marker="x", s=46,
                      color=MODEL_HEX[m], zorder=4)
            for _, r in g.iterrows():
                rows.append({"dataset": M.name(k), "model": LET[m],
                             "minutil_pct": r["rel_minutil_pct"],
                             "status": r["status"], "n_huis": r["n_huis"],
                             "runtime_s": r["runtime_s"]})
        a.set_xscale("log")
        a.set_yscale("log")
        lo = H["rel_minutil_pct"].min()
        a.set_xticks([v for v in (20, 10, 5, 2, 1, 0.5, 0.2) if v >= lo * 0.99])
        a.xaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:g}%"))
        a.xaxis.set_minor_formatter(plt.NullFormatter())
        a.yaxis.set_minor_formatter(plt.NullFormatter())
        a.invert_xaxis()
        a.set_title(M.name(k), fontsize=14)
        a.tick_params(labelsize=11)
    for a in list(axs.flat)[n:]:
        a.set_visible(False)
    axs.flat[0].legend(fontsize=11, loc="upper left")
    return _save(fig, "huim_mining", pd.DataFrame(rows))


def ch_main(M, w, h) -> Path:
    c = M.cmp[(M.cmp["wap_set"] == "all") & (M.cmp["dataset"] != "hus")].copy()
    c["row"] = c["dataset_name"] + "\n" + c["test_set"]
    rows = list(dict.fromkeys(c["row"]))
    fig = fig_pt(w, h)
    ax = fig.add_axes([0.2, 0.1, 0.77, 0.86])
    y = np.arange(len(rows))[::-1]
    hh = 0.26
    for off, pre, col, lab in ((hh, "k-NN", C_S0, "k-NN nghiên cứu (S0)"),
                               (0, ms.TRANSFER, C_TR, ms.TRANSFER),
                               (-hh, "Tối ưu", C_OPT, "Tối ưu lại trên bộ đó")):
        v = np.array([c[(c["row"] == r) & c["method"].str.startswith(pre)]
                      ["mean_err_m"].astype(float).max() for r in rows])
        ok = ~np.isnan(v)
        ax.barh((y + off)[ok], v[ok], hh, color=col, label=lab)
        for yy, vv in zip((y + off)[ok], v[ok]):
            ax.text(vv, yy, f" {vn(vv)}", va="center", fontsize=11.5)
    ax.set_yticks(y)
    ax.set_yticklabels(rows, fontsize=12.5)
    ax.set_xlabel("Sai số trung bình (m)")
    ax.set_xlim(0, ax.get_xlim()[1] * 1.12)
    ax.grid(axis="y", visible=False)
    ax.legend(loc="lower right", fontsize=12)
    return _save(fig, "main_error", c.drop(columns=["row"]))


def ch_cdf(M, w, h) -> Path:
    pairs = [(k, t) for k in M.keys if k != "hus" for t in M.tests(k)]
    fig = fig_pt(w, h)
    n = len(pairs)
    ncol = min(4, n)
    nrow = int(np.ceil(n / ncol))
    axs = fig.subplots(nrow, ncol, squeeze=False,
                       gridspec_kw={"hspace": 0.62, "wspace": 0.3})
    rows = []
    for a, (k, t) in zip(axs.flat, pairs):
        xmax = 0
        for f, col, lab in (("study_knn", C_S0, "k-NN nghiên cứu"),
                            ("optimised", C_OPT, "Mô hình tối ưu")):
            e = np.sort(pd.read_csv(rd.ROOT / k / "predictions" /
                                    f"{t}__{f}.csv")["err_m"].to_numpy())
            a.plot(e, np.arange(1, len(e) + 1) / len(e) * 100, color=col, lw=2.2,
                   label=lab)
            xmax = max(xmax, np.percentile(e, 97))
            rows.append({"dataset": M.name(k), "test_set": t, "model": lab,
                         "p50": np.median(e), "p90": np.percentile(e, 90)})
        a.set_xlim(0, xmax * 1.05 if xmax else 1)
        a.set_ylim(0, 100)
        a.set_title(f"{M.name(k)} · {t}", fontsize=12.5)
        a.tick_params(labelsize=10.5)
        a.set_xlabel("Sai số (m)", fontsize=11)
    for a in list(axs.flat)[n:]:
        a.set_visible(False)
    axs.flat[0].legend(fontsize=10.5, loc="lower right")
    axs.flat[0].set_ylabel("% vân tay ≤ x", fontsize=11)
    return _save(fig, "error_cdf", pd.DataFrame(rows))


def ch_controls(M, w, h) -> Path:
    d = M.hv[(M.hv["family"] == "huim_union") & (M.hv["dataset"] != "hus")].copy()
    prim = d.groupby("dataset")["test_set"].transform("first")
    d = d[d["test_set"] == prim]
    d["row"] = (d["dataset_name"] + " · " + d["utility_model"] + " ("
                + d["n_waps"].astype(str) + " AP)")
    fig = fig_pt(w, h)
    ax = fig.add_axes([0.32, 0.1, 0.65, 0.86])
    y = np.arange(len(d))[::-1]
    for col, c, mk, lab in (("opt_mean_err_huim_m", C_OPT, "o", "HUIM (union)"),
                            ("opt_mean_err_top_mi_m", C_TR, "s", "Top-n theo MI"),
                            ("opt_mean_err_top_support_m", "#1BAF7A", "D",
                             "Top-n theo tần suất")):
        ax.scatter(d[col], y, s=70, color=c, marker=mk, label=lab, zorder=3,
                   edgecolor="white", linewidth=1)
    ax.scatter(d["opt_mean_err_all_waps_m"], y, s=110, marker="|",
               color="#14223A", label="Mọi AP", zorder=4)
    ax.set_yticks(y)
    ax.set_yticklabels(d["row"], fontsize=11.5)
    ax.set_xlabel("Sai số trung bình của mô hình tối ưu (m)")
    ax.set_xlim(left=0)
    ax.grid(axis="y", visible=False)
    ax.legend(fontsize=11, loc="lower right")
    return _save(fig, "huim_vs_controls", d.drop(columns=["row"]))


def ch_curve(M, w, h) -> Path:
    cv = M.cv[M.cv["dataset"] != "hus"]
    keys = list(dict.fromkeys(cv["dataset"]))
    fig = fig_pt(w, h)
    n = len(keys)
    ncol = min(3, n)
    nrow = int(np.ceil(n / ncol))
    axs = fig.subplots(nrow, ncol, squeeze=False,
                       gridspec_kw={"hspace": 0.62, "wspace": 0.3})
    for a, k in zip(axs.flat, keys):
        g = cv[cv["dataset"] == k]
        for ser, c, mk, lab in (("HUIM B", C_OPT, "o", "HUIM Model B"),
                                ("control_top_mi", C_TR, "s", "Top-MI"),
                                ("control_top_support", "#1BAF7A", "D",
                                 "Top-support"),
                                ("control_random", C_S0, "^", "Ngẫu nhiên")):
            hh = g[g["series"] == ser].sort_values("budget_n")
            if len(hh):
                a.plot(hh["budget_n"], hh["opt_mean_err_m"], color=c, lw=2,
                       marker=mk, ms=4.5, label=lab)
        a.axhline(g["all_waps_mean_err_m"].iloc[0], color="#14223A", lw=1.2,
                  ls=":", label="Mọi AP")
        if g["budget_n"].max() > 30:
            ms.log_counts(a, g["budget_n"])
        a.set_ylim(bottom=0)
        a.set_title(f"{g['dataset_name'].iloc[0]} · {g['test_set'].iloc[0]}",
                    fontsize=12.5)
        a.tick_params(labelsize=10.5)
        a.set_xlabel("Số AP giữ lại", fontsize=11)
    for a in list(axs.flat)[n:]:
        a.set_visible(False)
    axs.flat[0].legend(fontsize=10, loc="upper right")
    axs.flat[0].set_ylabel("Sai số TB (m)", fontsize=11)
    return _save(fig, "error_vs_aps", cv)


# ------------------------------------------------------------------ builder
class Builder:
    def __init__(self, M: rd.MD, bg: dict[str, Path]) -> None:
        self.M, self.bg = M, bg
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(1440), Pt(810)
        self.R = rd.Refs()
        self.F = Formulas(SLIDES / "deck_multi_assets" / "formulas")
        self.slides: list[tuple[str, object]] = []
        self.real = [k for k in M.keys if k != "hus"]
        self.tests = [(k, t) for k in self.real for t in M.tests(k)]

    def slide(self, bg: str, name: str):
        s = new_slide(self.prs, self.bg[bg])
        self.slides.append((name, s))
        return s

    def source(self, s, text, x, y, w, size=14):
        return tb(s, x, y, w, 36, "{i:Nguồn:} " + text, size=size, color=MUTED,
                  italic=True, spacing=1.0)

    def sub(self, s, text, y=112):
        tb(s, 48, y, 1340, 30, text, size=19, color=MUTED)

    def fin(self, k, t):
        return self.M.row(k, t, "Tối ưu PHASE I3")

    def s0(self, k, t):
        return self.M.row(k, t, "k-NN")

    def trf(self, k, t):
        return self.M.row(k, t, ms.TRANSFER)

    # ----------------------------------------------------------- 1 cover
    def s_cover(self):
        s = self.slide("cover", "Bìa")
        tb(s, 46, 70, 680, 330, ["HUIM & IPS", "trên nhiều", "{g:bộ dữ liệu}"],
           size=84, bold=True, spacing=0.94, shadow=True)
        tb(s, 50, 410, 610, 130,
           f"Chỉ thay dữ liệu, giữ nguyên phương pháp: HUIM chọn AP, k-NN và quy "
           f"trình tối ưu sai số theo mét — đánh giá tòa nhà, tầng, khoảng cách và "
           f"thời gian trên {len(self.real)} bộ dữ liệu Wi-Fi.", size=23,
           color=SOFT, spacing=1.12)
        card(s, 50, 612, 470, 48, fill=GOLD, line=None, radius=0.5)
        tb(s, 50, 612, 470, 48, f"MỞ RỘNG ĐA DỮ LIỆU  ·  {dt.date.today():%d/%m/%Y}",
           size=19, color=INK, bold=True, align="c", anchor="m")
        tb(s, 52, 680, 640, 60, "Tiếp nối Version 1 và Version 2 (UJIIndoorLoc). "
           "Mọi số liệu sinh tự động từ results/multi_dataset/.", size=16.5,
           color=MUTED, spacing=1.1)

    # ----------------------------------------------------------- 2 goals
    def s_goals(self):
        s = self.slide("plain", "Mục tiêu")
        title(s, "Mục tiêu và {g:câu hỏi}", size=58)
        self.sub(s, "Một kết quả đo trên MỘT bộ dữ liệu có thể chỉ là đặc thù của bộ "
                    "đó — nên cùng một chương trình được chạy trên mọi bộ trong thư mục")
        qs = [("CH1", "Áp dụng được?", "Quy trình HUIM → chọn AP → định vị chạy trên "
                                       "dữ liệu khác mà không phải sửa phương pháp?"),
              ("CH2", "Kết quả ra sao?", "Độ chính xác tòa nhà, tầng, sai số mét và "
                                         "thời gian định vị trên từng bộ?"),
              ("CH3", "Tối ưu còn hiệu quả?", f"Quy trình {ms.M_S0} → {ms.M_SEL} còn "
                                              f"giúp? Cấu hình {ms.M_SEL} dùng lại "
                                              "nguyên xi được?"),
              ("CH4", "HUIM hơn đối chứng?", "Tập AP do HUIM chọn có tốt hơn cách "
                                             "chọn đơn giản cùng số AP?")]
        for i, (n, h, t) in enumerate(qs):
            x = 46 + i * 340
            card(s, x, 170, 322, 330, radius=0.06)
            tb(s, x + 22, 188, 280, 50, f"{{g:{n}}}", size=40)
            tb(s, x + 22, 246, 280, 60, h, size=23, bold=True)
            tb(s, x + 22, 300, 280, 190, t, size=19, color=SOFT, spacing=1.1)
        card(s, 46, 530, 1348, 230, radius=0.04, line=GOLD)
        tb(s, 78, 552, 400, 40, "{g:NGUYÊN TẮC: CHỈ ĐỔI DỮ LIỆU}", size=24)
        tb(s, 78, 598, 1290, 160, [
            "Gọi lại đúng mã nguồn và hằng số của nghiên cứu: chuẩn hoá RSSI, ba mô hình "
            "p(i), EFIM, lưới minUtil và ngân sách, quy tắc chọn AP, lưới k-NN, các bước "
            "S0–S4.",
            "Chỉ quyết định riêng những gì dữ liệu buộc phải có: giá trị ‘không phát "
            "hiện’, đơn vị toạ độ, nhãn vị trí khi chỉ có một tầng, cách chia fold."],
           size=19.5, bullet="–", gap=10, spacing=1.1)

    # ----------------------------------------------------------- 3 process
    def s_exp(self):
        M = self.M
        s = self.slide("exp", "Quy trình")
        tb(s, 31, 46, 1360, 90, "Quy trình thực nghiệm {g:đa dữ liệu}", size=62,
           bold=True, shadow=True)
        n_runs = sum(len(M.runs[k]) for k in self.real if k != "uji")
        kp = [(225, str(len(self.real)), "bộ dữ liệu (thêm HUS gốc để đối chiếu)"),
              (378, str(len(self.tests)), "tập kiểm tra độc lập"),
              (534, str(n_runs), "lần chạy EFIM mới")]
        for y, v, lab in kp:
            tb(s, 70, y - 6, 250, 70, f"{{g:{v}}}", size=56)
            tb(s, 72, y + 70, 250, 50, lab, size=17, spacing=1.0)
        tb(s, 373, 188, 700, 34, "CÙNG MỘT QUY TRÌNH CHO MỌI BỘ", size=26, bold=True)
        steps = [(463, "01", "CHUẨN HOÁ\nDỮ LIỆU", "về lược đồ UJI: cột AP, 100 = "
                                                  "không thấy, X/Y, tòa/tầng"),
                 (668, "02", "GIAO DỊCH\nUTILITY", "q(i,T) = RSSI chuẩn hoá; p(i) "
                                                   "theo Model A/B/C"),
                 (883, "03", "KHAI PHÁ\nHUIM (EFIM)", "θ từ 20% giảm dần, ngân sách "
                                                      "600 s / 1 triệu HUI"),
                 (1089, "04", "ĐỊNH VỊ\nk-NN + TỐI ƯU", "k-NN nghiên cứu; S0 → S4 chọn "
                                                       "trên fold của train"),
                 (1294, "05", "ĐÁNH GIÁ", "tòa nhà, tầng, sai số mét, thời gian; HUIM "
                                          "vs đối chứng")]
        for cx, n, h, t in steps:
            tb(s, cx - 95, 270, 190, 60, f"{{g:{n}}}", size=46, align="c")
            tb(s, cx - 95, 388, 190, 64, h, size=19, bold=True, align="c", spacing=1.0)
            tb(s, cx - 84, 462, 168, 140, t, size=16, align="c", spacing=1.06)
        tb(s, 412, 676, 170, 50, "{g:KẾT QUẢ}", size=32)
        tb(s, 622, 652, 750, 100, "Mọi tham số chọn chỉ trên tập huấn luyện của "
           "từng bộ; tập kiểm tra chỉ dùng để chấm điểm, một lần cho mỗi bước.",
           size=19.5, anchor="m", spacing=1.08)

    # ----------------------------------------------------------- 4 datasets
    def s_datasets(self):
        M = self.M
        s = self.slide("plain", "Các bộ dữ liệu")
        title(s, f"{len(self.real)} bộ dữ liệu {{g:rất khác nhau}}", size=56)
        empty = M.DS[0]["summary"].get("empty_folders_in_dataset_root", [])
        self.sub(s, f"Từ {M.ap_range[0]} đến {M.ap_range[1]} AP, từ 1 phòng tới "
                    f"{max(M.by[k]['summary']['n_buildings'] for k in self.real)} "
                    "tòa nhà, test trùng vị trí "
                    "train hoặc ở vị trí mới hoàn toàn")
        rows = [["Bộ dữ liệu", "Train", "AP (thấy ở train)", "Vị trí", "Tòa / tầng",
                 "Tập kiểm tra (vị trí trùng train)"]]
        for k in self.real:
            sm = M.by[k]["summary"]
            tst = "\n".join(f"{n}: {vi(v['rows'])} dòng "
                            f"({v['positions_also_in_training']}/{v['positions']})"
                            for n, v in sm["tests"].items())
            rows.append([M.name(k), vi(sm["train"]["rows"]),
                         f"{sm['n_wap_columns']} ({sm['train']['aps_ever_detected']})",
                         str(sm["train"]["positions"]),
                         f"{sm['n_buildings']} / {sm['n_building_floor_classes']}",
                         tst])
        table(s, 46, 160, [270, 130, 200, 110, 140, 498], rows, size=16.5,
              row_h=64)
        foot = [f"{{g:Nguồn:}} UJIIndoorLoc {self.R.c('torres2014')}; UJI_LIB "
                f"{self.R.c('mendoza2018')}; DLULIB, HCXY, HUS không có tài liệu "
                "kèm theo."]
        if empty:
            foot.append(f"Thư mục rỗng (cả trong zip): {', '.join(empty)} — có thể "
                        "nằm ở các phần zip sau (…-1-002.zip).")
        tb(s, 46, 668, 1348, 100, foot, size=17, color=SOFT, gap=6, spacing=1.08)

    # ----------------------------------------------------------- 5 schema
    def s_schema(self):
        M = self.M
        s = self.slide("plain", "Lược đồ chung")
        title(s, "Đưa mọi bộ về {g:một lược đồ chung}", size=56)
        self.sub(s, "src/multi_datasets.py — mọi quyết định riêng được ghi vào "
                    "dataset_summary.json của từng bộ")
        sent = {"uji": "100", "uji_lib3": "100", "dlulib": "−100 → 100",
                "hcxy": "100", "hus_dedup": "không có", "hus": "không có"}
        why = {"uji": "test: người đo khác", "uji_lib3": "test trải đủ 25 tháng",
               "dlulib": "test đo lại đúng RP", "hcxy": "RP test không có ở train",
               "hus_dedup": "không có cột người/thiết bị"}
        rows = [["Bộ dữ liệu", "‘Không thấy’", "Hệ số ra mét", "Nhãn vị trí",
                 "Fold kiểm định (vì sao)"]]
        for k in self.real:
            sm = M.by[k]["summary"]
            g = sm["cv_group"].split(" - ")[0]
            rows.append([M.name(k), sent[k],
                         vn(sm["ground_factor"], 6) if sm["ground_factor"] != 1
                         else "1 (đã là mét)",
                         "tòa_tầng" if sm["n_building_floor_classes"] > 1 else "RP",
                         f"{g.replace('none', 'KFold ngẫu nhiên')} ({why.get(k, '')})"])
        table(s, 46, 160, [270, 170, 200, 170, 538], rows, size=17, row_h=58)
        card(s, 46, 540, 1348, 220, radius=0.04)
        tb(s, 70, 556, 1300, 34, "{g:Vì sao nhãn là RP khi chỉ có một tầng?}", size=21)
        tb(s, 70, 596, 1300, 160, [
            f"Model B, C dùng thông tin tương hỗ I(A;Y) {self.R.c('shannon1948')} với "
            "Y = tòa_tầng. Một tầng ⇒ Y không đổi ⇒ I(A;Y) = 0 với mọi AP ⇒ p(i) "
            "không xác định.",
            "Khi đó dùng lớp vị trí mịn nhất mà dữ liệu có: điểm tham chiếu (RP). "
            "Độ chính xác tòa/tầng chỉ báo cáo ở bộ có nhiều hơn một lớp."],
           size=19, bullet="–", gap=8, spacing=1.1)

    # ----------------------------------------------------------- 6 dlulib
    def s_dlulib(self):
        M = self.M
        if "dlulib" not in M.keys:
            return
        s = self.slide("plain", "DLULIB")
        title(s, "DLULIB: 17 AP {g:dùng chung} ba tầng", size=56)
        self.sub(s, "Ba file theo tầng chỉ được ghép thành một bài toán định vị tầng "
                    "nếu cột AP_k là cùng một AP vật lý ở cả ba file")
        T = rd.dlulib_floor_table(M)
        p = ch_dlulib(M, T, 1348, 330)
        chart_card(s, p, 46, 160, 1348, 340)
        home = T.groupby("strongest_floor")["AP"].apply(lambda x: ", ".join(x))
        tb(s, 46, 524, 1348, 230, [
            "Mỗi AP có một ‘tầng nhà’ nơi nó được thấy nhiều và mạnh nhất, yếu dần ở "
            "tầng khác — đúng như một AP cố định lan sóng qua sàn bê tông.",
            "AP mạnh nhất theo tầng: " + "; ".join(f"{{g:tầng {f}}}: {v}"
                                                   for f, v in home.items()) + ".",
            "−100 dBm là giá trị ‘không thấy’ (số ô −100 lớn gấp hàng nghìn lần ô "
            "−99) ⇒ đổi thành 100 như UJIIndoorLoc. Các file *_Mean không dùng vì "
            "nghiên cứu huấn luyện trên mẫu gốc."],
           size=19, bullet="–", gap=10, spacing=1.1)

    # ----------------------------------------------------------- 7 leakage
    def s_leak(self):
        M = self.M
        if "hus" not in M.keys or "hus_dedup" not in M.keys:
            return
        s = self.slide("plain", "Rò rỉ HUS")
        title(s, "HUS: tập test {g:chép từ tập train}", size=56)
        self.sub(s, "Kiểm tra trước khi tin: đếm dòng test trùng hệt (cùng vị trí, cùng "
                    "mọi RSSI) một dòng train")
        b = M.B["hus"]
        key = b.wap_cols + ["X", "Y", "FLOOR"]
        trk = b.train[key].drop_duplicates()
        twin = b.train.duplicated(subset=key, keep=False).mean()
        stats = []
        for t, te in b.tests.items():
            m = te[key].merge(trk, on=key, how="left", indicator=True)
            stats.append((t, int((m["_merge"] == "both").sum()), len(te)))
        cards_ = [(f"{h}/{n}", f"dòng {t} có bản sao y hệt trong train")
                  for t, h, n in stats]
        cards_.append((pc(twin, 1), "dòng train có ‘bản sao tự nhiên’ trong chính "
                                    "tập train"))
        for i, (v, lab) in enumerate(cards_):
            x = 46 + i * 452
            card(s, x, 166, 430, 170, radius=0.06)
            tb(s, x + 24, 182, 390, 80, f"{{g:{v}}}", size=54)
            tb(s, x + 24, 264, 390, 60, lab, size=18, color=SOFT, spacing=1.05)
        r_raw, r_ded = self.fin("hus", "testcase1"), self.fin("hus_dedup", "testcase1")
        r_raw2, r_ded2 = self.fin("hus", "testcase2"), self.fin("hus_dedup", "testcase2")
        rows = [["Mô hình tối ưu, sai số TB", "testcase1", "testcase2"],
                ["HUS nguyên bản (còn bản sao)", f"{vn(r_raw['mean_err_m'])} m",
                 f"{vn(r_raw2['mean_err_m'])} m"],
                ["HUS bỏ dòng chép (dùng để kết luận)", f"{vn(r_ded['mean_err_m'])} m",
                 f"{vn(r_ded2['mean_err_m'])} m"]]
        table(s, 46, 370, [520, 200, 200], rows, size=18, row_h=56,
              highlight={2: DARK_GREEN})
        n_rm = len(M.B["hus"].train) - len(M.B["hus_dedup"].train)
        tb(s, 1000, 370, 394, 300, [
            "Còn bản sao ⇒ k-NN chỉ việc tìm lại chính nó: sai số ‘gần 0’ là "
            "{g:học thuộc}, không phải định vị.",
            f"Sửa: bỏ đúng {vi(n_rm)} dòng train mà hai tập test chép lại "
            "(một bản cho mỗi dòng test)."], size=18.5, bullet="–", gap=10,
            spacing=1.1)
        rows = [["Bộ · tập test", "Dòng test trùng train"]]
        for k in [k for k in self.real if k != "hus_dedup"]:
            bb = M.B[k]
            kk = bb.wap_cols + ["X", "Y", "FLOOR"]
            tk = bb.train[kk].drop_duplicates()
            for t, te in bb.tests.items():
                m = te[kk].merge(tk, on=kk, how="left", indicator=True)
                rows.append([f"{M.name(k)} · {t}",
                             f"{int((m['_merge'] == 'both').sum())} / {vi(len(te))}"])
        tb(s, 46, 580, 920, 30, "{g:Các bộ khác: không có dòng test nào trùng train}",
           size=18)
        tb(s, 46, 614, 920, 150, "   ·   ".join(f"{a}: {b_}" for a, b_ in rows[1:]),
           size=16, color=SOFT, spacing=1.15)

    # ----------------------------------------------------------- 8 mapping
    def s_mapping(self):
        R = self.R
        s = self.slide("model", "Ánh xạ HUIM")
        tb(s, 50, 52, 1300, 90, "Vân tay Wi-Fi thành {g:giao dịch utility}", size=60,
           bold=True, shadow=True)
        boxes = [((68, 343), "GIAO DỊCH T", "một vân tay\n= một lần quét Wi-Fi"),
                 ((411, 669), "ITEM i", "một AP được\nphát hiện"),
                 ((736, 1012), "UTILITY", "u(i,T) = q(i,T) × p(i)\nq: RSSI chuẩn hoá\n"
                                          "p: giá trị của AP"),
                 ((1080, 1354), "HUI", "nhóm AP có tổng\nutility ≥ minUtil")]
        for (x0, x1), h, t in boxes:
            tb(s, x0 + 8, 176 if h == "UTILITY" else 190, x1 - x0 - 16, 30,
               f"{{g:{h}}}", size=22, align="c")
            tb(s, x0 + 8, 214 if h == "UTILITY" else 228, x1 - x0 - 16, 130, t,
               size=19, align="c", spacing=1.06)
        tb(s, 60, 389, 1310, 34, "{g:BA MÔ HÌNH UTILITY — DÙNG CHUNG q(i,T), CHỈ "
           "KHÁC p(i) — ÁP NGUYÊN CHO MỌI BỘ}", size=22, align="c")
        cols = [((51, 441), "MODEL A", "TÍN HIỆU",
                 "p(i) = 1\nđối chứng: utility chỉ do RSSI\nvà đồng xuất hiện"),
                ((441, 895), "MODEL B", "MUTUAL INFORMATION",
                 "p(i) ∝ I(A_i; Y)\nY = tòa_tầng (UJI, DLULIB)\nY = RP (bộ chỉ có một tầng)"),
                ((895, 1372), "MODEL C", "MI × ĐỘ HIẾM",
                 "p(i) ∝ I(A_i; Y) · log(N / (1 + support))\nphạt AP ở đâu cũng thấy")]
        for (x0, x1), h, sb, t in cols:
            tb(s, x0 + 10, 456, x1 - x0 - 20, 34, f"{{g:{h}}}", size=26, align="c")
            tb(s, x0 + 10, 496, x1 - x0 - 20, 34, sb, size=22, align="c")
            tb(s, x0 + 14, 552, x1 - x0 - 28, 110, t, size=18, align="c",
               spacing=1.08)
        tb(s, 90, 684, 1260, 80,
           f"{{g:Nguồn:}} utility của tập mục — Yao, Hamilton & Butz {R.c('yao2004')}; "
           f"thông tin tương hỗ — Shannon {R.c('shannon1948')}; IDF — Spärck Jones "
           f"{R.c('sparck1972')}. Ánh xạ vân tay → giao dịch là thiết kế của bài tìm "
           "hiểu.", size=18, align="c", anchor="m", spacing=1.08)

    # ----------------------------------------------------------- 9 formulas
    def s_formulas(self):
        M, R = self.M, self.R
        if "dlulib" not in M.keys:
            return
        s = self.slide("plain", "Công thức trên dữ liệu thật")
        title(s, "Công thức áp lên {g:một vân tay thật}", size=56)
        b = M.B["dlulib"]
        tr = b.train
        prep = rd.fit_preprocessor(tr)
        y = rd.class_label(b, tr)
        P = {m: rd.build_external_utility(m[0], tr, prep, y, "training")
             for m in MODELS}
        row = tr.iloc[[0]]
        tx = prep.transactions(row)[0][0]
        self.sub(s, f"DLULIB, dòng {int(row['SRC_ROW'].iloc[0])} của "
                    f"{row['SRC_FILE'].iloc[0].split('/')[-1]} — tính bằng chính mã "
                    f"nguồn của nghiên cứu")
        card(s, 46, 160, 640, 600, radius=0.04)
        tb(s, 70, 176, 600, 30, "1 · Chuẩn hoá RSSI (utility nội)", size=20,
           color=GOLD, bold=True)
        self.F.put(s, r"q(i,T)=\mathrm{round}\left(1+\frac{RSSI-RSSI_{\min}}"
                      r"{RSSI_{\max}-RSSI_{\min}}\times 99\right)", 70, 218, size=24,
                   max_w=590)
        tb(s, 70, 300, 600, 60, f"RSSI_min = {prep.rssi_min_train:g} dBm, RSSI_max "
           f"= {prep.rssi_max_train:g} dBm (học trên train DLULIB)".replace("-", "−"),
           size=17, color=SOFT)
        tb(s, 70, 360, 600, 30, "2 · Utility của AP, của tập, ngưỡng", size=20,
           color=GOLD, bold=True)
        self.F.put(s, r"u(i,T)=q(i,T)\,p(i),\qquad u(X)=\sum_{T\supseteq X}"
                      r"\sum_{i\in X}u(i,T)", 70, 402, size=23, max_w=590)
        self.F.put(s, r"X\in\mathrm{HUIs}\ \Leftrightarrow\ u(X)\geq\mathrm{minUtil}"
                      r"=\theta\sum_T TU(T)", 70, 470, size=23, max_w=590)
        w0 = list(tx)[0]
        q0 = tx[w0]
        r0 = int(row[w0].iloc[0])
        lo, hi = prep.rssi_min_train, prep.rssi_max_train
        f_ = lambda v: f"{v:g}".replace("-", "−")
        tb(s, 70, 548, 600, 200, [
            f"{{g:Kiểm tay}} {w0.replace('WAP_', '')}: RSSI {f_(r0)} dBm",
            f"q = 1 + ({f_(r0)} − ({f_(lo)})) / ({f_(hi)} − ({f_(lo)})) × 99 = "
            f"{vn(1 + (r0 - lo) / (hi - lo) * 99, 3)} → {q0}",
            f"u (Model B) = {q0} × {vn(P['B_discriminative'].external_utility[w0], 4)}"
            f" = {vn(q0 * P['B_discriminative'].external_utility[w0], 2)}"],
           size=17.5, gap=6, spacing=1.08)
        self.source(s, f"Yao, Hamilton & Butz {R.c('yao2004')}; minUtil theo tỉ lệ — "
                       f"Liu, Liao & Choudhary {R.c('liu2005')}", 70, 722, 600)
        rows = [["AP", "RSSI", "q", "p Model B", "u Model B"]]
        for w, q in tx.items():
            pb = P["B_discriminative"].external_utility[w]
            rows.append([w.replace("WAP_", ""), f_(int(row[w].iloc[0])), str(q),
                         vn(pb, 4), vn(q * pb, 2)])
        tu = sum(q * P["B_discriminative"].external_utility[w] for w, q in tx.items())
        rows.append(["{g:TU(T)}", "", "", "", f"{{g:{vn(tu, 2)}}}"])
        table(s, 714, 160, [130, 120, 90, 160, 180], rows, size=15.5,
              row_h=min(40, 600 / len(rows)))

    # ----------------------------------------------------------- 10 params
    def s_params(self):
        M = self.M
        s = self.slide("plain", "Tham số p(i)")
        title(s, "Ba mô hình {g:p(i)} trên từng bộ", size=56)
        self.sub(s, "Ước lượng chỉ trên tập huấn luyện; H(Y) = độ bất định của vị trí, "
                    "max I(A;Y) = phần bất định mà AP tốt nhất gỡ được")
        rows = [["Bộ dữ liệu", "Nhãn Y", "H(Y) bit", "max I(A;Y) bit",
                 "3 AP có I lớn nhất", "Tổng utility A / B / C"]]
        for k in self.real:
            d = M.diag[k]
            b = d["B_discriminative"]
            rows.append([M.name(k), "tòa_tầng" if M.by[k]["summary"]
                         ["n_building_floor_classes"] > 1 else "RP",
                         vn(b["label_entropy_H_Y_bits"], 3), vn(b["mi_bits_max"], 3),
                         ", ".join(t["wap"].replace("WAP_", "")
                                   for t in b.get("top10_waps_by_mi", [])[:3]),
                         " / ".join(vi(d[m]["total_database_utility"])
                                    for m in MODELS)])
        table(s, 46, 160, [260, 130, 130, 170, 290, 368], rows, size=16.5, row_h=60)
        y0 = 160 + 60 * len(rows) + 30
        self.F.put(s, r"p_B(i)=\frac{\max(I(A_i;Y),10^{-6})}{\max_j I(A_j;Y)},\qquad "
                      r"I(A;Y)=\sum_a\sum_y P(a,y)\log_2\frac{P(a,y)}{P(a)P(y)}",
                   60, y0, size=24, max_w=1300)
        tb(s, 60, y0 + 90, 1300, 60, "A_i có 6 trạng thái: không thấy và 5 khoảng "
           "RSSI (−104/−90/−80/−70/−60/0 dBm). Model C nhân thêm log(N/(1+support)).",
           size=17.5, color=SOFT)

    # ----------------------------------------------------------- 11 mining
    def s_mining(self):
        M = self.M
        s = self.slide("plain", "Khai phá HUIM")
        title(s, "Khai phá HUIM: {g:vùng chuyển pha} lặp lại", size=56)
        self.sub(s, "EFIM, θ = 20% giảm dần; dừng ở lần đầu chạm 600 s hoặc 1.000.000 "
                    "HUI; tinh chỉnh tối đa 3 lần (×: lần chạy chạm ngân sách)")
        p = ch_mining(M, 860, 560)
        chart_card(s, p, 46, 160, 870, 600)
        rows = [["Bộ · Model", "θ chính", "HUI", "AP", "EFIM (s)"]]
        for k in self.real:
            if k == "uji":
                continue
            H = M.runs[k]
            for m in MODELS:
                g = H[(H["utility_model"] == m) & (H["status"] == "COMPLETED")
                      & (H["n_huis"] > 0)]
                if not len(g):
                    continue
                r = g.loc[g["rel_minutil_pct"].idxmin()]
                rows.append([f"{M.name(k).split(' (')[0]} · {LET[m]}",
                             mu(r["rel_minutil_pct"]) + "%", vi(r["n_huis"]),
                             f"{int(r['unique_waps_in_huis'])}/"
                             f"{int(r['n_wap_columns_total'])}",
                             vn(r["runtime_s"], 1)])
        table(s, 936, 160, [170, 70, 100, 70, 70], rows, size=13,
              row_h=min(34, 600 / len(rows)), pad=4)

    # ----------------------------------------------------------- 12 protocol
    def s_protocol(self):
        R = self.R
        s = self.slide("plain", "Tối ưu S0-S4")
        title(s, f"Tối ưu sai số: {{g:cùng quy trình {ms.M_SEL}}}", size=56)
        self.sub(s, "Chọn trên 6 fold kiểm định lấy từ tập huấn luyện của TỪNG bộ — "
                    "tập kiểm tra chỉ chấm điểm sau khi đã chọn xong")
        steps = [("S0", "k-NN của nghiên cứu", "dùng như bộ ước lượng toạ độ"),
                 ("S1", "biểu diễn × khoảng cách", "positive / powed / exponential × "
                                                   "Euclid / Manhattan / Bray–Curtis / cosine"),
                 ("S2", "số láng giềng, trọng số", "k = 1…30, w = 1/dᵖ, p = 0…4"),
                 ("S3", "dò lệch thiết bị", "δ ∈ ±3, ±6, ±9 dB; giữ δ có láng giềng "
                                            "gần nhất"),
                 ("S4", "ghép hai biểu diễn", "trung bình toạ độ của hai mô hình")]
        for i, (n, h, t) in enumerate(steps):
            y = 160 + i * 98
            card(s, 46, y, 640, 86, radius=0.08)
            tb(s, 66, y + 14, 70, 56, f"{{g:{n}}}", size=34)
            tb(s, 150, y + 12, 520, 30, h, size=20, bold=True)
            tb(s, 150, y + 44, 520, 40, t, size=16, color=SOFT)
        tb(s, 46, 660, 640, 100, "S3, S4 chỉ được nhận nếu giảm sai số dev ≥ 0,05 m. "
           f"Ngoài ra: áp NGUYÊN XI cấu hình UJIIndoorLoc ({ms.M_SEL}) lên từng bộ.",
           size=17, color=SOFT, spacing=1.08)
        card(s, 714, 160, 680, 300, radius=0.04)
        tb(s, 738, 174, 640, 30, "Công thức (nguồn)", size=20, color=GOLD, bold=True)
        self.F.put(s, r"\mathrm{powed}=\left(\frac{RSSI+105}{105}\right)^{e},\ \ "
                      r"\mathrm{exp.}=\frac{e^{(RSSI+105)/24}}{e^{105/24}}", 738, 214,
                   size=22, max_w=630)
        self.F.put(s, r"d_{BC}(a,b)=\frac{\sum_i|a_i-b_i|}{\sum_i|a_i+b_i|},\qquad "
                      r"\hat{x}=\frac{\sum_j w_j x_j}{\sum_j w_j},\ w_j=\frac{1}{d_j^{\,p}}",
                   738, 300, size=22, max_w=630)
        tb(s, 738, 392, 640, 60, f"Torres-Sospedra et al. {R.c('torres2015')}; Bray & "
           f"Curtis {R.c('bray1957')}; Shepard {R.c('shepard1968')}", size=15,
           color=MUTED, italic=True)
        rows = [["Bộ dữ liệu", "Fold nhóm theo"]]
        for k in self.real:
            g = self.M.by[k]["summary"]["cv_group"].split(" - ")[0]
            rows.append([self.M.name(k), g.replace("none", "KFold ngẫu nhiên")])
        table(s, 714, 480, [380, 300], rows, size=16, row_h=min(46, 280 / len(rows)))

    # ----------------------------------------------------------- 13 divider
    def s_divider(self):
        s = self.slide("cover", "Kết quả")
        tb(s, 46, 70, 680, 120, "{g:Kết quả}", size=96, bold=True, shadow=True)
        tb(s, 46, 200, 650, 200, ["Tòa nhà, tầng,", "khoảng cách, thời gian"],
           size=54, bold=True, spacing=0.98, shadow=True)
        better = [(k, t) for k, t in self.tests
                  if self.fin(k, t)["mean_err_m"] < self.s0(k, t)["mean_err_m"]]
        tb(s, 50, 430, 600, 120, f"Mô hình tối ưu giảm sai số so với k-NN của nghiên "
           f"cứu ở {len(better)}/{len(self.tests)} tập kiểm tra.", size=26,
           color=SOFT, spacing=1.1)

    # ----------------------------------------------------------- 14 table
    def s_main_table(self):
        M = self.M
        s = self.slide("plain", "Bảng kết quả")
        title(s, "Tòa nhà, tầng, {g:sai số} và {g:thời gian}", size=54)
        self.sub(s, "Mọi AP; ‘—’ = bộ chỉ có một tòa / một tầng; thời gian = ms cho mỗi "
                    "vân tay trên tập test chính")
        rows = [["Bộ · tập test", "Tòa", "Tầng S0 → tối ưu", "Sai số TB: S0",
                 f"UJI {ms.M_SEL}", "Tối ưu lại", "Trung vị", "ms/vân tay"]]
        hi = {}
        for k, t in self.tests:
            a, f, tr_ = self.s0(k, t), self.fin(k, t), self.trf(k, t)
            nb, nf = a["n_buildings"], a["n_floor_classes"]
            rows.append([f"{M.name(k)} · {t}",
                         rd.acc(f["building_accuracy"], nb),
                         (f"{pc(a['floor_accuracy'])} → {{g:{pc(f['floor_accuracy'])}}}"
                          if nf > 1 else "—"),
                         f"{vn(a['mean_err_m'])} m",
                         "—" if pd.isna(tr_["mean_err_m"])
                         else f"{vn(tr_['mean_err_m'])} m",
                         f"{{g:{vn(f['mean_err_m'])} m}}", f"{vn(f['median_err_m'])} m",
                         "" if pd.isna(f["ms_per_fp"]) else vn(f["ms_per_fp"], 2)])
        table(s, 46, 160, [330, 110, 250, 140, 140, 150, 120, 108], rows, size=16,
              row_h=min(62, 560 / len(rows)))

    # ----------------------------------------------------------- 15 chart
    def s_main_chart(self):
        s = self.slide("plain", "Biểu đồ sai số")
        title(s, "Tối ưu lại {g:thắng} dùng lại nguyên xi", size=56)
        self.sub(s, "Sai số trung bình (m) trên từng tập kiểm tra: k-NN nghiên cứu, cấu "
                    "hình UJIIndoorLoc áp nguyên, và quy trình chọn chạy lại trên bộ đó")
        p = ch_main(self.M, 900, 590)
        chart_card(s, p, 46, 160, 910, 600)
        wins = [(k, t) for k, t in self.tests
                if self.trf(k, t)["mean_err_m"] <= self.fin(k, t)["mean_err_m"]]
        worse = [(k, t) for k, t in self.tests
                 if self.fin(k, t)["mean_err_m"] > self.s0(k, t)["mean_err_m"]]
        card(s, 980, 160, 414, 600, radius=0.05)
        tb(s, 1004, 180, 370, 40, "{g:ĐIỀU RÚT RA}", size=22)
        pts = ["Cái tổng quát hoá được là {g:quy trình chọn} trên fold của train, "
               "không phải bộ siêu tham số của UJIIndoorLoc.",
               f"Cấu hình {ms.M_SEL} áp nguyên bằng/thắng mô hình chọn lại ở "
               f"{len(wins)}/{len(self.tests)} tập (trên UJI hai mô hình là một)."]
        if worse:
            pts.append("Ngoại lệ: " + ", ".join(
                f"{self.M.name(k)} · {t} ({vn(self.s0(k, t)['mean_err_m'])} → "
                f"{vn(self.fin(k, t)['mean_err_m'])} m)" for k, t in worse)
                + " — fold kiểm định không mô phỏng được vị trí test mới.")
        tb(s, 1004, 226, 370, 520, pts, size=18, bullet="–", gap=14, spacing=1.1)

    # ----------------------------------------------------------- 16 selected
    def s_selected(self):
        M = self.M
        s = self.slide("plain", "Mô hình được chọn")
        title(s, "Mỗi bộ chọn {g:một mô hình khác}", size=56)
        self.sub(s, "Bước cuối được giữ trên dev (CHỌN) và sai số trên test (BÁO CÁO)")
        rows = [["Bộ dữ liệu", "Cấu hình được chọn", "Dev TB", "Test TB"]]
        for k in self.real:
            p = M.by[k]["pos"]
            A = M.abl(k)
            last = A[A["stage"] == p["stages"][-1]["stage"]]
            rows.append([M.name(k), p["selected_label"].replace("offset +-", "lệch ±")
                         .replace("ensemble with", "ghép"),
                         f"{vn(p['stages'][-1]['dev']['mean_ground_m'])} m",
                         " · ".join(f"{vn(v)} m" for v in last["test_mean_err_m"])])
        table(s, 46, 160, [260, 760, 140, 188], rows, size=16, row_h=56)
        uji = M.by.get("uji")
        rep = uji["pos"].get("reproduces_study") if uji else None
        if rep:
            card(s, 46, 590, 1348, 170, radius=0.05,
                 line=GOLD if rep["match"] else RGBColor(0xF0, 0x7A, 0x6A))
            tb(s, 72, 606, 1300, 34, "{g:KIỂM TRA TÁI LẬP UJIIndoorLoc qua code chung}",
               size=21)
            tb(s, 72, 648, 1300, 110, [
                f"Nghiên cứu gốc: {vn(rep['study_mean_ground_m'], 4)} m, "
                f"{pc(rep['study_joint_accuracy'])} đúng tòa–tầng · Code chung: "
                f"{vn(rep['generic_mean_ground_m'], 4)} m, "
                f"{pc(rep['generic_joint_accuracy'])} → "
                f"{'{g:TRÙNG KHỚP}' if rep['match'] else '{R:KHÁC}'}.",
                (f"PHASE I: {M.phase1['n_identical']}/{M.phase1['n_study']} điều kiện "
                 "chọn AP trùng tuyệt đối — phần mở rộng chỉ đổi dữ liệu, không đổi "
                 "phương pháp." if M.phase1 else
                 "Phần mở rộng chỉ đổi dữ liệu, không đổi phương pháp.")],
               size=18.5, bullet="–", gap=6, spacing=1.08)

    # ----------------------------------------------------------- 17 cdf
    def s_cdf(self):
        s = self.slide("plain", "Phân bố sai số")
        title(s, "Phân bố sai số {g:từng tập kiểm tra}", size=56)
        self.sub(s, "Đường càng lệch trái càng tốt: % vân tay có sai số ≤ x mét")
        p = ch_cdf(self.M, 1340, 580)
        chart_card(s, p, 46, 160, 1348, 600)

    # ----------------------------------------------------------- 18 controls
    def s_controls(self):
        M = self.M
        s = self.slide("plain", "HUIM vs đối chứng")
        title(s, "HUIM có hơn {g:đối chứng}?", size=56)
        self.sub(s, "Cùng mô hình tối ưu, cùng số AP: tập HUIM (union tốt nhất mỗi mô "
                    "hình) so với top-n theo MI và theo tần suất phát hiện")
        p = ch_controls(M, 900, 590)
        chart_card(s, p, 46, 160, 910, 600)
        H = M.hv[M.hv["dataset"] != "hus"]
        win = sum(r["opt_mean_err_huim_m"] < np.nanmin(
            [r["opt_mean_err_top_mi_m"], r["opt_mean_err_top_support_m"]]) - 1e-9
            for _, r in H.iterrows())
        card(s, 980, 160, 414, 600, radius=0.05)
        tb(s, 1004, 184, 370, 80, f"{{g:{win}/{len(H)}}}", size=58)
        tb(s, 1004, 270, 370, 60, "tập AP của HUIM thắng hẳn cả hai đối chứng cùng "
           "số AP", size=18, color=SOFT)
        tb(s, 1004, 360, 370, 380, [
            "Kết luận của nghiên cứu gốc lặp lại trên dữ liệu mới: HUIM tìm nhóm AP "
            "‘mạnh và hay đi cùng nhau’, nhưng không vì thế mà định vị tốt hơn.",
            "Bộ ít AP (" + ", ".join(
                f"{M.name(k).split(' (')[0]} {M.by[k]['summary']['n_wap_columns']}"
                for k in self.real if M.by[k]["summary"]["n_wap_columns"] <= 20)
            + "): union lấy gần hết AP, nên tác dụng giảm AP chỉ thấy ở top-n "
              "(slide sau)."], size=18, bullet="–", gap=12, spacing=1.1)

    # ----------------------------------------------------------- 19 curve
    def s_curve(self):
        M = self.M
        if not len(M.cv):
            return
        s = self.slide("plain", "Bao nhiêu AP")
        title(s, "Cần giữ {g:bao nhiêu AP}?", size=56)
        self.sub(s, f"Sai số của mô hình tối ưu theo số AP; bảng: số AP nhỏ nhất để sai "
                    f"số ≤ {vn(ms.CURVE_TOL, 1)} × sai số dùng mọi AP")
        p = ch_curve(M, 900, 590)
        chart_card(s, p, 46, 160, 910, 600)
        cv = M.cv[M.cv["dataset"] != "hus"]
        rows = [["Bộ", "Mọi AP", "HUIM B", "Top-MI", "Ngẫu nhiên"]]
        for k in dict.fromkeys(cv["dataset"]):
            g = cv[cv["dataset"] == k]
            lim = ms.CURVE_TOL * g["all_waps_mean_err_m"].iloc[0]
            cells = []
            for ser in ("HUIM B", "control_top_mi", "control_random"):
                h = g[(g["series"] == ser) & (g["opt_mean_err_m"] <= lim)]
                cells.append("—" if not len(h) else str(int(h["budget_n"].min())))
            rows.append([g["dataset_name"].iloc[0].split(" (")[0],
                         str(g["n_wap_columns"].iloc[0]), *cells])
        table(s, 976, 160, [150, 66, 66, 66, 70], rows, size=14,
              row_h=min(46, 560 / len(rows)), pad=4)

    # ----------------------------------------------------------- 20 timing
    def s_timing(self):
        M = self.M
        s = self.slide("plain", "Thời gian")
        title(s, "Thời gian {g:định vị} và {g:khai phá}", size=56)
        self.sub(s, f"Trung vị {rd.TIMING_REPEATS} lần đo trên tối đa "
                    f"{vi(rd.TIMING_MAX_ROWS)} vân tay test, mọi lõi CPU, máy không chạy "
                    "việc khác; EFIM chỉ chạy một lần khi xây dựng hệ thống")
        rows = [["Bộ · tập AP", "Số AP", "k-NN (ms/vt)", "Tối ưu (ms/vt)",
                 "CSDL (MB)", "EFIM (s)"]]
        for k in self.real:
            t = M.by[k]["timing"]
            if t is None:
                continue
            for _, r in t.iterrows():
                c = r["condition"]
                lab = ("mọi AP" if c == "baseline_all" else
                       "AP thấy ở train" if c == "baseline_detected" else
                       ("HUIM " + c.split("__")[1][0] + (" union" if "union" in c
                                                          else " top-n")))
                rows.append([f"{M.name(k).split(' (')[0]} · {lab}", str(int(r["n_waps"])),
                             vn(r["study_ms_per_fp"], 3), vn(r["opt_ms_per_fp"], 3),
                             vn(r["db_memory_mb"], 2),
                             "" if pd.isna(r["huim_runtime_s"])
                             else vn(r["huim_runtime_s"], 1)])
        if len(rows) == 1:
            tb(s, 46, 200, 1300, 40, "Chưa đo (chạy stage timing).", size=20)
            return
        n = len(rows)
        if n > 13:
            half = (n - 1 + 1) // 2
            for i, chunk in enumerate((rows[1:1 + half], rows[1 + half:])):
                table(s, 46 + i * 684, 160, [290, 66, 100, 110, 50, 52][:6],
                      [rows[0]] + chunk, size=12.5, row_h=min(40, 580 / (len(chunk) + 1)),
                      pad=4)
        else:
            table(s, 46, 160, [520, 120, 180, 180, 160, 188], rows, size=16,
                  row_h=min(46, 580 / n))

    # ----------------------------------------------------------- 21 lessons
    def s_lessons(self):
        M = self.M
        s = self.slide("plain", "Bài học")
        title(s, "Bài học {g:từ dữ liệu}", size=58)
        items = []
        if "uji_lib3" in M.keys:
            a, b = self.fin("uji_lib3", "Test-01"), self.fin("uji_lib3", "Test-02")
            items.append(("Vị trí mới khó hơn", f"UJI_LIB: Test-01 (đo lại vị trí train) "
                          f"{vn(a['mean_err_m'])} m, Test-02 (vị trí mới) "
                          f"{vn(b['mean_err_m'])} m — cùng một mô hình."))
        if "hus" in M.keys and "hus_dedup" in M.keys:
            a, b = self.fin("hus", "testcase1"), self.fin("hus_dedup", "testcase1")
            items.append(("Rò rỉ làm đẹp số liệu", f"HUS testcase1: {vn(a['mean_err_m'])} "
                          f"m khi còn bản sao, {vn(b['mean_err_m'])} m khi bỏ bản sao."))
        if "dlulib" in M.keys:
            a = self.fin("dlulib", "Testing")
            items.append(("Tầng có AP ‘riêng’ thì dễ", f"DLULIB đúng tầng "
                          f"{pc(a['floor_accuracy'])} chỉ với "
                          f"{M.by['dlulib']['summary']['n_wap_columns']} AP."))
        if "uji" in M.keys:
            a = self.fin("uji", "validationData")
            items.append(("UJIIndoorLoc khó nhất", f"{vn(a['mean_err_m'])} m: người đo "
                          "khác, điện thoại khác, sau 4 tháng, vị trí chưa khảo sát."))
        for i, (h, t) in enumerate(items):
            x = 46 + (i % 2) * 684
            y = 150 + (i // 2) * 300
            card(s, x, y, 664, 280, radius=0.05)
            tb(s, x + 26, y + 22, 620, 40, f"{{g:{h}}}", size=26)
            tb(s, x + 26, y + 82, 610, 180, t, size=21, spacing=1.12)

    # ----------------------------------------------------------- 22 verify
    def s_verify(self):
        M = self.M
        s = self.slide("plain", "Kiểm chứng")
        title(s, "Kiểm chứng {g:độc lập}", size=58)
        self.sub(s, "Người đọc tự chạy lại / tự tính lại được, không cần tin mã thực "
                    "nghiệm", y=116)
        tests = load_json(rd.ROOT / "efim_tests.json") if (rd.ROOT / "efim_tests.json"
                                                          ).exists() else None
        rep = M.by["uji"]["pos"].get("reproduces_study") if "uji" in M.keys else None
        cards_ = []
        if tests:
            cards_.append(("EFIM", f"{tests.get('passed', 0)}", "phép kiểm thử đạt",
                           "EFIM so với bộ duyệt vét cạn trên CSDL ngẫu nhiên; mô "
                           "hình utility; công thức.", "python -m pytest tests -q"))
        if rep and M.phase1:
            cards_.append(("Tái lập UJI",
                           f"{M.phase1['n_identical']}/{M.phase1['n_study']}",
                           "điều kiện PHASE I trùng",
                           f"PHASE I3 qua code chung: {vn(rep['generic_mean_ground_m'], 4)} "
                           f"m (gốc {vn(rep['study_mean_ground_m'], 4)} m).",
                           "run_multi_dataset.py --dataset uji"))
        if M.kit:
            cards_.append(("Excel", vi(M.kit["ok"]), "ô kiểm tra KHỚP",
                           f"{M.kit['bad']} ô lệch. Công thức Excel tính lại tòa, tầng, "
                           "sai số, dựng lại (x, y) từ láng giềng và trọng số.",
                           "verification/kiem_chung_da_dataset.xlsx"))
        w = (1348 - 32 * (len(cards_) - 1)) / max(len(cards_), 1)
        for i, (h, v, sb, t, cmd) in enumerate(cards_):
            x = 46 + i * (w + 32)
            card(s, x, 176, w, 470, radius=0.05)
            tb(s, x + 26, 196, w - 52, 30, h.upper(), size=21, color=GOLD, bold=True)
            tb(s, x + 26, 240, w - 52, 80, f"{{g:{v}}}", size=58)
            tb(s, x + 26, 326, w - 52, 30, sb, size=18, color=SOFT)
            tb(s, x + 26, 380, w - 52, 200, t, size=18.5, spacing=1.1)
            tb(s, x + 26, 590, w - 52, 50, cmd, size=14, color=MUTED, font="Consolas")
        tb(s, 46, 676, 1350, 90, "Mọi số trên slide được src/make_multi_dataset_deck.py "
           "đọc từ results/multi_dataset/ lúc dựng (không gõ tay); mỗi biểu đồ có file "
           "CSV trong figures/deck_multi/.", size=18, color=SOFT, spacing=1.1)

    # ----------------------------------------------------------- 23 conclusion
    def s_conclusion(self):
        better = [(k, t) for k, t in self.tests
                  if self.fin(k, t)["mean_err_m"] < self.s0(k, t)["mean_err_m"]]
        s = self.slide("stairs", "Kết luận")
        tb(s, 57, 58, 1000, 190, ["Kết luận và", "{g:hướng tiếp theo}"], size=72,
           bold=True, spacing=0.98, shadow=True)
        tb(s, 56, 296, 760, 480, [
            f"{{g:CH1:}} phương pháp chạy trên {len(self.real)} bộ dữ liệu "
            f"({self.M.ap_range[0]} → {self.M.ap_range[1]} AP, "
            f"{self.M.cls_range[0]} → {self.M.cls_range[1]} lớp tòa–tầng) mà không "
            "sửa thuật toán.",
            f"{{g:CH2–3:}} tối ưu sai số giảm lỗi so với k-NN gốc ở {len(better)}/"
            f"{len(self.tests)} tập test; tổng quát hoá là QUY TRÌNH CHỌN, không phải "
            "siêu tham số của UJI.",
            "{g:CH4:} HUIM vẫn không thắng đối chứng đơn giản cùng số AP.",
            "Kiểm tra dữ liệu là một phần của phương pháp (HUS rò rỉ). Tiếp theo: "
            "lấy đủ các thư mục còn thiếu (syl, tuji, UJI_LIB tầng 5)."],
           size=20, bullet="–", gap=14, spacing=1.1)

    # ----------------------------------------------------------- 24 refs
    def s_refs(self):
        R = self.R
        s = self.slide("plain", "Tài liệu tham khảo")
        title(s, "Tài liệu {g:tham khảo}", size=50)
        paras = [f"[{i}] {R.entry[k]}" for i, k in enumerate(R.order, 1)]
        tb(s, 46, 112, 1350, 640, paras, size=14.5 if len(paras) > 10 else 16,
           gap=8, spacing=1.04)
        tb(s, 46, 760, 1350, 30, "Mọi DOI đã kiểm tra qua Crossref (research/"
           "papers_table.csv, doi_verification_formulas.json, "
           "doi_verification_multi_dataset.json).", size=13, color=MUTED)

    def page_numbers(self):
        n = len(self.slides)
        for k, (_, s) in enumerate(self.slides, start=1):
            if k > 1:
                tb(s, 1330, 782, 90, 20, f"{k} / {n}", size=11.5, color=MUTED,
                   align="r")

    def build(self):
        for f in (self.s_cover, self.s_goals, self.s_exp, self.s_datasets,
                  self.s_schema, self.s_dlulib, self.s_leak, self.s_mapping,
                  self.s_formulas, self.s_params, self.s_mining, self.s_protocol,
                  self.s_divider, self.s_main_table, self.s_main_chart,
                  self.s_selected, self.s_cdf, self.s_controls, self.s_curve,
                  self.s_timing, self.s_lessons, self.s_verify, self.s_conclusion,
                  self.s_refs):
            n0 = len(self.slides)
            f()
            if len(self.slides) > n0:
                LOG.info("slide %2d  %s", len(self.slides), self.slides[-1][0])
        self.page_numbers()
        cp = self.prs.core_properties
        cp.title = "HUIM và IPS — mở rộng đa dữ liệu"
        cp.subject = "HUIM chọn AP + k-NN + tối ưu sai số trên nhiều bộ Wi-Fi fingerprint"
        return self.prs


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--png", type=Path)
    a = ap.parse_args()
    if not font_available():
        LOG.warning("Public Sans is not installed; PowerPoint will substitute a font")
    chart_style()
    bg = templates(V1_PDF, ASSETS, TEMPLATES)
    M = rd.MD()
    prs = Builder(M, bg).build()
    tmp = OUT_PPTX.with_name(OUT_PPTX.stem + ".__tmp__.pptx")
    prs.save(tmp)
    tmp.replace(OUT_PPTX)
    LOG.info("wrote %s (%d slides)", OUT_PPTX, len(prs.slides))
    if not a.no_pdf:
        export_pdf(OUT_PPTX, OUT_PDF, a.png)
        LOG.info("wrote %s", OUT_PDF)


if __name__ == "__main__":
    main()
