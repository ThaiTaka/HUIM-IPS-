"""Build verification/kiem_chung_thu_cong.xlsx - an Excel workbook in which
every reported number of the study is recomputed by Excel formulas from the
raw data, so it can be checked by hand, cell by cell.

Inputs (all produced by earlier phases; this script computes no result itself)
  ujiindoorloc/UJIndoorLoc/{training,validation}Data.csv
  results/metrics/huim_runs.csv
  results/ips_summary.json, results/positioning_optimisation.json
  results/predictions/val_predictions_study_knn.csv      (optimize_positioning.py)
  results/predictions/val_predictions_optimised.csv
  results/verification/*.csv / *.json                     (verify_huim.py,
                                                           verify_positioning.py)

Colour convention (from the workbook's own legend)
  blue text   = a value copied from a data or result file (an input)
  black text  = an Excel formula computed from those inputs
  green cell  = a check: KHỚP / ĐẠT means the recomputation agrees

Sheets
  0_HuongDan        how to use the workbook, step by step (Vietnamese)
  1_minUtil         minUtil% x total utility = absolute minUtil, every run
  2_HUI_mau         sample HUIs: utility from EFIM vs recomputed from the CSV
  3_HUI_tinh_tay    one HUI added up transaction by transaction
  4_kNN_tinh_tay    one validation fingerprint: representation, distances,
                    weights, coordinates and error, all as formulas
  5_DuDoan_kNN_cu   1,111 predictions of the study k-NN + accuracy / error
  6_DuDoan_ToiUu    1,111 predictions of the optimised model; (x, y) rebuilt
                    from the 12 neighbours by formula
  7_TongHop         Excel results next to the reported results
  8_HeSo_Mercator   the 0.766126 ground-distance factor from the latitude
  9_KiemTra_Python  PASS/FAIL table of verify_huim.py / verify_positioning.py

Run: python src/make_verification_kit.py  (then Excel recalculates on open)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).resolve().parent))
from utils import METRICS, PROJECT_ROOT, RESULTS, find_dataset, get_logger, load_json

LOG = get_logger("kit", "verification_kit.log")
OUT_DIR = PROJECT_ROOT / "verification"
OUT = OUT_DIR / "kiem_chung_thu_cong.xlsx"
PRED = RESULTS / "predictions"
VER = RESULTS / "verification"

FONT = "Times New Roman"
BLUE = Font(name=FONT, color="0000FF", size=10)
BLACK = Font(name=FONT, color="000000", size=10)
BOLD = Font(name=FONT, bold=True, size=10)
TITLE = Font(name=FONT, bold=True, size=13)
HEAD_FILL = PatternFill("solid", fgColor="DDE3EA")
CHECK_FILL = PatternFill("solid", fgColor="E2F0D9")
NOTE = Font(name=FONT, italic=True, color="555555", size=9)
THIN = Border(*(Side(style="thin", color="BBBBBB"),) * 4)
OK, BAD = "KHỚP", "LỆCH"


def q_formula(cell: str) -> str:
    """Excel version of q(i,T) = round(1 + (RSSI + 104)/104 * 99), 0 when the
    cell is 100. numpy rounds a .5 to the nearest EVEN integer and Excel's
    ROUND goes away from zero; the only RSSI that lands on .5 is -52 dBm
    (q = 50.5), so the formula reproduces numpy's rule explicitly."""
    x = f"(1+({cell}+104)/104*99)"
    return (f'=IF({cell}=100,0,IF(MOD({x},1)=0.5,'
            f'IF(MOD(INT({x}),2)=0,INT({x}),INT({x})+1),ROUND({x},0)))')


def pos_formula(cell: str) -> str:
    """The positive representation (not rounded): 0 when not detected,
    1 + (RSSI + 104)/104 * 99 otherwise."""
    return f"=IF({cell}=100,0,1+(MIN(0,MAX(-104,{cell}))+104)/104*99)"


def header(ws, row: int, cols: list[str], widths: list[float] | None = None):
    for j, h in enumerate(cols, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="center")
        c.border = THIN
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[L(j)].width = w


def put(ws, row, col, value, font=BLUE, fmt=None, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def check(ws, row, col, formula, fmt=None):
    return put(ws, row, col, formula, BLACK, fmt, CHECK_FILL)


# ----------------------------------------------------------------- sheets
def sheet_guide(wb, info: dict) -> None:
    ws = wb.active
    ws.title = "0_HuongDan"
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 120
    lines = [
        ("Bộ kiểm chứng thủ công - HUIM-IPS trên UJIIndoorLoc", TITLE),
        ("Mục đích: mọi con số chính của bài (minUtil, HUI, độ chính xác %, sai số mét) "
         "được Excel TÍNH LẠI bằng công thức từ dữ liệu thô. Bạn không phải tin script: "
         "bấm vào ô bất kỳ để xem công thức.", BLACK),
        ("Quy ước màu: chữ XANH DƯƠNG = số chép từ file dữ liệu/kết quả (đầu vào); "
         "chữ ĐEN = công thức Excel; ô NỀN XANH LÁ = ô kiểm tra (KHỚP/ĐẠT là đúng).", BOLD),
        ("", BLACK),
        ("THỨ TỰ KIỂM TRA ĐỀ NGHỊ", BOLD),
        ("Bước 1 - sheet 1_minUtil: minUtil tuyệt đối = minUtil(%) / 100 x tổng utility "
         "CSDL. Cột 'Kiểm tra' phải là KHỚP ở cả 37 lần khai phá. Ba dòng tô đậm là ba "
         "ngưỡng dùng trong các biểu đồ (A 1,9294%, B 2%, C 1,7321%).", BLACK),
        ("Bước 2 - sheet 2_HUI_mau: với từng HUI mẫu, utility do EFIM báo cáo phải bằng "
         "utility tính lại độc lập từ CSV (cột KHỚP) và phải >= minUtil (cột ĐẠT). "
         "Script verify_huim.py đã làm phép này cho TẤT CẢ HUI, và còn vét cạn mọi tổ hợp "
         "1-3 WAP để chứng minh EFIM không bỏ sót (xem sheet 9).", BLACK),
        ("Bước 3 - sheet 3_HUI_tinh_tay: cộng tay một HUI qua từng giao dịch: "
         "q = làm tròn(1 + (RSSI+104)/104 x 99), u(X,T) = tổng q x p, u(X) = tổng các "
         "u(X,T). So với số EFIM báo cáo và với minUtil.", BLACK),
        ("Bước 4 - sheet 4_kNN_tinh_tay: một fingerprint của validationData được định vị "
         "từng bước bằng công thức: biểu diễn tín hiệu -> khoảng cách tới các fingerprint "
         "huấn luyện -> láng giềng gần nhất -> trọng số -> toạ độ dự đoán -> sai số mét.",
         BLACK),
        ("Bước 5 - sheet 5 và 6: 1.111 dự đoán của mô hình cũ (k-NN k=1) và mô hình tối "
         "ưu. Mỗi dòng có công thức 'đúng tòa nhà?', 'đúng tầng?', sai số mét. Ô tổng hợp "
         "ở đầu sheet = AVERAGE / MEDIAN / PERCENTILE của các cột đó.", BLACK),
        ("Bước 6 - sheet 7_TongHop: so số Excel tính được với số trong báo cáo/biểu đồ. "
         "Tất cả phải KHỚP.", BLACK),
        ("Bước 7 - sheet 8_HeSo_Mercator: hệ số 0,766126 đổi 'mét toạ độ' sang mét thực "
         "địa được tính lại từ vĩ độ của khu vực (cos(vĩ độ)).", BLACK),
        ("", BLACK),
        ("HIỂU ĐÚNG CÁC CON SỐ", BOLD),
        (f"- Độ chính xác joint {info['joint_pct']} (mô hình cũ) nghĩa là "
         f"{info['joint_n']} / 1.111 fingerprint kiểm định được đoán đúng CẢ tòa nhà "
         "LẪN tầng. Đếm được bằng SUM trên cột 'Đúng cả hai?' của sheet 5. Mô hình "
         f"tối ưu (sheet 6): {info['opt_joint_n']} / 1.111 = {info['opt_joint_pct']}.",
         BLACK),
        (f"- Sai số trung bình: mô hình cũ {info['old_mean']} m, mô hình tối ưu "
         f"{info['opt_mean']} m (mét thực địa). Mô hình tối ưu là trung bình toạ độ "
         f"của {info['n_members']} member; mỗi member dịch RSSI truy vấn đi δ dB, lấy "
         f"{info['k']} láng giềng Bray-Curtis gần nhất, trọng số 1/d^{info['p']}.",
         BLACK),
        ("- Sai số mét của một fingerprint = khoảng cách thẳng giữa toạ độ thật và toạ độ "
         "dự đoán: căn((x1-x2)^2 + (y1-y2)^2) x 0,766126. 'Sai số trung bình' là "
         "AVERAGE của cột đó; 'trung vị' là MEDIAN.", BLACK),
        ("- Không có dữ liệu validation nào được dùng để chọn tham số: mô hình tối ưu được "
         "chọn trên các fold 'giữ lại người đo' của trainingData (xem "
         "results/positioning_optimisation.json, mục protocol).", BLACK),
        ("", BLACK),
        ("KIỂM TRA BẰNG PYTHON (độc lập với code thí nghiệm)", BOLD),
        ("python src/verify_huim.py                 - tính lại mọi HUI từ CSV, vét cạn tổ "
         "hợp 1-3 WAP", BLACK),
        ("python src/verify_positioning.py          - tính lại độ chính xác và sai số chỉ "
         "bằng thư viện chuẩn (csv + math)", BLACK),
        ("python src/verify_positioning.py --brute 25   - tìm lại láng giềng gần nhất "
         "trên toàn bộ 19.937 dòng cho 25 dòng ngẫu nhiên", BLACK),
        ("python src/verify_positioning.py --row 17     - in toàn bộ phép tính cho dòng 17 "
         "của validationData.csv", BLACK),
        ("", BLACK),
        ("Ghi chú: 'dòng' (csv_line) luôn là số dòng khi mở file CSV bằng Excel (dòng 1 là "
         "tiêu đề), nên có thể nhảy thẳng tới dòng đó để đối chiếu.", NOTE),
        ("Mở CSV gốc bằng Excel: file dùng dấu phẩy ngăn cột và dấu chấm thập phân. Nếu "
         "Excel đang dùng định dạng số kiểu Việt Nam, mở trực tiếp có thể dồn cả dòng vào "
         "một ô hoặc đọc sai số; khi đó dùng Data -> From Text/CSV, Delimiter = Comma, "
         "Locale = English (United States).", NOTE),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=2, value=t)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")


def sheet_minutil(wb, H: pd.DataFrame) -> dict:
    ws = wb.create_sheet("1_minUtil")
    header(ws, 1, ["Model", "Run tag", "minUtil (%)", "Tổng utility CSDL",
                   "minUtil tuyệt đối (huim_runs.csv)",
                   "minUtil tuyệt đối = C/100*D", "Kiểm tra", "Trạng thái",
                   "Số HUI", "Số WAP (union)"],
           [8, 34, 12, 18, 22, 22, 11, 26, 11, 12])
    primary = {}
    for m, g in H[(H.status == "COMPLETED") & (H.n_huis > 0)].groupby("utility_model"):
        primary[m] = g.loc[g.rel_minutil_pct.idxmin(), "run_tag"]
    rows = {}
    for i, (_, r) in enumerate(H.iterrows(), start=2):
        put(ws, i, 1, r.utility_model.split("_")[0])
        put(ws, i, 2, r.run_tag)
        put(ws, i, 3, float(r.rel_minutil_pct), fmt="0.0000")
        put(ws, i, 4, float(r.total_db_utility), fmt="#,##0.0000")
        put(ws, i, 5, float(r.abs_minutil), fmt="#,##0.0000")
        put(ws, i, 6, f"=C{i}/100*D{i}", BLACK, "#,##0.0000")
        check(ws, i, 7, f'=IF(ABS(F{i}-E{i})<=0.01,"{OK}","{BAD}")')
        put(ws, i, 8, r.status)
        put(ws, i, 9, int(r.n_huis), fmt="#,##0")
        put(ws, i, 10, int(r.unique_waps_in_huis))
        if r.run_tag in primary.values():
            for j in range(1, 11):
                ws.cell(row=i, column=j).font = Font(
                    name=FONT, bold=True, size=10,
                    color=ws.cell(row=i, column=j).font.color)
            rows[r.run_tag] = i
    n = len(H) + 2
    ws.cell(row=n + 1, column=2, value="Dòng in đậm: ngưỡng thấp nhất CHẠY XONG của "
            "mỗi model, là ngưỡng dùng trong các biểu đồ HUIM.").font = NOTE
    ws.freeze_panes = "A2"
    return rows


def sheet_hui_samples(wb, S: pd.DataFrame) -> None:
    ws = wb.create_sheet("2_HUI_mau")
    header(ws, 1, ["Run", "Itemset (các WAP)", "Độ dài", "u(X) do EFIM báo cáo",
                   "u(X) tính lại từ CSV (verify_huim.py)", "minUtil tuyệt đối",
                   "u(X) >= minUtil ?", "EFIM = tính lại ?", "Support EFIM",
                   "Support tính lại", "Support khớp ?"],
           [30, 60, 8, 18, 22, 16, 16, 15, 12, 12, 13])
    for i, (_, r) in enumerate(S.iterrows(), start=2):
        put(ws, i, 1, r.run)
        put(ws, i, 2, r.itemset)
        put(ws, i, 3, int(r.length))
        put(ws, i, 4, float(r.u_efim), fmt="#,##0.0000")
        put(ws, i, 5, float(r.u_recomputed), fmt="#,##0.0000")
        put(ws, i, 6, float(r.minutil_abs), fmt="#,##0.0000")
        check(ws, i, 7, f'=IF(E{i}>=F{i}*(1-1E-9),"ĐẠT","KHÔNG ĐẠT")')
        check(ws, i, 8, f'=IF(ABS(D{i}-E{i})<=1E-6*E{i},"{OK}","{BAD}")')
        put(ws, i, 9, int(r.support_efim))
        put(ws, i, 10, int(r.support_recomputed))
        check(ws, i, 11, f'=IF(I{i}=J{i},"{OK}","{BAD}")')
    n = len(S) + 3
    ws.cell(row=n, column=2, value="Mỗi run: 5 HUI utility lớn nhất, 5 HUI sát ngưỡng "
            "nhất (khó nhất), 5 HUI ngẫu nhiên. verify_huim.py kiểm tra TẤT CẢ HUI theo "
            "cùng cách.").font = NOTE
    ws.freeze_panes = "A2"


def sheet_hui_by_hand(wb, run_tag: str) -> None:
    ws = wb.create_sheet("3_HUI_tinh_tay")
    meta = load_json(VER / f"hui_breakdown_{run_tag}.json")
    B = pd.read_csv(VER / f"hui_breakdown_{run_tag}.csv")
    a, b = meta["itemset"]
    ws["A1"] = f"Cộng tay utility của HUI {{{a}, {b}}} - run {run_tag}"
    ws["A1"].font = TITLE
    ws["A2"] = ("Mỗi dòng là một giao dịch (fingerprint của trainingData.csv) chứa CẢ "
                "HAI WAP. q tính từ RSSI bằng công thức; p là trọng số ngoài của model "
                "(đầu vào, từ results/verification/p_values.csv).")
    ws["A2"].font = NOTE
    top = 4
    header(ws, top, ["Dòng trong trainingData.csv", f"RSSI {a}", f"q({a},T)",
                     f"p({a})", f"RSSI {b}", f"q({b},T)", f"p({b})",
                     "u(X,T) = q·p + q·p"], [16, 12, 11, 14, 12, 11, 14, 16])
    r0 = top + 1
    for i, (_, r) in enumerate(B.iterrows(), start=r0):
        put(ws, i, 1, int(r.csv_row))
        put(ws, i, 2, float(r[f"RSSI_{a}"]))
        put(ws, i, 3, q_formula(f"B{i}"), BLACK)
        put(ws, i, 4, float(r[f"p_{a}"]), fmt="0.000000")
        put(ws, i, 5, float(r[f"RSSI_{b}"]))
        put(ws, i, 6, q_formula(f"E{i}"), BLACK)
        put(ws, i, 7, float(r[f"p_{b}"]), fmt="0.000000")
        put(ws, i, 8, f"=C{i}*D{i}+F{i}*G{i}", BLACK, "#,##0.0000")
    r1 = r0 + len(B) - 1
    s = r1 + 2
    put(ws, s, 7, "u(X) = tổng cột H", BOLD)
    put(ws, s, 8, f"=SUM(H{r0}:H{r1})", BLACK, "#,##0.0000")
    put(ws, s + 1, 7, "Support = số dòng", BOLD)
    put(ws, s + 1, 8, f"=COUNT(A{r0}:A{r1})", BLACK)
    put(ws, s + 2, 7, "u(X) do EFIM báo cáo", BOLD)
    put(ws, s + 2, 8, float(meta["u_efim"]), fmt="#,##0.0000")
    put(ws, s + 3, 7, "Support do EFIM báo cáo", BOLD)
    put(ws, s + 3, 8, int(meta["support"]))
    put(ws, s + 4, 7, "minUtil tuyệt đối", BOLD)
    put(ws, s + 4, 8, float(meta["minutil_abs"]), fmt="#,##0.0000")
    put(ws, s + 5, 7, "Tổng của Excel = EFIM ?", BOLD)
    check(ws, s + 5, 8, f'=IF(ABS(H{s}-H{s + 2})<=1E-6*H{s + 2},"{OK}","{BAD}")')
    put(ws, s + 6, 7, "Support khớp ?", BOLD)
    check(ws, s + 6, 8, f'=IF(H{s + 1}=H{s + 3},"{OK}","{BAD}")')
    put(ws, s + 7, 7, "u(X) >= minUtil ?", BOLD)
    check(ws, s + 7, 8, f'=IF(H{s}>=H{s + 4}*(1-1E-9),"ĐẠT - là HUI",'
                        f'"KHÔNG ĐẠT")')
    ws.freeze_panes = ws.cell(row=r0, column=1)


def rep_formula(kind: str, cell: str) -> str:
    """Excel version of each signal representation (0 when not detected)."""
    c = f"MIN(0,MAX(-104,{cell}))"
    if kind == "positive":
        return f"=IF({cell}=100,0,1+({c}+104)/104*99)"
    if kind == "exponential":
        return f"=IF({cell}=100,0,EXP(({c}+105)/24)/EXP(105/24))"
    if kind == "powed":
        return f"=IF({cell}=100,0,(({c}+105)/105)^EXP(1))"
    raise ValueError(kind)


def model_members(cfg: dict) -> list[tuple[str, str, str]]:
    """(column prefix, representation, delta column) per ensemble member."""
    out = [("", cfg["representation"], "delta_db")]
    if cfg.get("ensemble_with"):
        out.append(("m2_", cfg["ensemble_with"], "m2_delta_db"))
    return out


def sheet_knn_by_hand(wb, tr: pd.DataFrame, va: pd.DataFrame,
                      study: pd.DataFrame, opt: pd.DataFrame, cfg: dict,
                      factor: float) -> int:
    """One validation fingerprint, located step by step with formulas, by the
    study's k-NN and by every member of the optimised model."""
    ws = wb.create_sheet("4_kNN_tinh_tay")
    wap = [c for c in tr.columns if c.startswith("WAP")]
    mem = model_members(cfg)
    # a typical case: the row whose optimised error is the median one
    o = opt.sort_values("err_ground_m").reset_index(drop=True)
    row = o.iloc[len(o) // 2]
    line = int(row.csv_line)
    s_row = study[study.csv_line == line].iloc[0]
    k, p = int(cfg["k"]), float(cfg["power"])
    study_line = int(s_row.nbr1_train_line)
    mem_lines = [[int(row[f"{pre}nbr{j}_train_line"]) for j in range(1, k + 1)]
                 for pre, _, _ in mem]
    rng = np.random.default_rng(7)
    rand_lines = [int(x) + 2 for x in rng.choice(len(tr), 2, replace=False)]
    cands = list(dict.fromkeys([study_line] + sum(mem_lines, []) + rand_lines))
    kinds = list(dict.fromkeys(["positive"] + [kd for _, kd, _ in mem]))
    q_raw = va.loc[line - 2, wap].to_numpy(float)
    used = [j for j in range(len(wap))
            if q_raw[j] != 100 or any(tr.loc[c - 2, wap[j]] != 100 for c in cands)]

    ws["A1"] = (f"Định vị tay fingerprint dòng {line} của validationData.csv "
                f"(dòng có sai số của mô hình tối ưu bằng trung vị)")
    ws["A1"].font = TITLE
    info = [("Tòa nhà / tầng thật", f"{int(row.true_building)} / {int(row.true_floor)}"),
            ("x thật (LONGITUDE)", float(row.true_x)),
            ("y thật (LATITUDE)", float(row.true_y)),
            ("k (số láng giềng mỗi member)", k),
            ("p (trọng số = 1/d^p)", p),
            ("Hệ số mét thực địa", factor)]
    for m, (pre, kind, dcol) in enumerate(mem, start=1):
        info.append((f"δ member {m} ({kind}), dB", float(row[dcol])))
    for i, (lab, v) in enumerate(info, start=2):
        put(ws, i, 1, lab, BOLD)
        put(ws, i, 2, v)
    KCELL, PCELL, FCELL = "$B$5", "$B$6", "$B$7"
    dcell = [f"$B${8 + m}" for m in range(len(mem))]
    ws["D2"] = ("Chỉ liệt kê các WAP được phát hiện ở truy vấn hoặc ở ít nhất một ứng "
                "viên: WAP vắng mặt ở cả hai bên cho |0-0| = 0 và 0+0 = 0, không đổi "
                "khoảng cách.")
    ws["D3"] = ("Ứng viên = láng giềng mà các mô hình đã chọn + 2 dòng ngẫu nhiên. Việc "
                "tìm trên TOÀN BỘ 19.937 dòng do verify_positioning.py --brute kiểm tra.")
    ws["D4"] = ("Mô hình tối ưu = trung bình toạ độ của các member; mỗi member dịch RSSI "
                "truy vấn đi δ dB, đổi sang biểu diễn của nó, lấy k láng giềng Bray-Curtis "
                "gần nhất, trọng số 1/d^p.")
    for c in ("D2", "D3", "D4"):
        ws[c].font = NOTE

    top = 9 + len(mem) + 1
    cols = ["WAP", "RSSI truy vấn", "positive (0..100)"]
    for m, (pre, kind, dcol) in enumerate(mem, start=1):
        cols += [f"RSSI dịch δ{m}", f"{kind} sau dịch δ{m}"]
    first_cand = len(cols) + 1
    width = 1 + len(kinds)
    for c in cands:
        cols += [f"RSSI dòng {c}"] + [f"{kd} dòng {c}" for kd in kinds]
    header(ws, top, cols, [11, 10, 11] + [10, 12] * len(mem)
           + [10] * (len(cols) - 3 - 2 * len(mem)))
    r0 = top + 1

    def c0(ci):                       # first (raw RSSI) column of candidate ci
        return first_cand + ci * width

    def rep_col(ci, kind):
        return c0(ci) + 1 + kinds.index(kind)

    for i, j in enumerate(used, start=r0):
        put(ws, i, 1, wap[j])
        put(ws, i, 2, float(q_raw[j]))
        put(ws, i, 3, rep_formula("positive", f"B{i}"), BLACK, "0.000")
        for m, (pre, kind, dcol) in enumerate(mem):
            cs, cr = 4 + 2 * m, 5 + 2 * m
            put(ws, i, cs, f"=IF(B{i}=100,100,MIN(0,MAX(-104,B{i}+{dcell[m]})))",
                BLACK)
            put(ws, i, cr, rep_formula(kind, f"{L(cs)}{i}"), BLACK, "0.000000")
        for ci, c in enumerate(cands):
            put(ws, i, c0(ci), float(tr.loc[c - 2, wap[j]]))
            for kd in kinds:
                put(ws, i, rep_col(ci, kd), rep_formula(kd, f"{L(c0(ci))}{i}"),
                    BLACK, "0.000000")
    r1 = r0 + len(used) - 1

    # ---- per-candidate summary rows (values sit in each candidate's raw column)
    s = r1 + 2
    lab = ["Dòng trong trainingData.csv", "Tòa nhà", "Tầng", "x", "y",
           "Manhattan với truy vấn positive (mô hình cũ)", "Hạng Manhattan"]
    for m, (pre, kind, dcol) in enumerate(mem, start=1):
        lab += [f"Bray-Curtis member {m} ({kind}, truy vấn dịch δ{m})",
                f"Hạng Bray-Curtis member {m}",
                f"Trọng số member {m} = 1/d^p nếu hạng <= k",
                f"Trọng số chuẩn hoá member {m}"]
    for t, text in enumerate(lab):
        put(ws, s + t, 1, text, BOLD)
    fc, lc = L(c0(0)), L(c0(len(cands) - 1))

    def row_rng(r):
        return f"${fc}${r}:${lc}${r}"

    pcol = {kd: None for kd in kinds}
    for ci, c in enumerate(cands):
        col = c0(ci)
        cl = L(col)
        put(ws, s, col, c)
        put(ws, s + 1, col, int(tr.loc[c - 2, "BUILDINGID"]))
        put(ws, s + 2, col, int(tr.loc[c - 2, "FLOOR"]))
        put(ws, s + 3, col, float(tr.loc[c - 2, "LONGITUDE"]), fmt="0.0000")
        put(ws, s + 4, col, float(tr.loc[c - 2, "LATITUDE"]), fmt="0.0000")
        pc = L(rep_col(ci, "positive"))
        put(ws, s + 5, col, f"=SUMPRODUCT(ABS($C${r0}:$C${r1}-{pc}{r0}:{pc}{r1}))",
            BLACK, "0.000")
        put(ws, s + 6, col, f"=RANK({cl}{s + 5},{row_rng(s + 5)},1)", BLACK)
        for m, (pre, kind, dcol) in enumerate(mem):
            b = s + 7 + 4 * m
            qc = L(5 + 2 * m)
            kc = L(rep_col(ci, kind))
            put(ws, b, col, f"=SUMPRODUCT(ABS(${qc}${r0}:${qc}${r1}-{kc}{r0}:{kc}{r1}))"
                            f"/SUMPRODUCT(${qc}${r0}:${qc}${r1}+{kc}{r0}:{kc}{r1})",
                BLACK, "0.000000")
            put(ws, b + 1, col, f"=RANK({cl}{b},{row_rng(b)},1)", BLACK)
            put(ws, b + 2, col, f"=IF({cl}{b + 1}<={KCELL},1/MAX({cl}{b},1E-9)"
                                f"^{PCELL},0)", BLACK, "0.000E+00")
            put(ws, b + 3, col, f"={cl}{b + 2}/SUM({row_rng(b + 2)})", BLACK,
                "0.000000")

    # ---- results ------------------------------------------------------------
    rr = s + len(lab) + 2
    res = [("MÔ HÌNH CŨ (k-NN k=1, positive + Manhattan)", None, None),
           ("Láng giềng gần nhất theo Excel (dòng)",
            f"=INDEX({row_rng(s)},MATCH(MIN({row_rng(s + 5)}),{row_rng(s + 5)},0))",
            None),
           ("Láng giềng mô hình đã dùng (file dự đoán)", study_line, "in")]
    res.append(("Kiểm tra", f'=IF(B{rr + 1}=B{rr + 2},"{OK}","{BAD}")', "chk"))
    res.append(("x dự đoán = x của láng giềng đó",
                f"=INDEX({row_rng(s + 3)},MATCH(B{rr + 1},{row_rng(s)},0))", None))
    res.append(("y dự đoán",
                f"=INDEX({row_rng(s + 4)},MATCH(B{rr + 1},{row_rng(s)},0))", None))
    res.append(("Sai số (mét thực địa)",
                f"=SQRT(($B$3-B{rr + 4})^2+($B$4-B{rr + 5})^2)*{FCELL}", None))
    res.append(("Sai số trong file dự đoán", float(s_row.err_ground_m), "in"))
    res.append(("Kiểm tra", f'=IF(ABS(B{rr + 6}-B{rr + 7})<1E-6,"{OK}","{BAD}")',
                "chk"))
    res.append(("MÔ HÌNH TỐI ƯU", None, None))
    xr, yr = [], []
    for m in range(len(mem)):
        b = s + 7 + 4 * m
        t = rr + len(res)
        res.append((f"x member {m + 1} = SUMPRODUCT(trọng số chuẩn hoá, x)",
                    f"=SUMPRODUCT({row_rng(b + 3)},{row_rng(s + 3)})", None))
        res.append((f"y member {m + 1}",
                    f"=SUMPRODUCT({row_rng(b + 3)},{row_rng(s + 4)})", None))
        xr.append(f"B{t}")
        yr.append(f"B{t + 1}")
    tx = rr + len(res)
    res.append(("x dự đoán = trung bình các member", f"=AVERAGE({','.join(xr)})",
                None))
    res.append(("y dự đoán", f"=AVERAGE({','.join(yr)})", None))
    res.append(("x dự đoán trong file", float(row.pred_x), "in"))
    res.append(("y dự đoán trong file", float(row.pred_y), "in"))
    res.append(("Kiểm tra toạ độ",
                f'=IF(AND(ABS(B{tx}-B{tx + 2})<1E-6,ABS(B{tx + 1}-B{tx + 3})<1E-6),'
                f'"{OK}","{BAD}")', "chk"))
    res.append(("Sai số (mét thực địa)",
                f"=SQRT(($B$3-B{tx})^2+($B$4-B{tx + 1})^2)*{FCELL}", None))
    res.append(("Sai số trong file dự đoán", float(row.err_ground_m), "in"))
    te = rr + len(res)
    res.append(("Kiểm tra", f'=IF(ABS(B{te - 2}-B{te - 1})<1E-6,"{OK}","{BAD}")',
                "chk"))
    for t, (text, v, kind) in enumerate(res):
        put(ws, rr + t, 1, text, BOLD)
        if v is None:
            continue
        if kind == "chk":
            check(ws, rr + t, 2, v)
        else:
            put(ws, rr + t, 2, v, BLUE if kind == "in" else BLACK, "0.0000")
    ws.column_dimensions["A"].width = 46
    ws.freeze_panes = ws.cell(row=r0, column=2)
    return line


def predictions_sheet(wb, name: str, P: pd.DataFrame, factor_ref: str,
                      reported: dict, cfg: dict | None = None,
                      train: pd.DataFrame | None = None) -> dict:
    """1,111 rows with formulas, plus a summary block above them. With `cfg`
    (the optimised model) each (x, y) is rebuilt from the neighbours of every
    ensemble member: mean over members of sum(w*x)/sum(w), w = 1/d^p."""
    ws = wb.create_sheet(name)
    top = 15
    rebuild = cfg is not None
    mem = model_members(cfg) if rebuild else []
    k = int(cfg["k"]) if rebuild else 0
    power = float(cfg["power"]) if rebuild else 0.0
    cols = ["Dòng validationData.csv", "Tòa nhà thật", "Tầng thật",
            "Tòa nhà dự đoán", "Tầng dự đoán", "Đúng tòa nhà? (1/0)",
            "Đúng tầng? (1/0)", "Đúng cả hai? (1/0)", "x thật", "y thật"]
    if rebuild:
        dpos = {}
        for m, (pre, kind, dcol) in enumerate(mem, start=1):
            cols.append(f"δ member {m} (dB)")
            dpos[pre] = len(cols)
        cols += ["x dự đoán (file)", "y dự đoán (file)",
                 "x dự đoán = TB member(Σw·x/Σw)", "y dự đoán (dựng lại)",
                 "(x,y) khớp file? (1/0)"]
        cxf, cyf, cxr, cyr, cmt = range(len(cols) - 4, len(cols) + 1)
    else:
        cols += ["Láng giềng (dòng trainingData.csv)", "x dự đoán", "y dự đoán"]
        cxr, cyr = len(cols) - 1, len(cols)
    cols += ["Sai số (m toạ độ chiếu)", "Sai số (m thực địa)"]
    cpe, cge = len(cols) - 1, len(cols)
    blocks = {}
    for pre, kind, dcol in mem:
        m = "2" if pre else "1"
        blocks[pre] = len(cols) + 1
        for j in range(1, k + 1):
            cols += [f"M{m} láng giềng {j}: dòng", f"M{m} d{j}", f"M{m} x{j}",
                     f"M{m} y{j}", f"M{m} w{j}=1/d^p"]
    header(ws, top, cols)
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[L(j)].width = 11
    r0 = top + 1
    for i, (_, r) in enumerate(P.iterrows(), start=r0):
        put(ws, i, 1, int(r.csv_line))
        put(ws, i, 2, int(r.true_building))
        put(ws, i, 3, int(r.true_floor))
        put(ws, i, 4, int(r.pred_building))
        put(ws, i, 5, int(r.pred_floor))
        put(ws, i, 6, f"=IF(B{i}=D{i},1,0)", BLACK)
        put(ws, i, 7, f"=IF(C{i}=E{i},1,0)", BLACK)
        put(ws, i, 8, f"=F{i}*G{i}", BLACK)
        put(ws, i, 9, float(r.true_x), fmt="0.0000")
        put(ws, i, 10, float(r.true_y), fmt="0.0000")
        if rebuild:
            parts_x, parts_y = [], []
            for pre, kind, dcol in mem:
                put(ws, i, dpos[pre], float(r[dcol]))
                b0 = blocks[pre]
                wx, wy, ww = [], [], []
                for j in range(k):
                    c = b0 + 5 * j
                    ln = int(r[f"{pre}nbr{j + 1}_train_line"])
                    put(ws, i, c, ln)
                    put(ws, i, c + 1, float(r[f"{pre}nbr{j + 1}_distance"]),
                        fmt="0.000000")
                    put(ws, i, c + 2, float(train.loc[ln - 2, "LONGITUDE"]),
                        fmt="0.0000")
                    put(ws, i, c + 3, float(train.loc[ln - 2, "LATITUDE"]),
                        fmt="0.0000")
                    put(ws, i, c + 4, f"=1/MAX({L(c + 1)}{i},1E-9)^{power:g}",
                        BLACK, "0.000E+00")
                    ww.append(f"{L(c + 4)}{i}")
                    wx.append(f"{L(c + 4)}{i}*{L(c + 2)}{i}")
                    wy.append(f"{L(c + 4)}{i}*{L(c + 3)}{i}")
                parts_x.append(f"({'+'.join(wx)})/({'+'.join(ww)})")
                parts_y.append(f"({'+'.join(wy)})/({'+'.join(ww)})")
            n = len(mem)
            put(ws, i, cxf, float(r.pred_x), fmt="0.0000")
            put(ws, i, cyf, float(r.pred_y), fmt="0.0000")
            put(ws, i, cxr, f"=({'+'.join(parts_x)})/{n}", BLACK, "0.0000")
            put(ws, i, cyr, f"=({'+'.join(parts_y)})/{n}", BLACK, "0.0000")
            put(ws, i, cmt, f"=IF(AND(ABS({L(cxr)}{i}-{L(cxf)}{i})<1E-6,"
                            f"ABS({L(cyr)}{i}-{L(cyf)}{i})<1E-6),1,0)", BLACK)
        else:
            put(ws, i, 11, int(r.nbr1_train_line))
            put(ws, i, cxr, float(r.pred_x), fmt="0.0000")
            put(ws, i, cyr, float(r.pred_y), fmt="0.0000")
        put(ws, i, cpe, f"=SQRT((I{i}-{L(cxr)}{i})^2+(J{i}-{L(cyr)}{i})^2)",
            BLACK, "0.0000")
        put(ws, i, cge, f"={L(cpe)}{i}*{factor_ref}", BLACK, "0.0000")
    r1 = r0 + len(P) - 1
    pe, ge = L(cpe), L(cge)

    ws["A1"] = name.split("_", 1)[1].replace("_", " ")
    ws["A1"].font = TITLE
    pe_rep = reported["positioning_error"]
    summary = [
        ("Số fingerprint", f"=COUNT(A{r0}:A{r1})", None, "0"),
        ("Độ chính xác tòa nhà", f"=AVERAGE(F{r0}:F{r1})",
         reported["building_accuracy"], "0.0000%"),
        ("Độ chính xác tầng", f"=AVERAGE(G{r0}:G{r1})",
         reported["floor_accuracy"], "0.0000%"),
        ("Độ chính xác joint (tòa nhà + tầng)", f"=AVERAGE(H{r0}:H{r1})",
         reported["joint_building_floor_accuracy"], "0.0000%"),
        ("Số dòng đúng cả hai", f"=SUM(H{r0}:H{r1})", None, "0"),
        ("Sai số TB (m toạ độ chiếu)", f"=AVERAGE({pe}{r0}:{pe}{r1})",
         pe_rep["mean_error"], "0.0000"),
        ("Sai số TB (m thực địa)", f"=AVERAGE({ge}{r0}:{ge}{r1})",
         pe_rep["mean_error_ground_m"], "0.0000"),
        ("Trung vị (m thực địa)", f"=MEDIAN({ge}{r0}:{ge}{r1})",
         pe_rep["median_error_ground_m"], "0.0000"),
        ("Phân vị 90 (m thực địa)", f"=PERCENTILE({ge}{r0}:{ge}{r1},0.9)",
         pe_rep["p90_error_ground_m"], "0.0000"),
        ("Tỉ lệ fingerprint sai số <= 8 m", f"=COUNTIF({ge}{r0}:{ge}{r1},\"<=8\")"
                                             f"/COUNT({ge}{r0}:{ge}{r1})", None,
         "0.00%"),
    ]
    if rebuild:
        summary.append(("Số dòng (x,y) dựng lại khớp file",
                        f"=SUM({L(cmt)}{r0}:{L(cmt)}{r1})", None, "0"))
    put(ws, 2, 1, "Chỉ số", BOLD)
    put(ws, 2, 2, "Excel tính", BOLD)
    put(ws, 2, 3, "Báo cáo (JSON)", BOLD)
    put(ws, 2, 4, "Kiểm tra", BOLD)
    cells = {}
    for t, (text, f, rep, fmt) in enumerate(summary, start=3):
        put(ws, t, 1, text, BOLD)
        put(ws, t, 2, f, BLACK, fmt)
        cells[text] = f"'{name}'!B{t}"
        if rep is not None:
            put(ws, t, 3, float(rep), fmt=fmt)
            check(ws, t, 4, f'=IF(ABS(B{t}-C{t})<=0.0001,"{OK}","{BAD}")')
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = ws.cell(row=r0, column=2)
    return cells


def sheet_summary(wb, cells_old: dict, cells_new: dict, rep_old: dict,
                  rep_new: dict) -> None:
    ws = wb.create_sheet("7_TongHop")
    header(ws, 1, ["Chỉ số", "Mô hình cũ: Excel", "Mô hình cũ: báo cáo",
                   "Khớp?", "Mô hình tối ưu: Excel", "Mô hình tối ưu: báo cáo",
                   "Khớp?"], [36, 16, 16, 9, 18, 18, 9])
    keys = [("Độ chính xác tòa nhà", "building_accuracy", "0.00%"),
            ("Độ chính xác tầng", "floor_accuracy", "0.00%"),
            ("Độ chính xác joint (tòa nhà + tầng)",
             "joint_building_floor_accuracy", "0.00%"),
            ("Sai số TB (m thực địa)", "mean_error_ground_m", "0.000"),
            ("Trung vị (m thực địa)", "median_error_ground_m", "0.000"),
            ("Phân vị 90 (m thực địa)", "p90_error_ground_m", "0.000")]
    for i, (lab, key, fmt) in enumerate(keys, start=2):
        put(ws, i, 1, lab, BOLD)
        for c0, cells, rep in ((2, cells_old, rep_old), (5, cells_new, rep_new)):
            v = rep[key] if key in rep else rep["positioning_error"][key]
            put(ws, i, c0, f"={cells[lab]}", Font(name=FONT, color="008000",
                                                  size=10), fmt)
            put(ws, i, c0 + 1, float(v), fmt=fmt)
            check(ws, i, c0 + 2, f'=IF(ABS({L(c0)}{i}-{L(c0 + 1)}{i})<=0.0001,'
                                 f'"{OK}","{BAD}")')
    ws.cell(row=10, column=1, value="Chữ xanh lá = liên kết tới ô tổng hợp của sheet 5 "
            "và 6 (không phải số gõ tay).").font = NOTE


def sheet_mercator(wb, tr: pd.DataFrame) -> str:
    ws = wb.create_sheet("8_HeSo_Mercator")
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 18
    rows = [
        ("LATITUDE nhỏ nhất của trainingData (m, Web Mercator)",
         float(tr.LATITUDE.min()), "in"),
        ("LATITUDE lớn nhất", float(tr.LATITUDE.max()), "in"),
        ("Bán kính trục lớn WGS84 R (m)", 6378137.0, "in"),
        ("y tâm = (min + max) / 2", "=(B2+B3)/2", None),
        ("Vĩ độ (độ) = DEGREES(2*ATAN(EXP(y/R)) - PI()/2)",
         "=DEGREES(2*ATAN(EXP(B5/B4))-PI()/2)", None),
        ("Hệ số mét thực địa = COS(vĩ độ)", "=COS(RADIANS(B6))", None),
        ("Hệ số dùng trong code (ips_model.py)", 0.766126, "in"),
        ("Kiểm tra (làm tròn 6 chữ số)", f'=IF(ROUND(B7,6)=B8,"{OK}","{BAD}")',
         "chk"),
    ]
    put(ws, 1, 1, "Vì sao nhân 0,766126: toạ độ của UJIIndoorLoc là Web Mercator "
        "(EPSG:3857), phóng đại khoảng cách 1/cos(vĩ độ) lần ở vĩ độ ~40°.", NOTE)
    for i, (lab, v, kind) in enumerate(rows, start=2):
        put(ws, i, 1, lab, BOLD)
        if kind == "chk":
            check(ws, i, 2, v)
        else:
            put(ws, i, 2, v, BLUE if kind == "in" else BLACK,
                "0.000000" if i >= 6 else "#,##0.0000")
    return "'8_HeSo_Mercator'!$B$7"


def sheet_python_checks(wb) -> None:
    ws = wb.create_sheet("9_KiemTra_Python")
    header(ws, 1, ["Nguồn", "Run / mô hình", "Phép kiểm tra", "Kết quả",
                   "Chi tiết"], [24, 36, 40, 10, 110])
    i = 2
    for fname in ("huim_checks.json", "positioning_checks.json"):
        p = VER / fname
        if not p.exists():
            continue
        for c in load_json(p)["checks"]:
            put(ws, i, 1, fname, BLACK)
            put(ws, i, 2, c.get("run", ""), BLACK)
            put(ws, i, 3, c["check"], BLACK)
            put(ws, i, 4, c["result"], BOLD,
                fill=CHECK_FILL if c["result"] == "PASS" else
                PatternFill("solid", fgColor="F4CCCC"))
            det = c.get("detail") or (f"tính lại {c.get('recomputed')} / báo cáo "
                                      f"{c.get('reported')}")
            put(ws, i, 5, det, BLACK)
            i += 1
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------- main
def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = find_dataset()
    tr = pd.read_csv(paths["trainingData.csv"])
    va = pd.read_csv(paths["validationData.csv"])
    H = pd.read_csv(METRICS / "huim_runs.csv")
    study = pd.read_csv(PRED / "val_predictions_study_knn.csv")
    opt = pd.read_csv(PRED / "val_predictions_optimised.csv")
    S = pd.read_csv(VER / "hui_samples.csv")
    isum = load_json(RESULTS / "ips_summary.json")
    rep_old = isum["conditions"]["baseline_all_520"]["result"]["knn"]
    po = load_json(RESULTS / "positioning_optimisation.json")
    rep_new = po["validation"]["selected"]
    cfg = po["selected_config"]

    def vn(x, d=2):
        return f"{x:.{d}f}".replace(".", ",")

    wb = Workbook()
    sheet_guide(wb, {
        "joint_pct": vn(rep_old["joint_building_floor_accuracy"] * 100) + "%",
        "joint_n": int(round(rep_old["joint_building_floor_accuracy"] * len(va))),
        "opt_joint_pct": vn(rep_new["joint_building_floor_accuracy"] * 100) + "%",
        "opt_joint_n": int(round(rep_new["joint_building_floor_accuracy"] * len(va))),
        "old_mean": vn(rep_old["positioning_error"]["mean_error_ground_m"]),
        "opt_mean": vn(rep_new["positioning_error"]["mean_error_ground_m"]),
        "n_members": len(model_members(cfg)), "k": int(cfg["k"]),
        "p": f"{float(cfg['power']):g}"})
    primary = sheet_minutil(wb, H)
    sheet_hui_samples(wb, S)
    b_run = [t for t in primary if t.startswith("B_")][0]
    sheet_hui_by_hand(wb, b_run)
    factor_ref = "'8_HeSo_Mercator'!$B$7"
    line = sheet_knn_by_hand(wb, tr, va, study, opt, cfg,
                             float(rep_old["positioning_error"]
                                   ["ground_correction_factor"]))
    c_old = predictions_sheet(wb, "5_DuDoan_kNN_cu", study, factor_ref, rep_old)
    c_new = predictions_sheet(wb, "6_DuDoan_ToiUu", opt, factor_ref, rep_new,
                              cfg=cfg, train=tr)
    sheet_summary(wb, c_old, c_new, rep_old, rep_new)
    sheet_mercator(wb, tr)
    sheet_python_checks(wb)
    wb.calculation.fullCalcOnLoad = True
    tmp = OUT.with_name(OUT.stem + ".__tmp__.xlsx")
    wb.save(tmp)
    tmp.replace(OUT)
    LOG.info("wrote %s (hand-located validation line %d)", OUT, line)
    excel_recalculate(OUT)


def excel_recalculate(path: Path) -> None:
    """openpyxl writes formulas without their values. If Excel is installed,
    let it calculate and save once, so the file also shows its numbers in
    viewers that do not calculate (and the check below can read them).
    Without Excel the workbook still calculates itself when opened."""
    import os
    import subprocess
    if os.name != "nt":
        LOG.info("not Windows: Excel will calculate the formulas on open")
        return
    ps = ("$x = New-Object -ComObject Excel.Application; $x.Visible = $false; "
          "$x.DisplayAlerts = $false; try { $w = $x.Workbooks.Open('%s'); "
          "$x.CalculateFull(); $w.Save(); $w.Close($true) } finally { $x.Quit() }"
          % str(path).replace("'", "''"))
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                       capture_output=True, text=True, timeout=600)
    if r.returncode != 0:
        LOG.info("Excel not available (%s): formulas calculate on open",
                 (r.stderr or "").strip().splitlines()[-1:] or "no message")
        return
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    bad, n_ok = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and (v.startswith("#") or v == BAD
                                           or v == "KHÔNG ĐẠT"):
                    bad.append(f"{ws.title}!{c.coordinate}={v}")
                elif v in (OK, "ĐẠT", "ĐẠT - là HUI"):
                    n_ok += 1
    if bad:
        raise SystemExit(f"workbook has {len(bad)} failing cells: {bad[:10]}")
    LOG.info("Excel recalculated the workbook: %d checks KHỚP/ĐẠT, no error "
             "or mismatch cell", n_ok)


if __name__ == "__main__":
    main()
