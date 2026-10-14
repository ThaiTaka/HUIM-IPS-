"""Build the Version 2 of the Canva deck "HUIM và IPS".

Version 1 is the user's Canva export "HUIM và IPS.pdf" (22 pages, project state
of 2026-10-08). Version 2 keeps its design - every background is a page of
that PDF with the text removed (src/deck_v2_kit.templates) - and rewrites the
content from the current results:

  * every number, table and chart is read from results/ at build time
    (the same loaders as src/make_huim_report_docx.py, so the deck and the
    report "Báo Cáo Thuật Toán HUIM" cannot disagree);
  * every formula carries its author / source; sources are the
    Crossref-verified bibliography of the project plus the eight works checked
    in research/doi_verification_formulas.json; a formula designed by this
    study says so;
  * new slides compare Version 1 and Version 2 and correct the statements of
    the Version 1 slides that were wrong.

Inputs  HUIM và IPS.pdf (Version 1, design source), results/**, figures/fig09
        csv, research/doi_verification_formulas.json, tests/test_efim.py
Outputs slides/HUIM_va_IPS_V2.pptx, "HUIM và IPS - V2.pdf" (via PowerPoint),
        figures/deck_v2/*.png + .csv (every chart with the rows it plots),
        slides/v2_assets/ (backgrounds, formula images)

Run: python src/make_deck_v2.py            (Windows + PowerPoint for the PDF)
     python src/make_deck_v2.py --no-pdf   (PPTX only)
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import make_huim_report_docx as hr                      # noqa: E402  (loaders)
from deck_v2_kit import (CARD_FILL, CARD_LINE, GOLD, HEAD_FILL, INK,  # noqa: E402
                         MODEL_HEX, MUTED, SOFT, WHITE, Formulas, add_rich, arrow,
                         card,
                         chart_card, chart_style, fig_pt, font_available,
                         new_slide, picture, rule, table, tb, templates, title)
from pptx import Presentation                                    # noqa: E402
from pptx.dml.color import RGBColor                              # noqa: E402
from pptx.enum.shapes import MSO_SHAPE                           # noqa: E402
from pptx.util import Pt                                         # noqa: E402
from utils import FIGURES, PROJECT_ROOT, RESULTS, SLIDES, get_logger, load_json  # noqa: E402

import matplotlib.pyplot as plt                                  # noqa: E402
import matplotlib.ticker as mticker                              # noqa: E402

LOG = get_logger("deck_v2", "deck_v2.log")
V1_PDF = PROJECT_ROOT / "HUIM và IPS.pdf"
OUT_PPTX = SLIDES / "HUIM_va_IPS_V2.pptx"
OUT_PDF = PROJECT_ROOT / "HUIM và IPS - V2.pdf"
ASSETS = SLIDES / "v2_assets"
CHARTS = FIGURES / "deck_v2"

vn, vi, pc, mu = hr.vn, hr.vi, hr.pc, hr.mu
SUP = str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹")


def sup(x) -> str:
    """Exponent typed with superscript digits (Public Sans has them)."""
    return str(int(x)).translate(SUP)
LETTER = hr.LETTER
MODELS = ["A_signal", "B_discriminative", "C_rare_discriminative"]
MNAME = {"A_signal": "A (p = 1)", "B_discriminative": "B (MI)",
         "C_rare_discriminative": "C (MI × IDF)"}

# Version 1 pages used as backgrounds: name -> (page, [bbox of the V1 chart /
# screenshot images that this version replaces])
TEMPLATES = {"cover": (1, []), "what": (2, []), "ips": (3, []),
             "pdef": (4, [(43, 208, 1391, 740)]), "model": (8, []),
             "efim": (14, []), "exp": (16, []),
             "res_m": (18, [(762, 265, 1422, 595)]),
             "res_cost": (19, [(120, 114, 1193, 500)]),
             "res_fr": (20, [(24, 117, 816, 684)]),
             "plain": (21, []), "stairs": (22, [])}


# --------------------------------------------------------------- references
# Short in-slide labels. The full entries come from the Crossref-verified
# bibliography of the project (make_huim_report_docx.REFS) or from
# research/doi_verification_formulas.json.
SHORT = {
    "agrawal1993": ("Agrawal, Imieliński & Swami, SIGMOD 1993", "170035.170072"),
    "yao2004": ("Yao, Hamilton & Butz, SDM 2004", "1.9781611972740.51"),
    "liu2005": ("Liu, Liao & Choudhary, PAKDD 2005", "11430919_79"),
    "zida2017": ("Zida, Fournier-Viger, Lin, Wu & Tseng, KAIS 2017", "s10115-016-0986-0"),
    "zhang2018": ("Zhang, Almpanidis, Wang & Liu, ESWA 2018", "j.eswa.2018.02.008"),
    "gan2021": ("Gan et al., IEEE TKDE 2021", "TKDE.2019.2942594"),
    "fv2019": ("Fournier-Viger et al., 2019", "978-3-030-04921-8_1"),
    "spmf2016": ("Fournier-Viger et al., SPMF 2016", "978-3-319-46131-1_8"),
    "bahl2000": ("Bahl & Padmanabhan, INFOCOM 2000", "INFCOM.2000.832252"),
    "chen2006": ("Chen, Yang, Yin & Chai, IEEE TKDE 2006", "TKDE.2006.112"),
    "song2016": ("Song, Wang & Yuan, ISPRS IJGI 2016", "ijgi5100189"),
    "torres2014": ("Torres-Sospedra et al., IPIN 2014", "IPIN.2014.7275492"),
    "khalaj2017": ("Khalajmehrabadi, Gatsis & Akopian, IEEE COMST 2017",
                   "COMST.2017.2671454"),
    "kaem2004": ("Kaemarungsi & Krishnamurthy, INFOCOM 2004", "INFCOM.2004.1356988"),
    "torres2015": ("Torres-Sospedra et al., ESWA 2015", "j.eswa.2015.08.013"),
    "adler2015": ("Adler et al., IPIN 2015", "IPIN.2015.7346749"),
    "shannon1948": ("Shannon, Bell Syst. Tech. J. 1948", "Shannon1948"),
    "sparck1972": ("Spärck Jones, J. Documentation 1972", "SparckJones1972"),
    "bray1957": ("Bray & Curtis, Ecol. Monogr. 1957", "BrayCurtis1957"),
    "shepard1968": ("Shepard, ACM Nat. Conf. 1968", "Shepard1968"),
    "snyder1987": ("Snyder, USGS Prof. Paper 1395, 1987", "Snyder1987"),
    "spearman1904": ("Spearman, Am. J. Psychol. 1904", "Spearman1904"),
    "jaccard1912": ("Jaccard, New Phytologist 1912", "Jaccard1912"),
    "efron1979": ("Efron, Ann. Statist. 1979", "Efron1979"),
}


def _crossref_entry(m: dict) -> str:
    def name(full: str) -> str:
        parts = full.split()
        if not parts:
            return full
        fam = parts[-1] if not full.isupper() else parts[-1].title()
        given = parts[:-1]
        if full.isupper():
            given = [g.title() for g in given]
            if len(parts) >= 3 and parts[-2].upper() == "SPARCK":
                fam, given = "Spärck Jones", given[:-1]
        ini = " ".join(g if g.endswith(".") else g[0] + "." for g in given)
        return f"{ini} {fam}".strip()
    au = [name(a) for a in m["authors"].split("; ")]
    au = au[0] if len(au) == 1 else ", ".join(au[:-1]) + " and " + au[-1]
    t = m["title"]
    t = t.capitalize() if t.isupper() else t
    t = t.rstrip(".")
    src = " ".join((m["container"] or m.get("publisher") or "").split()).rstrip(" -")
    if src.endswith(" on"):
        src = src[:-3]
    bits = [f'{au}, "{t}," {src}']
    if m.get("volume"):
        bits.append(f"vol. {m['volume']}")
    if m.get("issue"):
        bits.append(f"no. {m['issue']}")
    if m.get("page"):
        bits.append(f"pp. {m['page'].replace('-', '–')}")
    bits.append(str(m["year"]))
    return ", ".join(bits) + f". doi:{m['doi']}"


class Refs:
    """Numbers references in order of first citation."""

    def __init__(self) -> None:
        extra = load_json(PROJECT_ROOT / "research" / "doi_verification_formulas.json")
        self.checked = extra["checked_utc"][:10]
        self.entry, self.kind = {}, {}
        for key, (_, tag) in SHORT.items():
            if tag in extra["sources"]:
                self.entry[key] = _crossref_entry(extra["sources"][tag])
                self.kind[key] = "new"
            else:
                hit = [r for r in hr.REFS if tag in r]
                if len(hit) != 1:
                    raise KeyError(f"reference {key} not found in the verified list")
                self.entry[key] = hit[0]
                self.kind[key] = "project"
        self.order: list[str] = []

    def n(self, key: str) -> int:
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1

    def cite(self, *keys: str) -> str:
        return "[" + "], [".join(str(self.n(k)) for k in keys) + "]"

    def who(self, key: str) -> str:
        return f"{SHORT[key][0]} [{self.n(key)}]"


# --------------------------------------------------------------------- data
class Deck:
    """All numbers of the deck, read from the result files."""

    def __init__(self) -> None:
        self.D = D = hr.Data()
        self.S, self.cs = D.S, D.S["cross_split"]
        self.tr = D.S["training"]
        self.H, self.I, self.AB, self.CN = D.H, D.I, D.AB, D.CN
        self.base = D.base
        self.prim = {r["utility_model"]: r for _, r in D.primary.iterrows()}
        self.tot = D.totals
        self.O = D.O
        self.coord = D.coord["conclusion"]
        self.PO = D.PO
        self.s0 = D.AB[D.AB["stage"] == "S0"].iloc[0]
        self.sel = D.AB.iloc[-1]
        self.best_val = D.AB.loc[D.AB["val_mean_ground_m"].idxmin()]
        self.BS = hr.bootstrap(D)
        self.CDF = hr.cdf_table(D)
        self.U = hr.union_vs_controls(D)
        self.M = hr.matched_budgets(D)
        self.O5 = D.O["observation_5_huim_vs_controls"]["verdict_inputs"]
        self.O6 = D.O["observation_6_strong_vs_discriminative"]
        self.O9 = D.O["observation_9_ranking_vs_size"]
        self.O3 = D.O["observation_3_models_select_different_aps"]
        self.cost = D.COST.set_index("condition")
        self.hsum = D.hsum
        import re
        self.n_folds = int(re.search(r"GroupKFold\((\d+)\)",
                                     D.PO["protocol"]["selection_split"]).group(1))
        self.t_lim, self.max_huis = D.hsum["time_limit_s"], D.hsum["max_huis"]
        self.vchk = {n: load_json(RESULTS / "verification" / n)
                     for n in ("huim_checks.json", "positioning_checks.json")}
        self.excel = hr.excel_checks()
        self.fig09 = pd.read_csv(FIGURES / "fig09_feature_reduction_vs_accuracy.csv")
        self._v1_meta()
        self._example()
        self._efim_example()

    def _v1_meta(self) -> None:
        import pymupdf
        doc = pymupdf.open(V1_PDF)
        c = doc.metadata.get("creationDate", "")          # D:YYYYMMDDhhmmss
        self.v1_pages = doc.page_count
        self.v1_author = doc.metadata.get("author", "")
        self.v1_date = f"{c[8:10]}/{c[6:8]}/{c[2:6]}" if len(c) >= 10 else ""
        doc.close()

    def _example(self) -> None:
        """The fingerprint of the Version 1 slides 'Bước 1-2', recomputed."""
        from preprocess_uji import fit_preprocessor, joint_label, load_splits
        from utility_models import build_external_utility
        tr, _ = load_splits()
        prep = fit_preprocessor(tr)
        y = joint_label(tr)
        self.p = {m: build_external_utility(m[0], tr, prep, y, "training")
                  for m in MODELS}
        wap = prep.wap_cols
        W = tr[wap].to_numpy()
        det = W != 100
        # Version 1 used the row WAP255/WAP284/WAP310/WAP366: find it again
        want = ["WAP255", "WAP284", "WAP310", "WAP366"]
        idx = [wap.index(w) for w in want]
        hit = np.where((det.sum(axis=1) == len(want)) & det[:, idx].all(axis=1))[0]
        self.ex_index = int(hit[0])
        self.ex_line = self.ex_index + 2                    # CSV line, header = 1
        row = tr.iloc[[self.ex_index]]
        self.ex_tx = prep.transactions(row)[0][0]
        self.ex_raw = {w: int(row[w].iloc[0]) for w in self.ex_tx}
        self.ex_tu = {m: sum(q * self.p[m].external_utility[i]
                             for i, q in self.ex_tx.items()) for m in MODELS}
        self.prep = prep
        self.n_wap = len(wap)
        sup = det.sum(axis=0)
        self.max_sup_share = sup.max() / len(tr)
        self.max_sup_wap = wap[int(sup.argmax())]
        empty = tr[det.sum(axis=1) == 0]
        self.empty_by_phone = empty["PHONEID"].value_counts()
        self.n_distinct_items = int((sup > 0).sum())

    def _efim_example(self) -> None:
        sys.path.insert(0, str(PROJECT_ROOT / "tests"))
        from huim.brute_force import itemset_utility, twu
        from huim.efim import mine_efim
        from test_efim import RUNNING_DB as DB, RUNNING_P as P
        self.ex_mu = 30
        self.ex_db, self.ex_p = DB, P
        self.ex_twu = {i: twu(i, DB, P) for i in sorted(P)}
        order = sorted((i for i in self.ex_twu if self.ex_twu[i] >= self.ex_mu),
                       key=lambda i: (self.ex_twu[i], i))
        self.ex_order = order
        self.ex_su0 = {it: sum(sum(q * P[i] for i, q in t.items()
                                   if i in order and order.index(i) >= z)
                               for t in DB if it in t)
                       for z, it in enumerate(order)}
        self.ex_huis, self.ex_st = mine_efim(DB, P, self.ex_mu)
        self.u = lambda *xs: itemset_utility(tuple(xs), DB, P)

    # helpers
    def run(self, model, pct):
        return self.D.run(model, pct)

    def cond(self, name):
        return self.D.cond(name)


# ------------------------------------------------------------------- charts
def _save(fig, name: str, data: pd.DataFrame) -> Path:
    CHARTS.mkdir(parents=True, exist_ok=True)
    p = CHARTS / f"{name}.png"
    fig.savefig(p)
    plt.close(fig)
    data.to_csv(CHARTS / f"{name}.csv", index=False)
    return p


def _minutil_axis(ax):
    ticks = [20, 10, 5, 3, 2, 1.5]
    ax.set_xscale("log")
    ax.xaxis.set_major_locator(mticker.FixedLocator(ticks))
    ax.xaxis.set_major_formatter(mticker.FixedFormatter([mu(t) + "%" for t in ticks]))
    ax.xaxis.set_minor_locator(mticker.NullLocator())
    ax.set_xlim(23, 1.38)


def _model_lines(ax, H, col, label=True):
    for m in MODELS:
        g = H[H["utility_model"] == m].sort_values("rel_minutil_pct")
        c = MODEL_HEX[m]
        ok = g["status"] == "COMPLETED"
        ax.plot(g["rel_minutil_pct"], g[col], color=c, lw=2,
                label=f"Mô hình {MNAME[m]}" if label else None)
        ax.scatter(g.loc[ok, "rel_minutil_pct"], g.loc[ok, col], color=c, s=34, zorder=3)
        ax.scatter(g.loc[~ok, "rel_minutil_pct"], g.loc[~ok, col], s=70, zorder=4,
                   facecolors="white", edgecolors=c, linewidths=2)


def ch_hui(d: Deck, w, h) -> Path:
    H = d.H
    fig = fig_pt(w, h)
    a1, a2 = fig.subplots(2, 1)
    _model_lines(a1, H, "n_huis")
    _model_lines(a2, H, "unique_waps_in_huis", label=False)
    a1.set_yscale("symlog", linthresh=1)
    a1.set_ylabel("Số HUI (symlog)")
    a1.set_title("(a) Số HUI bùng nổ khi hạ minUtil", loc="left")
    a2.set_ylabel("Số WAP trong hợp các HUI")
    a2.set_title("(b) Số WAP chỉ lên tới " + f"{int(d.U['n_waps'].min())}–"
                 f"{int(d.U['n_waps'].max())} ở lần chạy xong", loc="left")
    for ax in (a1, a2):
        _minutil_axis(ax)
        ax.set_xlabel("minUtil (% tổng utility, giảm dần sang phải)")
    a1.scatter([], [], s=70, facecolors="white", edgecolors="#5B6676", linewidths=2,
               label="Điểm rỗng: dừng ở giới hạn 600 s / 1.000.000 HUI")
    a1.legend(loc="upper left")
    fig.tight_layout(h_pad=1.2)
    return _save(fig, "d01_minutil_hui_wap",
                 H[["utility_model", "rel_minutil_pct", "status", "n_huis",
                    "unique_waps_in_huis"]])


def ch_cost2(d: Deck, w, h) -> Path:
    rows = [d.run(m, 2.0) for m in MODELS]
    df = pd.DataFrame({"model": MODELS, "runtime_s": [r["runtime_s"] for r in rows],
                       "n_huis": [r["n_huis"] for r in rows]})
    fig = fig_pt(w, h)
    ax = fig.subplots()
    x = np.arange(3)
    ax.bar(x, df["runtime_s"], 0.55, color=[MODEL_HEX[m] for m in MODELS])
    for xi, v, n in zip(x, df["runtime_s"], df["n_huis"]):
        ax.text(xi, v + df["runtime_s"].max() * 0.02, f"{vn(v, 1)} s",
                ha="center", fontsize=16, fontweight="bold")
    ax.set_xticks(x)
    ax.set_xticklabels([f"Mô hình {MNAME[m]}\n{vi(n)} HUI" for m, n in
                        zip(MODELS, df["n_huis"])], fontsize=14)
    ax.set_ylabel("Thời gian chạy EFIM (giây)")
    ax.set_title("Thời gian chạy EFIM ở minUtil = 2% tổng utility", loc="left")
    ax.grid(axis="x", alpha=0)
    ax.set_ylim(0, df["runtime_s"].max() * 1.15)
    fig.tight_layout()
    return _save(fig, "d02_runtime_at_2pct", df)


def ch_runtime_mem(d: Deck, w, h) -> Path:
    H = d.H
    fig = fig_pt(w, h)
    a1, a2 = fig.subplots(2, 1)
    _model_lines(a1, H, "runtime_s")
    _model_lines(a2, H, "peak_memory_delta_mb", label=False)
    a1.set_yscale("log")
    a1.set_ylabel("Thời gian (s, log)")
    a1.set_title("(a) Thời gian khai phá", loc="left")
    a2.set_ylabel("Bộ nhớ đỉnh tăng thêm (MB)")
    a2.set_title("(b) Bộ nhớ đỉnh", loc="left")
    for ax in (a1, a2):
        _minutil_axis(ax)
        ax.set_xlabel("minUtil (% tổng utility, giảm dần sang phải)")
    a1.legend(loc="upper left")
    fig.tight_layout(h_pad=1.2)
    return _save(fig, "d03_minutil_runtime_memory",
                 H[["utility_model", "rel_minutil_pct", "status", "runtime_s",
                    "peak_memory_delta_mb", "search_nodes_visited"]])


def ch_union(d: Deck, w, h) -> Path:
    U = d.U
    fig = fig_pt(w, h)
    ax = fig.subplots()
    series = [("huim_union", "HUIM union", "#E2A72E"),
              ("top_mi", "Top-n theo MI", "#17797F"),
              ("top_support", "Top-n theo số lần phát hiện", "#14223A"),
              ("random_mean", "Ngẫu nhiên (TB 3 seed)", "#AEB6C2")]
    x = np.arange(len(U))
    bw = 0.2
    for k, (col, lab, c) in enumerate(series):
        xs = x + (k - 1.5) * bw
        ax.bar(xs, U[col] * 100, bw, color=c, label=lab)
        for xi, v in zip(xs, U[col] * 100):
            ax.text(xi, v + 0.6, vn(v, 1), ha="center", fontsize=12,
                    fontweight="bold" if col == "huim_union" else "normal")
    b = float(d.base["joint_building_floor_accuracy"]) * 100
    ax.axhline(b, color="#14223A", lw=1.4, ls="--")
    ax.text(len(U) - 0.5, b + 0.8, f"Baseline 520 WAP: {vn(b)}%", ha="right", fontsize=13)
    ax.set_xticks(x)
    ax.set_xticklabels([f"Mô hình {LETTER[m]} · {n} WAP\nminUtil {mu(p)}%"
                        for m, n, p in zip(U["model"], U["n_waps"], U["minutil_pct"])],
                       fontsize=13.5)
    lo = float(U[[s[0] for s in series]].min().min()) * 100
    ax.set_ylim(max(0, lo - 12), 95)
    ax.set_ylabel("Độ chính xác joint (%)")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.17), ncol=4, fontsize=12.5)
    ax.grid(axis="x", alpha=0)
    fig.tight_layout()
    return _save(fig, "d04_union_vs_controls", U)


def ch_featred(d: Deck, w, h) -> Path:
    F = d.fig09.copy()
    fig = fig_pt(w, h)
    ax = fig.subplots()
    rnd = (F[F["family"] == "control_random"]
           .groupby("n_waps_selected", as_index=False)
           .agg(feature_reduction_pct=("feature_reduction_pct", "first"),
                joint=("joint_building_floor_accuracy", "mean")))
    ax.plot(rnd["feature_reduction_pct"], rnd["joint"] * 100, color="#AEB6C2", lw=1.6,
            marker="x", ms=6, label="Ngẫu nhiên (TB 3 seed)")
    for fam, lab, c, mk in (("control_top_mi", "Top-n theo MI", "#17797F", "s"),
                            ("control_top_support", "Top-n theo số lần phát hiện",
                             "#14223A", "^")):
        g = F[F["family"] == fam].sort_values("feature_reduction_pct")
        ax.plot(g["feature_reduction_pct"], g["joint_building_floor_accuracy"] * 100,
                color=c, lw=1.6, marker=mk, ms=5, label=lab)
    for m in MODELS:
        g = F[(F["family"] == "huim_topn") & (F["utility_model"] == m)].sort_values(
            "feature_reduction_pct")
        p = d.prim[m]["rel_minutil_pct"]
        ax.plot(g["feature_reduction_pct"], g["joint_building_floor_accuracy"] * 100,
                color=MODEL_HEX[m], lw=2, marker="o", ms=5,
                label=f"HUIM top-n {LETTER[m]} (minUtil {mu(p)}%)")
        u = F[(F["family"] == "huim_union") & (F["utility_model"] == m)]
        ax.scatter(u["feature_reduction_pct"], u["joint_building_floor_accuracy"] * 100,
                   marker="D", s=58, facecolors="white", edgecolors=MODEL_HEX[m],
                   linewidths=1.8, zorder=4)
    ax.scatter([], [], marker="D", s=58, facecolors="white", edgecolors="#5B6676",
               linewidths=1.8, label="HUIM union (minUtil là núm duy nhất)")
    b = float(d.base["joint_building_floor_accuracy"]) * 100
    ax.axhline(b, color="black", lw=1.2, ls="--", label=f"Baseline 520 WAP ({vn(b)}%)")
    ax.set_xlabel("Giảm chiều (%) = (1 − số WAP chọn / 520) × 100")
    ax.set_ylabel("Độ chính xác joint (%)")
    ax.set_ylim(0, 100)
    ax.legend(loc="lower left", fontsize=11.5)
    fig.tight_layout()
    return _save(fig, "d05_feature_reduction_vs_joint", F)


def ch_ips_cost(d: Deck, w, h) -> Path:
    C = d.D.COST.sort_values("n_features")
    fig = fig_pt(w, h)
    a1, a2, a3 = fig.subplots(1, 3)
    styles = (("baseline", "Baseline", "*", "black"),
              ("control_top_mi", "Top-n MI", "s", "#17797F"),
              ("control_top_support", "Top-n phát hiện", "^", "#14223A"))
    a1.plot(C["n_features"], C["predict_ms_per_fingerprint"], color="#C8CDD5", lw=1, zorder=1)
    for fam, lab, mk, c in styles:
        g = C[C["family"] == fam]
        a1.scatter(g["n_features"], g["predict_ms_per_fingerprint"], marker=mk, s=40,
                   color=c, label=lab, zorder=3)
        a3.scatter(g["predict_ms_per_fingerprint"], g["joint_building_floor_accuracy"] * 100,
                   marker=mk, s=40, color=c, zorder=3)
    g = C[C["family"].str.startswith("huim")]
    a1.scatter(g["n_features"], g["predict_ms_per_fingerprint"], s=44, color="#E2A72E",
               label="HUIM", zorder=4)
    a3.scatter(g["predict_ms_per_fingerprint"], g["joint_building_floor_accuracy"] * 100,
               s=44, color="#E2A72E", zorder=4)
    a1.set_xlabel("Số WAP")
    a1.set_ylabel("ms / fingerprint (1 lõi)")
    a1.set_title("(a) Thời gian định vị", loc="left")
    a1.legend(fontsize=10.5, loc="upper left")
    for col, lab, ls, c in (("db_memory_float64_mb", "float64", "-", "#14223A"),
                            ("db_memory_float32_mb", "float32", "--", "#17797F"),
                            ("db_memory_int8_mb", "int8", ":", "#E2A72E")):
        a2.plot(C["n_features"], C[col], ls=ls, color=c, lw=2, marker="o", ms=3, label=lab)
    a2.set_xlabel("Số WAP")
    a2.set_ylabel("CSDL vân tay (MB)")
    a2.set_title("(b) Bộ nhớ CSDL", loc="left")
    a2.legend(fontsize=10.5, loc="upper left")
    a3.set_xlabel("ms / fingerprint")
    a3.set_ylabel("Joint (%)")
    a3.set_title("(c) Chi phí – độ chính xác", loc="left")
    for ax in (a1, a2, a3):
        ax.tick_params(labelsize=11)
        ax.xaxis.label.set_size(12)
        ax.yaxis.label.set_size(12)
        ax.title.set_size(13)
    fig.tight_layout(w_pad=1.2)
    return _save(fig, "d06_ips_time_memory",
                 C[["condition", "family", "n_features", "predict_ms_per_fingerprint",
                    "db_memory_float64_mb", "db_memory_float32_mb", "db_memory_int8_mb",
                    "joint_building_floor_accuracy"]])


V1_SETS = ["baseline_all_520", "control_top_support__n200", "control_top_mi__n91",
           "huim_union__C_rare_discriminative__1.7321pct",
           "huim_union__A_signal__1.9294pct", "huim_union__B_discriminative__2pct"]


def _set_name(d: Deck, c: str) -> str:
    r = d.cond(c)
    if c == "baseline_all_520":
        return "Baseline\n520 WAP"
    if c.startswith("control"):
        return f"Top-{int(r['n_waps_selected'])}\n" + ("MI" if "_mi_" in c else "phát hiện")
    return f"HUIM {LETTER[r['utility_model']]}\n{int(r['n_waps_selected'])} WAP"


def ch_v1_error(d: Deck, w, h) -> Path:
    rows = []
    for c in V1_SETS:
        r = d.cond(c)
        rows.append({"condition": c, "median": r["pos_median_error_ground_m"],
                     "mean": r["pos_mean_error_ground_m"],
                     "p90": r["pos_p90_error_ground_m"]})
    df = pd.DataFrame(rows)
    fig = fig_pt(w, h)
    ax = fig.subplots()
    x = np.arange(len(df))
    bw = 0.26
    for k, (col, lab, c) in enumerate((("median", "Trung vị", "#17797F"),
                                       ("mean", "Trung bình", "#E2A72E"),
                                       ("p90", "Phân vị 90", "#14223A"))):
        xs = x + (k - 1) * bw
        ax.bar(xs, df[col], bw, color=c, label=lab)
        for xi, v in zip(xs, df[col]):
            ax.text(xi, v + 0.8, vn(v, 1), ha="center", fontsize=9.5, rotation=90,
                    va="bottom")
    ax.axhspan(7, 8, color="#2E7D32", alpha=0.18, lw=0, label="Mức tham chiếu 7–8 m")
    ax.set_xticks(x)
    ax.set_xticklabels([_set_name(d, c) for c in df["condition"]], fontsize=11)
    ax.set_ylabel("Sai số (m thực địa)")
    ax.set_ylim(0, float(df["p90"].max()) * 1.28)
    ax.legend(loc="upper left", ncol=2, fontsize=11)
    ax.grid(axis="x", alpha=0)
    ax.set_title("Sai số định vị Version 1 (k-NN, k = 1)", loc="left", fontsize=14)
    fig.tight_layout()
    return _save(fig, "d07_v1_error_by_set", df)


def ch_stages(d: Deck, w, h) -> Path:
    A = d.AB
    fig = fig_pt(w, h)
    ax = fig.subplots()
    x = np.arange(len(A))
    bw = 0.37
    ax.bar(x - bw / 2, A["dev_mean_ground_m"], bw, color="#AEB6C2",
           label="Dev: 6 fold giữ lại người đo (dùng để CHỌN)")
    ax.bar(x + bw / 2, A["val_mean_ground_m"], bw, color="#14223A",
           label="validationData (chỉ để BÁO CÁO)")
    for xi, v in zip(x - bw / 2, A["dev_mean_ground_m"]):
        ax.text(xi, v + 0.2, vn(v), ha="center", fontsize=12)
    for xi, v in zip(x + bw / 2, A["val_mean_ground_m"]):
        ax.text(xi, v + 0.2, vn(v), ha="center", fontsize=12.5, fontweight="bold")
    ax.axhspan(7, 8, color="#2E7D32", alpha=0.16, lw=0)
    ax.text(len(A) - 0.42, 7.5, "mức tham chiếu 7–8 m", va="center", ha="right",
            fontsize=12, color="#1B5E20")
    lab = {"S0": "S0\nk-NN V1", "S1": "S1\nkhoảng cách", "S2": "S2\n+ k, 1/d$^p$",
           "S3": "S3\n+ dò lệch", "S4": "S4\n+ ensemble"}
    ax.set_xticks(x)
    ax.set_xticklabels([lab.get(s, s) for s in A["stage"]], fontsize=13)
    ax.set_ylabel("Sai số trung bình (m thực địa)")
    ax.set_ylim(0, float(A[["dev_mean_ground_m", "val_mean_ground_m"]].max().max()) * 1.18)
    ax.legend(loc="upper right", fontsize=12.5)
    ax.grid(axis="x", alpha=0)
    fig.tight_layout()
    return _save(fig, "d08_v2_stages", A)


def ch_cdf(d: Deck, w, h) -> Path:
    D = d.D
    fig = fig_pt(w, h)
    ax = fig.subplots()
    for name, P, c, lw, ls in (("Version 1", D.Pold, "#8A94A0", 2.2, "--"),
                               ("Version 2", D.Pnew, "#14223A", 2.8, "-")):
        e = np.sort(P["err_ground_m"].to_numpy())
        ax.plot(e, np.arange(1, len(e) + 1) / len(e) * 100, color=c, lw=lw, ls=ls,
                label=f"{name}: TB {vn(e.mean())} m · trung vị {vn(np.median(e))} m")
    ax.axvspan(7, 8, color="#2E7D32", alpha=0.16, lw=0)
    ax.text(7.5, 3, "7–8 m", ha="center", fontsize=12, color="#1B5E20")
    ax.set_xlim(0, 40)
    ax.set_ylim(0, 100)
    ax.set_xlabel("Sai số định vị (m thực địa)")
    ax.set_ylabel("Tỉ lệ fingerprint có sai số ≤ x (%)")
    ax.legend(loc="lower right", fontsize=13)
    fig.tight_layout()
    data = pd.DataFrame({"err_v1_sorted": np.sort(D.Pold["err_ground_m"]),
                         "err_v2_sorted": np.sort(D.Pnew["err_ground_m"])})
    return _save(fig, "d09_v1_v2_error_cdf", data)


def ch_sets(d: Deck, w, h) -> Path:
    C = d.CN.set_index("condition").loc[hr.SET_ORDER].reset_index()
    labels = [hr.set_label(d.D, r) for _, r in C.iterrows()]
    fig = fig_pt(w, h)
    ax = fig.subplots()
    y = np.arange(len(C))[::-1]
    ax.axvspan(7, 8, color="#2E7D32", alpha=0.16, lw=0)
    for yi, a, b in zip(y, C["study_mean_ground_m"], C["opt_mean_ground_m"]):
        ax.plot([b, a], [yi, yi], color="#C8CDD5", lw=3, zorder=1)
        ax.scatter(a, yi, color="#8A94A0", s=70, zorder=3)
        ax.scatter(b, yi, color="#14223A", s=70, zorder=4)
        ax.text(a + 0.5, yi, vn(a), va="center", fontsize=12, color="#5B6676")
        ax.text(b - 0.5, yi, vn(b), va="center", ha="right", fontsize=12,
                fontweight="bold")
    ax.scatter([], [], color="#8A94A0", s=70, label="Version 1 (k-NN, k = 1)")
    ax.scatter([], [], color="#14223A", s=70, label="Version 2")
    ax.set_yticks(y)
    ax.set_yticklabels(labels, fontsize=12)
    ax.set_xlim(0, float(C["study_mean_ground_m"].max()) * 1.15)
    ax.set_xlabel("Sai số trung bình (m thực địa); dải xanh = 7–8 m")
    ax.legend(loc="upper right", fontsize=12.5)
    ax.grid(axis="y", alpha=0)
    fig.tight_layout()
    return _save(fig, "d10_v2_by_wap_set", C)


# ------------------------------------------------------------------- slides
class Builder:
    def __init__(self, d: Deck, bg: dict[str, Path]) -> None:
        self.d, self.bg = d, bg
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(1440), Pt(810)
        self.R = Refs()
        self.F = Formulas(ASSETS / "formulas")
        self.slides: list[tuple[str, object]] = []

    def slide(self, bg: str, name: str):
        s = new_slide(self.prs, self.bg[bg])
        self.slides.append((name, s))
        return s

    # ---- shared pieces
    def source(self, s, text, x, y, w, size=14.5, color=MUTED, align="l"):
        return tb(s, x, y, w, 40, "{i:Nguồn:} " + text, size=size, color=color,
                  italic=True, align=align, spacing=1.0)

    def fcard(self, s, x, y, w, h, head, tex_lines, note, src, fsize=26):
        """Formula card: gold header, formula image(s), note, source."""
        card(s, x, y, w, h, fill=CARD_FILL, line=CARD_LINE, radius=0.05)
        tb(s, x + 22, y + 16, w - 44, 30, head, size=20, color=GOLD, bold=True)
        yy = y + 58
        for tex in tex_lines:
            fw, fh = self.F.put(s, tex, x + 22, yy, size=fsize, max_w=w - 44,
                                align="c", box_w=w - 44)
            yy += fh + 8
        if note:
            tb(s, x + 22, yy + 4, w - 44, 60, note, size=16.5, color=SOFT, spacing=1.05)
        self.source(s, src, x + 22, y + h - 44, w - 44, size=14)

    def page_numbers(self):
        n = len(self.slides)
        for k, (_, s) in enumerate(self.slides, start=1):
            if k == 1:
                continue
            tb(s, 1330, 782, 90, 20, f"{k} / {n}", size=11.5, color=MUTED, align="r")

    # ---------------------------------------------------------------- 1 cover
    def s_cover(self):
        d = self.d
        s = self.slide("cover", "Bìa")
        tb(s, 46, 62, 660, 330, ["High Utility", "Itemset Mining",
                                 {"text": "{g:(HUIM)}", "size": 83}],
           size=90, bold=True, spacing=0.92, shadow=True)
        tb(s, 50, 410, 610, 120,
           "Ứng dụng chọn Access Point cho định vị trong nhà bằng vân tay Wi-Fi "
           "— thực nghiệm trên UJIIndoorLoc", size=25, color=SOFT, spacing=1.12)
        card(s, 50, 612, 400, 48, fill=GOLD, line=None, radius=0.5)
        tb(s, 50, 612, 400, 48, f"VERSION 2  ·  cập nhật {dt.date.today():%d/%m/%Y}",
           size=19, color=INK, bold=True, align="c", anchor="m")
        tb(s, 52, 680, 620, 60,
           f"Thay thế Version 1 (bản Canva ngày {d.v1_date}). Mọi số liệu sinh tự "
           f"động từ results/ của dự án.", size=16.5, color=MUTED, spacing=1.1)

    # ---------------------------------------------------------- 2 what's new
    def s_new(self):
        d, s0, sel = self.d, self.d.s0, self.d.sel
        s = self.slide("plain", "Version 2 có gì mới?")
        title(s, "{g:Version 2} có gì mới?", size=60)
        tb(s, 48, 118, 1300, 30, f"So với bản trình bày Version 1 (PDF Canva ngày "
           f"{d.v1_date}, {d.v1_pages} trang) — số liệu trên {vi(len(d.D.Pnew))} "
           f"fingerprint của validationData", size=20, color=MUTED)
        k = [("Sai số trung bình", f"{vn(s0['val_mean_ground_m'])} → "
              f"{{g:{vn(sel['val_mean_ground_m'])}}} m", "mét thực địa"),
             ("Sai số trung vị", f"{vn(s0['val_median_ground_m'])} → "
              f"{{g:{vn(sel['val_median_ground_m'])}}} m", "mét thực địa"),
             ("Độ chính xác tòa nhà + tầng",
              f"{pc(s0['val_joint_building_floor_accuracy'])} → "
              f"{{g:{pc(sel['val_joint_building_floor_accuracy'])}}}", "joint accuracy"),
             ("Fingerprint sai số ≤ 8 m", f"{pc(s0['val_share_within_8m'], 1)} → "
              f"{{g:{pc(sel['val_share_within_8m'], 1)}}}", "tỉ lệ trên tập kiểm định")]
        for i, (h, v, sub) in enumerate(k):
            x = 46 + i * 340
            card(s, x, 168, 322, 172, radius=0.08)
            tb(s, x + 20, 184, 282, 28, h, size=18, color=SOFT, bold=True)
            tb(s, x + 20, 224, 290, 60, v, size=29, bold=True)
            tb(s, x + 20, 292, 282, 30, sub, size=15, color=MUTED, italic=True)
        n_runs = len(d.H)
        g = d.PO["selected_config"]
        card(s, 46, 368, 660, 395, radius=0.04)
        tb(s, 70, 386, 610, 30, "GIỮ NGUYÊN TỪ VERSION 1", size=20, color=GOLD, bold=True)
        tb(s, 70, 428, 615, 330, [
            f"Dữ liệu UJIIndoorLoc, ánh xạ HUIM ↔ IPS, q(i,T) và ba mô hình p(i)",
            f"EFIM và {n_runs} lần khai phá; {len(d.I)} điều kiện đánh giá IPS",
            "Giao thức chống rò rỉ: mọi tham số học trên trainingData",
            "Kết luận về HUIM: tập WAP do HUIM chọn vẫn thua top-n theo MI"],
           size=20.5, bullet="–", gap=12, color=WHITE)
        card(s, 734, 368, 660, 395, radius=0.04)
        tb(s, 758, 386, 610, 30, "THAY ĐỔI VÀ BỔ SUNG", size=20, color=GOLD, bold=True)
        tb(s, 758, 428, 615, 330, [
            f"Bộ ước lượng toạ độ mới: {g['representation']} + Bray-Curtis, k = {g['k']}, "
            f"w = 1/d{sup(g['power'])}, dò lệch ±{g['offset_range']} dB, ensemble",
            f"Chọn mô hình trên {d.n_folds} fold giữ lại người đo (GroupKFold theo "
            f"USERID)",
            f"Kiểm chứng độc lập: {d.vchk['huim_checks.json']['n_checks']} phép HUIM, "
            f"{d.vchk['positioning_checks.json']['n_checks']} phép định vị, Excel tính tay",
            f"Mọi công thức ghi rõ tác giả / nguồn; đính chính {len(self.fix_rows()) - 1} "
            f"chỗ sai của slide V1",
            "Biểu đồ vẽ lại, in minUtil; số theo định dạng Việt Nam (1.234,56)"],
           size=19, bullet="–", gap=9, color=WHITE)

    # ------------------------------------------------------------ 3 HUIM là gì
    def s_what(self):
        R = self.R
        s = self.slide("what", "HUIM là gì?")
        tb(s, 46, 62, 520, 110, "HUIM là gì?", size=82, bold=True, shadow=True)
        tb(s, 46, 196, 500, 300, [
            "High Utility Itemset Mining {g:(HUIM)}: tìm các tập mục có "
            "{g:tổng utility cao} — kết hợp số lượng và giá trị của từng mục.",
            {"text": "Khác Frequent Itemset Mining chỉ đếm tần suất xuất hiện, "
                     "HUIM ưu tiên giá trị mang lại.", "gap": 14}], size=26,
           spacing=1.12)
        tb(s, 81, 552, 420, 34, "{g:Ví dụ siêu thị}", size=26)
        tb(s, 78, 598, 425, 120, ["nhiều chai nước suối → lời thấp mỗi chai",
                                  "1 chai nước hoa → lời cao"], size=23, gap=6)
        tb(s, 645, 113, 330, 34, "FREQUENT ITEMSET MINING", size=22, bold=True)
        tb(s, 1030, 112, 345, 34, "HIGH UTILITY ITEMSET MINING", size=20.5,
           bold=True, color=INK)
        rows = [((214, 280), "MỤC TIÊU", "Tập mục xuất hiện\nnhiều lần",
                 "Tập mục có tổng\nutility cao"),
                ((358, 426), "TIÊU CHÍ", "sup(X) ≥ minSup", "u(X) ≥ minUtil"),
                ((489, 555), "KẾT QUẢ", "Phổ biến nhưng có thể\nlợi nhuận thấp",
                 "Ít phổ biến nhưng\nlợi nhuận cao"),
                ((629, 698), "VÍ DỤ", "Bánh mì & sữa\n(rất thường xuyên)",
                 "Rượu vang & phô mai\n(ít nhưng lời cao)")]
        for (y0, y1), lab, a, b in rows:
            ly = {"MỤC TIÊU": 230, "TIÊU CHÍ": 368, "KẾT QUẢ": 502, "VÍ DỤ": 644}[lab]
            tb(s, 935, ly, 150, 26, f"{{g:{lab}}}", size=19, align="c")
            tb(s, 640, y0, 268, y1 - y0, a, size=23, align="c", anchor="m", spacing=1.05)
            tb(s, 1085, y0, 300, y1 - y0, b, size=23, align="c", anchor="m", spacing=1.05)
        self.source(s, f"FIM — {R.who('agrawal1993')}; HUIM — {R.who('yao2004')}",
                    46, 768, 1350, size=14)

    # -------------------------------------------------------------- 4 HUIM-IPS
    def s_ips(self):
        s = self.slide("ips", "HUIM và IPS")
        tb(s, 46, 74, 600, 240, ["HUIM và", "Hệ thống Định vị",
                                 "{g:Trong nhà (IPS)}"], size=64, bold=True,
           spacing=0.98, shadow=True)
        tb(s, 46, 362, 560, 340, [
            "{g:Wi-Fi fingerprinting} ghi cường độ tín hiệu (RSSI) từ nhiều AP tại "
            "các vị trí đã biết.",
            {"text": "Số AP rất lớn ({b:520} cột WAP) nhưng không phải AP nào cũng hữu "
                     "ích cho định vị.", "gap": 18},
            {"text": "Câu hỏi: nhóm AP có {g:tổng utility cao} do HUIM tìm ra có phải "
                     "tập đặc trưng tốt? — trả lời bằng thực nghiệm có đối chứng.",
             "gap": 18}], size=23, spacing=1.12)
        for x, t in ((805, "VỊ TRÍ A"), (1021, "VỊ TRÍ B"), (1234, "VỊ TRÍ C")):
            tb(s, x - 10, 97, 100, 26, t, size=20)
        for y, t in ((302, "AP 1"), (384, "AP 2"), (472, "AP 3"), (595, "AP N")):
            tb(s, 636, y, 70, 30, t, size=23, bold=True)
        tb(s, 683, 694, 160, 24, "Cường độ tín hiệu", size=17)
        tb(s, 888, 694, 60, 24, "Thấp", size=17)
        tb(s, 1300, 694, 60, 24, "Cao", size=17)

    # --------------------------------------------------- 5 HUIM definitions
    def s_defs(self):
        d, R = self.d, self.R
        s = self.slide("plain", "Định nghĩa toán học")
        title(s, "Định nghĩa toán học của {g:HUIM}", size=56)
        W2, H2 = 664, 272
        self.fcard(s, 46, 128, W2, H2, "1 · Utility của mục và của tập mục",
                   [r"u(i,T)=q(i,T)\cdot p(i)\qquad u(X,T)=\sum_{i\in X}u(i,T)",
                    r"u(X)=\sum_{T\in D,\;X\subseteq T}u(X,T)"],
                   "q: lượng (utility nội) · p: giá trị đơn vị (utility ngoài)",
                   f"{R.who('yao2004')}; tổng quan {R.who('gan2021')}")
        tot = d.tot
        self.fcard(s, 730, 128, W2, H2, "2 · Tập hữu ích cao (HUI) và minUtil",
                   [r"X\in\mathrm{HUIs}\Leftrightarrow u(X)\geq\mathrm{minUtil}",
                    r"\mathrm{minUtil}=\theta\cdot\sum_{T\in D}\mathrm{TU}(T)"],
                   f"θ theo % vì tổng utility chênh {vn(tot.max() / tot.min())} lần "
                   f"giữa ba mô hình p(i)",
                   f"HUI — {R.who('yao2004')}; θ theo % tổng utility — quy ước của "
                   f"bài tìm hiểu")
        self.fcard(s, 46, 418, W2, H2, "3 · Cận trên TWU (phản đơn điệu)",
                   [r"\mathrm{TU}(T)=\sum_{i\in T}u(i,T)",
                    r"\mathrm{TWU}(X)=\sum_{T\supseteq X}\mathrm{TU}(T)\;\geq\;u(Y)"
                    r"\quad\forall\,Y\supseteq X"],
                   "TWU < minUtil ⇒ loại X và mọi tập chứa X",
                   R.who("liu2005"))
        self.fcard(s, 730, 418, W2, H2, "4 · Độ hỗ trợ (Frequent Itemset Mining)",
                   [r"\mathrm{sup}(X)=\left|\{\,T\in D:\ X\subseteq T\,\}\right|"],
                   "FIM coi mọi mục ngang nhau; HUIM thêm lượng q và giá trị p",
                   R.who("agrawal1993"))
        ub, ubd, ubdf = d.u("b"), d.u("b", "d"), d.u("b", "d", "f")
        tb(s, 46, 708, 1350, 70,
           f"{{g:Utility không đơn điệu}} (ví dụ chuẩn của EFIM {R.cite('zida2017')}): "
           f"u({{b}}) = {vi(ub)} < u({{b,d}}) = {vi(ubd)} nhưng u({{b,d,f}}) = "
           f"{vi(ubdf)} < u({{b,d}}) ⇒ không cắt tỉa kiểu Apriori bằng u(X), phải dùng "
           f"cận trên (TWU, lu, su).", size=18.5, color=WHITE, spacing=1.1)

    # ------------------------------------------------- 6 utility modelling
    def s_model(self):
        R = self.R
        s = self.slide("model", "Mô hình hóa Utility trong IPS")
        tb(s, 50, 52, 1100, 90, "Mô hình hóa {g:Utility} trong IPS", size=67,
           bold=True, shadow=True)
        boxes = [((68, 343), "GIAO DỊCH T", "một fingerprint\n= một lần quét Wi-Fi"),
                 ((411, 669), "ITEM i", "một WAP được phát\nhiện (RSSI ≠ 100)"),
                 ((736, 1012), "UTILITY", "u(i,T) = q(i,T) × p(i)\nRSSI chuẩn hoá ×\ngiá trị của AP"),
                 ((1080, 1354), "HUI", "nhóm WAP có tổng\nutility ≥ minUtil")]
        for (x0, x1), h, t in boxes:
            tb(s, x0 + 8, 176 if h == "UTILITY" else 190, x1 - x0 - 16, 30,
               f"{{g:{h}}}", size=22, align="c")
            tb(s, x0 + 8, 214 if h == "UTILITY" else 228, x1 - x0 - 16, 130, t,
               size=20, align="c", spacing=1.08)
        tb(s, 60, 389, 1310, 34, "{g:BA MÔ HÌNH UTILITY — DÙNG CHUNG q(i,T), CHỈ KHÁC p(i)}",
           size=24, align="c")
        cols = [((51, 441), "MODEL A", "TÍN HIỆU (SIGNAL)",
                 "p(i) = 1: utility chỉ do cường\nđộ RSSI và đồng xuất hiện"),
                ((441, 895), "MODEL B", "MUTUAL INFORMATION",
                 "p(i) ∝ I(WAP i; Y): AP cho biết nhiều\nvề tòa nhà–tầng thì giá trị cao"),
                ((895, 1372), "MODEL C", "MI × ĐỘ HIẾM",
                 "p(i) ∝ I(WAP i; Y) · log(N / (1 + support)):\nphạt AP có mặt ở khắp nơi")]
        for (x0, x1), h, sub, t in cols:
            tb(s, x0 + 10, 456, x1 - x0 - 20, 34, f"{{g:{h}}}", size=26, align="c")
            tb(s, x0 + 10, 496, x1 - x0 - 20, 34, sub, size=23, align="c")
            tb(s, x0 + 14, 566, x1 - x0 - 28, 90, t, size=18.5, align="c", spacing=1.08)
        tb(s, 90, 684, 1260, 80,
           "{g:MỤC TIÊU:} kiểm tra nhóm AP có tổng utility cao có phải tập đặc trưng "
           "tốt cho định vị không — đo độ chính xác tòa nhà/tầng, sai số mét, thời "
           "gian, bộ nhớ, và so có đối chứng ở cùng số WAP.", size=20, align="c",
           anchor="m", spacing=1.08)

    # --------------------------------------------------------- 7 three p(i)
    def s_pdef(self):
        d, R = self.d, self.R
        s = self.slide("pdef", "Ba cách định nghĩa p(i)")
        tb(s, 56, 40, 1100, 110, "Ba cách định nghĩa p(i)", size=80, color=GOLD,
           shadow=True)
        dB = d.p["B_discriminative"].diagnostics
        specs = [
            ("A", "Model A", "Signal", "A_signal", r"p(i)=1",
             "Chỉ nhìn cường độ sóng và mức đồng xuất hiện: u(X) lớn khi các AP mạnh "
             "và hay xuất hiện cùng nhau.",
             f"Tổng utility CSDL: {vi(d.tot['A_signal'])}", "Đối chứng",
             "điều kiện đối chứng do bài tìm hiểu đặt ra (utility chỉ còn q)"),
            ("B", "Model B", "Discriminative", "B_discriminative",
             r"p(i)=\dfrac{\max\left(I(A_i;Y),\,10^{-6}\right)}{\max_j\,I(A_j;Y)}",
             "WAP nào giúp biết đang ở tòa nhà–tầng nào (Y) thì có giá trị cao.",
             f"Tổng utility CSDL: {vi(d.tot['B_discriminative'])} · "
             f"p trung bình {vn(dB['p_mean'], 3)}", "Model chính",
             f"I(·;·) — {R.who('shannon1948')}; chọn AP theo thông tin tương hỗ — "
             f"{R.who('chen2006')}; chuẩn hoá về (0, 1] — bài tìm hiểu"),
            ("C", "Model C", "Rare-discriminative", "C_rare_discriminative",
             r"p(i)\ \propto\ I(A_i;Y)\cdot\log\dfrac{N}{1+\mathrm{support}_i}",
             "Như B, nhưng phạt WAP có mặt ở khắp nơi (thừa số kiểu IDF).",
             f"Tổng utility CSDL: {vi(d.tot['C_rare_discriminative'])}",
             "Biến thể thứ ba",
             f"log(N / số lần xuất hiện) là IDF — {R.who('sparck1972')}; ghép với MI "
             f"— bài tìm hiểu"),
        ]
        cw, ch, y0 = 436, 588, 182
        for k, (L, name, sub, m, tex, desc, fact, role, src) in enumerate(specs):
            x = 46 + k * (cw + 16)
            card(s, x, y0, cw, ch, fill=WHITE, line=None, radius=0.035)
            card(s, x + 22, y0 + 22, 72, 72, fill=MODEL_HEX[m], line=None,
                 shape=MSO_SHAPE.OVAL)
            tb(s, x + 22, y0 + 22, 72, 72, L, size=40, bold=True, align="c", anchor="m")
            tb(s, x + 110, y0 + 30, cw - 130, 30, name, size=22, bold=True, color=INK)
            tb(s, x + 110, y0 + 60, cw - 130, 30, sub, size=19, color="#5B6676")
            self.F.put(s, tex, x + 22, y0 + 122, size=25, color="#14223A",
                       max_w=cw - 44, align="c", box_w=cw - 44)
            tb(s, x + 22, y0 + 232, cw - 44, 110, desc, size=18, color=INK, spacing=1.1)
            tb(s, x + 22, y0 + 350, cw - 44, 30, fact, size=15.5, color="#3B4656")
            tb(s, x + 22, y0 + 392, cw - 44, 34, role, size=21, bold=True,
               color=MODEL_HEX[m])
            tb(s, x + 22, y0 + 440, cw - 44, 140, "{i:Nguồn:} " + src, size=14.5,
               color="#4A5566", italic=True, spacing=1.05)

    # ------------------------------------------------ 8 preprocessing overview
    def s_prep(self):
        d, cs = self.d, self.d.cs
        s = self.slide("plain", "Tổng quan 5 bước tiền xử lý")
        title(s, "Tổng quan: {g:5 bước} tiền xử lý", size=58)
        tb(s, 48, 116, 1300, 30, "Từ một dòng CSV đến đầu vào của thuật toán EFIM",
           size=21, color=MUTED)
        steps = [("01", "Làm sạch", "100 = không phát hiện; kiểm tra RSSI nằm trong "
                                    "[−104, 0] dBm"),
                 ("02", "Chuẩn hoá RSSI", "q(i,T) ∈ [1, 100], RSSI min/max học trên "
                                          "trainingData"),
                 ("03", "Tạo transaction", "item = WAP; ô ≠ 100 → WAP có mặt trong "
                                           "giao dịch"),
                 ("04", "Trọng số p(i)", "Model A / B / C, chỉ ước lượng trên "
                                         "trainingData"),
                 ("05", "Bên trong EFIM", "TWU, sắp thứ tự, cắt tỉa TWU, gộp giao "
                                          "dịch")]
        cw, gap = 248, 33
        for k, (n, h, t) in enumerate(steps):
            x = 40 + k * (cw + gap)
            card(s, x, 176, cw, 250, radius=0.07)
            tb(s, x + 18, 190, cw - 36, 40, f"{{g:{n}}}", size=30)
            tb(s, x + 18, 236, cw - 36, 34, h, size=21, bold=True)
            tb(s, x + 18, 276, cw - 36, 140, t, size=17.5, color=SOFT, spacing=1.08)
            if k < 4:
                arrow(s, x + cw + 4, 290, 25, 20)
        ndp = d.tr["not_detected_probe"]
        tb(s, 52, 462, 1330, 300, [
            f"{{g:Đầu vào:}} trainingData.csv ({vi(cs['train_rows'])} dòng) và "
            f"validationData.csv ({vi(cs['validation_rows'])} dòng), mỗi file "
            f"{d.n_wap} cột WAP.",
            f"{{g:Đầu ra:}} danh sách transaction {{WAP: q}} và vector p(i) cho từng WAP.",
            f"{{g:Không làm đầy (impute) giá trị nào:}} 100 = không phát hiện = item vắng "
            f"mặt; {vn(ndp['share_eq_100_pct'])}% số ô WAP bằng 100, và "
            f"{vi(ndp['count_in_1_to_99'])} ô nằm trong 1…99.",
            "{g:Chống rò rỉ:} mọi tham số học từ dữ liệu (RSSI min/max, p(i)) chỉ lấy từ "
            "trainingData; hàm ước lượng p(i) báo lỗi nếu nhận dữ liệu khác."],
           size=20.5, bullet="–", gap=12)

    # ----------------------------------------------------- 9 step 1: RSSI
    def s_step1(self):
        d, R = self.d, self.R
        pp = d.hsum["preprocessor"]
        s = self.slide("plain", "Bước 1: chuẩn hoá RSSI")
        tb(s, 46, 30, 1360, 150, ["Bước 1: chuẩn hoá RSSI sang",
                                  "“positive representation”"], size=58, color=GOLD,
           spacing=0.98, shadow=True)
        tb(s, 46, 196, 640, 560, [
            "Không phát hiện (100) → 0 (“không có thông tin”), tức WAP vắng mặt khỏi "
            "giao dịch.",
            "Đã phát hiện: đưa RSSI về [0, 1] theo min/max của trainingData, rồi ánh "
            "xạ sang [1, 100].",
            "Làm tròn thành số nguyên dương; không ô đã phát hiện nào bị làm tròn về 0.",
            "q(i,T) là số nguyên dương → EFIM cần utility dương cho các cận trên.",
            f"Q_MAX = {pp['q_max']} chọn để dễ đọc, không phải siêu tham số được tinh "
            f"chỉnh."], size=22, bullet="•", gap=12)
        card(s, 720, 190, 674, 190, radius=0.05)
        tb(s, 742, 204, 630, 30, "Công thức (utility nội)", size=19, color=GOLD, bold=True)
        self.F.put(s, r"q(i,T)=\mathrm{round}\left(1+\dfrac{\mathrm{RSSI}_i-"
                      r"\mathrm{RSSI}_{\min}}{\mathrm{RSSI}_{\max}-\mathrm{RSSI}_{\min}}"
                      r"\times 99\right)", 742, 240, size=24, max_w=630, align="c",
                   box_w=630)
        self.source(s, f"biểu diễn “positive” — {R.who('torres2015')}; thang 1…100 "
                       f"và làm tròn — bài tìm hiểu", 742, 336, 630, size=14)
        rows = [["WAP", "RSSI thô (dBm)", "q(i,T)"]]
        for w_, q in d.ex_tx.items():
            rows.append([w_, vi(d.ex_raw[w_]).replace("-", "−"), str(q)])
        table(s, 720, 404, [230, 222, 222], rows, size=19, row_h=40)
        tb(s, 720, 404 + 40 * len(rows) + 12, 674, 90,
           f"Dòng {vi(d.ex_line)} của trainingData.csv (chỉ số pandas {vi(d.ex_index)}). "
           f"RSSI_min = {vi(pp['rssi_min_train']).replace('-', '−')}, RSSI_max = "
           f"{vi(pp['rssi_max_train'])} ước lượng trên {vi(pp['fitted_on_n_rows'])} "
           f"dòng trainingData.", size=15.5, color=MUTED, spacing=1.08)

    # ------------------------------------------- 10 step 2: transactions
    def s_step2(self):
        d = self.d
        db = d.hsum["database"]
        s = self.slide("plain", "Bước 2: transaction")
        tb(s, 46, 30, 1360, 150, ["Bước 2: mỗi fingerprint thành một",
                                  "utility transaction"], size=58, color=GOLD,
           spacing=0.98, shadow=True)
        n_det = len(d.ex_tx)
        tx = ", ".join(f"{w}: {q}" for w, q in d.ex_tx.items())
        flow = [("Dòng CSV " + vi(d.ex_line), f"{d.n_wap} cột WAP\n{n_det} cột ≠ 100, "
                 f"{d.n_wap - n_det} cột = 100"),
                ("Transaction", "{" + tx + "}"),
                ("Sau khi có p(i)", f"Model A: TU = {vn(d.ex_tu['A_signal'], 0)}\n"
                 f"Model B: TU ≈ {vn(d.ex_tu['B_discriminative'], 1)}\n"
                 f"Model C: TU ≈ {vn(d.ex_tu['C_rare_discriminative'], 1)}")]
        for k, (h, t) in enumerate(flow):
            x = 46 + k * 462
            card(s, x, 186, 420, 170, radius=0.07,
                 line=[CARD_LINE, GOLD, "#5FC98B"][k], lw=1.8)
            tb(s, x + 18, 200, 384, 30, h, size=20, bold=True, color=GOLD, align="c")
            tb(s, x + 18, 238, 384, 110, t, size=18.5, align="c", spacing=1.1)
            if k < 2:
                arrow(s, x + 426, 260, 30, 22)
        rows = [["Thống kê (trainingData)", "Giá trị"],
                ["Số dòng (fingerprint)", vi(db["n_rows"])],
                ["Transaction không rỗng", vi(db["n_non_empty_transactions"])],
                ["Transaction rỗng (không phát hiện AP nào)",
                 vi(db["n_empty_fingerprints_dropped"])],
                ["Độ dài trung bình (không rỗng)",
                 vn(db["detections_per_transaction_mean"])],
                ["Số item thực sự xuất hiện", f"{d.n_distinct_items} / {d.n_wap}"]]
        table(s, 46, 390, [520, 200], rows, size=18.5, row_h=56)
        ph = ", ".join(f"PHONEID {int(k)}: {int(v)}" for k, v in d.empty_by_phone.items())
        tb(s, 800, 396, 594, 360, [
            f"{vi(db['n_empty_fingerprints_dropped'])} dòng rỗng tập trung ở vài thiết bị "
            f"({ph}) — nghi lỗi thu thập; vẫn giữ cho bộ phân lớp để mọi điều kiện thấy "
            f"cùng dữ liệu.",
            f"Không WAP nào xuất hiện trong hơn {pc(d.max_sup_share, 1)} fingerprint "
            f"({d.max_sup_wap} là WAP xuất hiện nhiều nhất).",
            "TU(T) = Σ q(i,T)·p(i): cùng một giao dịch có utility rất khác nhau giữa ba "
            "mô hình — chỉ vì p(i)."], size=18.5, bullet="•", gap=12)

    # ------------------------------------------------ 11 step 3: p(i) table
    def s_step3(self):
        d, R = self.d, self.R
        dB = d.p["B_discriminative"].diagnostics
        top = dB.get("top10_waps_by_mi", [{}])[0]
        s = self.slide("plain", "Bước 3: external utility")
        tb(s, 46, 30, 1360, 90, "Bước 3: external utility p(i) — 3 mô hình", size=56,
           color=GOLD, shadow=True)
        rows = [["Mô hình", "p(i)", "Tổng utility CSDL", "Vai trò"],
                ["A – Signal", "p(i) = 1 với mọi WAP", vi(d.tot["A_signal"]),
                 "Đối chứng"],
                ["B – Discriminative", "max(I(A_i;Y), 10⁻⁶) / max_j I(A_j;Y)",
                 vi(d.tot["B_discriminative"]), "Mô hình chính"],
                ["C – Rare-discriminative", "∝ I(A_i;Y) × log(N / (1 + support_i))",
                 vi(d.tot["C_rare_discriminative"]), "Phạt WAP có mặt khắp nơi"]]
        table(s, 46, 128, [300, 470, 230, 348], rows, size=18.5, row_h=46)
        card(s, 46, 340, 820, 270, radius=0.05)
        tb(s, 68, 354, 780, 30, "Entropy và thông tin tương hỗ", size=20, color=GOLD,
           bold=True)
        self.F.put(s, r"H(Y)=-\sum_{y}P(y)\log_2 P(y)", 68, 394, size=24, max_w=780,
                   align="c", box_w=780)
        self.F.put(s, r"I(A_i;Y)=\sum_{a}\sum_{y}P(a,y)\,\log_2\dfrac{P(a,y)}{P(a)\,P(y)}",
                   68, 446, size=24, max_w=780, align="c", box_w=780)
        self.source(s, f"{R.who('shannon1948')}; ước lượng từ histogram đồng thời "
                       f"(6 trạng thái tín hiệu × {d.tr['joint_building_floor_n_classes']} "
                       f"lớp) — bài tìm hiểu", 68, 556, 780, size=14)
        edges = ", ".join(vi(e).replace("-", "−") for e in dB["rssi_bin_edges_dbm"])
        tb(s, 892, 344, 502, 420, [
            f"Y = BUILDINGID_FLOOR ({d.tr['joint_building_floor_n_classes']} lớp), "
            f"H(Y) = {vn(dB['label_entropy_H_Y_bits'], 3)} bit.",
            f"6 trạng thái: “không phát hiện” + 5 khoảng dBm, biên [{edges}].",
            f"WAP thông tin nhất: {top.get('wap', '')} với {vn(dB['mi_bits_max'], 3)} bit "
            f"= {pc(dB['mi_bits_max_as_fraction_of_H_Y'], 1)} của H(Y).",
            f"Sàn 10⁻⁶ giữ p > 0 vì EFIM cần utility dương; p trung bình (B) = "
            f"{vn(dB['p_mean'], 3)}."], size=18, bullet="•", gap=10)
        tb(s, 46, 632, 1350, 130, [
            f"{{g:Hệ quả 1:}} tổng utility chênh {vn(d.tot.max() / d.tot.min())} lần dù cùng "
            f"dữ liệu và cùng q ⇒ minUtil phải biểu diễn theo % tổng utility.",
            f"{{g:Hệ quả 2 (hạn chế lớn nhất):}} trên {d.O6['n_detected_waps_compared']} WAP, "
            f"hạng theo MI và hạng theo số lần phát hiện tương quan Spearman ρ = "
            f"{vn(d.O6['spearman_support_vs_mi_over_all_detected_waps'], 4)} "
            f"{R.cite('spearman1904')} ⇒ Model B gần như là tần suất phát hiện được "
            f"tham số hoá lại."], size=18.5, gap=8, spacing=1.08)

    # --------------------------------------- 12 fingerprint -> HUIM problem
    def s_mapping(self):
        d, R = self.d, self.R
        s = self.slide("plain", "Từ fingerprint đến HUIM")
        title(s, "Từ fingerprint Wi-Fi đến {g:bài toán HUIM}", size=56)
        tb(s, 48, 112, 1300, 30, "Fingerprint không có sẵn q và p, nên phải tự định "
           "nghĩa — đây là lựa chọn thiết kế của bài tìm hiểu", size=20, color=MUTED)
        rows = [["Khái niệm HUIM", "Item", "Transaction", "q(i,T)", "p(i)", "HUI"],
                ["Trong IPS", "Một WAP", "Một lần đo\n(fingerprint)",
                 "RSSI chuẩn hoá\n1…100", "{g:Model A, B hoặc C}", "Nhóm WAP\nutility cao"]]
        table(s, 46, 160, [230, 180, 230, 240, 250, 220], rows, size=19, row_h=62)
        tb(s, 48, 316, 1340, 40, "{g:Chỉ thay p(i), giữ nguyên mọi thứ khác:} cùng dữ "
           "liệu, cùng q, cùng EFIM, cùng bộ định vị.", size=21)
        card(s, 46, 380, 660, 380, radius=0.04)
        tb(s, 70, 398, 620, 30, "TỪ HUI ĐẾN TẬP WAP — 2 QUY TẮC", size=20, color=GOLD,
           bold=True)
        tb(s, 70, 440, 615, 300, [
            "{b:Union:} lấy mọi WAP xuất hiện trong ít nhất một HUI — minUtil là tham số "
            "duy nhất.",
            "{b:Top-n:} lấy n WAP có điểm đóng góp cao nhất (công thức bên phải).",
            "Cả hai quy tắc cố định TRƯỚC khi có bất kỳ kết quả độ chính xác nào.",
            "Lần khai phá bị dừng giữa chừng có tập HUI lệch theo thứ tự DFS → không "
            "dùng để xếp hạng WAP."], size=18.5, bullet="–", gap=10)
        self.fcard(s, 734, 380, 660, 380, "Điểm đóng góp của WAP i",
                   [r"\mathrm{contrib}(i)=\sum_{X\in\mathrm{HUIs},\;i\in X}"
                    r"\dfrac{u(X)}{|X|}"],
                   "Mỗi HUI chia đều utility cho các WAP thành viên; WAP có tổng "
                   "đóng góp cao xếp trước.",
                   "quy tắc xếp hạng do bài tìm hiểu đặt ra (không lấy từ tài liệu)",
                   fsize=30)

    # ----------------------------------------------------------- 13 EFIM
    def s_efim(self):
        R = self.R
        s = self.slide("efim", "Thuật toán EFIM")
        tb(s, 42, 40, 900, 110, "Thuật toán {g:EFIM}", size=82, bold=True, shadow=True)
        tb(s, 48, 150, 1340, 34, f"Zida, Fournier-Viger, Lin, Wu & Tseng — Knowledge "
           f"and Information Systems, 2017 {R.cite('zida2017')}", size=22, color=SOFT)
        cards = [(51, "01", "Tiền xử lý", "Tính TU, TWU; bỏ mục TWU < minUtil; sắp mục "
                                          "theo TWU {g:tăng} dần"),
                 (326, "02", "Chiếu CSDL", "Xuống nút α ∪ {z}: mỗi giao dịch chỉ giữ "
                                           "phần đứng sau z"),
                 (617, "03", "Duyệt sâu (DFS)", "Cây liệt kê tập mục; chỉ mở nhánh có "
                                                "su ≥ minUtil"),
                 (908, "04", "Tính utility", "Tính u(X) chính xác cho từng tập được "
                                             "duyệt"),
                 (1165, "05", "Xuất kết quả", "Giữ các tập có u(X) {g:≥} minUtil")]
        for x, n, h, t in cards:
            w = 230 if x < 900 else 205
            tb(s, x + 24, 232, 80, 36, f"{{g:{n}}}", size=28)
            tb(s, x + 24, 272, w, 34, h, size=22, bold=True)
            tb(s, x + 24, 310, w, 95, t, size=16.5, spacing=1.06)
        self.F.put(s, r"u(X)=\sum_{T\supseteq X}u(X,T)", 912, 440, size=22,
                   color="#F2BE3A", max_w=212, align="c", box_w=212)
        tb(s, 1182, 449, 173, 46, "{g:HUIs}", size=34, align="c")
        tb(s, 112, 630, 230, 120, ["{g:TỐI ƯU}", "{g:EFIM}"], size=42, spacing=0.95)
        tb(s, 380, 622, 440, 30, "{g:HAI CẬN TRÊN lu VÀ su}", size=21)
        tb(s, 380, 660, 440, 100, "local utility và subtree utility loại sớm mục / "
           "nhánh không thể đạt minUtil (slide sau)", size=18.5, spacing=1.08)
        tb(s, 883, 622, 490, 30, "{g:CHIẾU & GỘP GIAO DỊCH}", size=21)
        tb(s, 883, 660, 480, 100, "giao dịch chiếu giống hệt nhau được gộp; Fast "
           "Utility Counting tính lu, su trong một lượt quét", size=18.5, spacing=1.08)

    # -------------------------------------------------- 14 EFIM bounds
    def s_bounds(self):
        d, R = self.d, self.R
        s = self.slide("plain", "EFIM: hai cận trên")
        title(s, "EFIM: {g:hai cận trên} lu và su", size=56)
        card(s, 46, 128, 760, 560, radius=0.035)
        tb(s, 70, 144, 720, 30, "Công thức", size=20, color=GOLD, bold=True)
        y = 188
        for tex in (r"\mathrm{re}(\alpha,T)=\sum_{i\in T,\;i\succ\alpha}u(i,T)",
                    r"\mathrm{lu}(\alpha,z)=\sum_{T\supseteq\alpha\cup\{z\}}"
                    r"\left[u(\alpha,T)+\mathrm{re}(\alpha,T)\right]",
                    r"\mathrm{su}(\alpha,z)=\sum_{T\supseteq\alpha\cup\{z\}}\left[u(\alpha,T)"
                    r"+u(z,T)+\sum_{i\in T,\;i\succ z}u(i,T)\right]"):
            _, fh = self.F.put(s, tex, 70, y, size=24, max_w=712)
            y += fh + 16
        tb(s, 70, y + 4, 712, 250, [
            "Chỉ giữ z có lu(α, z) ≥ minUtil (tập Secondary); z khác bị loại khỏi CẢ "
            "cây con của α.",
            "Chỉ mở nhánh α ∪ {z} khi su(α, z) ≥ minUtil (tập Primary).",
            "Tại gốc: lu(∅, z) = TWU(z) ⇒ cắt tỉa TWU là trường hợp riêng.",
            "re(α, T): utility còn lại của các mục đứng sau α trong T."],
           size=18.5, bullet="–", gap=8)
        self.source(s, f"{R.who('zida2017')} — cả hai cận trên và Fast Utility Counting",
                    70, 646, 720, size=14)
        tw, order, su0 = d.ex_twu, d.ex_order, d.ex_su0
        tb(s, 836, 128, 560, 60, f"Ví dụ chạy chuẩn trong {R.cite('zida2017')} — "
           f"{len(d.ex_db)} giao dịch, minUtil = {d.ex_mu}", size=18.5, color=GOLD,
           bold=True, spacing=1.05)
        rows = [["Mục", "TWU", "Thứ tự", "su gốc", "Vai trò"]]
        for i in order:
            rows.append([i, vi(tw[i]), str(order.index(i) + 1), vi(su0[i]),
                         "Primary" if su0[i] >= d.ex_mu else "Chỉ Secondary"])
        table(s, 836, 196, [80, 100, 100, 110, 170], rows, size=17, row_h=38)
        st = d.ex_st
        n_sub = 2 ** len(d.ex_p) - 1
        tb(s, 836, 196 + 38 * len(rows) + 14, 560, 200, [
            f"EFIM trả về {len(d.ex_huis)} HUI nhưng chỉ duyệt {st.n_nodes_visited} / "
            f"{n_sub} tập mục khả dĩ; {st.n_candidates_pruned_by_su} nhánh bị loại nhờ su.",
            f"Lý do chọn EFIM: đánh giá thực nghiệm độc lập {R.cite('zhang2018')}; SPMF "
            f"{R.cite('spmf2016')} cần Java nên EFIM được cài lại bằng Python và đối chiếu "
            f"oracle vét cạn."], size=16.5, bullet="•", gap=8, spacing=1.06)

    # --------------------------------------------- 15 inside EFIM / TWU table
    def s_inside(self):
        d = self.d
        s = self.slide("plain", "Tiền xử lý bên trong EFIM")
        title(s, "Tiền xử lý {g:bên trong EFIM}", size=56)
        tb(s, 48, 112, 1300, 30, "Chạy ở đầu mỗi lần khai phá, trước khi tìm kiếm theo "
           "chiều sâu", size=20, color=MUTED)
        tb(s, 46, 170, 700, 560, [
            "Kiểm tra mọi p(i) > 0 và q > 0; thiếu thì báo lỗi.",
            "Tính TU(T) = Σ q·p cho mỗi giao dịch; TWU(i) = tổng TU của các giao dịch chứa i.",
            "Cắt tỉa TWU: bỏ item có TWU < minUtil.",
            "Sắp item theo TWU tăng dần (bằng nhau thì theo nhãn).",
            "Dựng CSDL đã cắt tỉa: bỏ giao dịch rỗng, gộp giao dịch giống hệt nhau "
            "(có cộng dồn số lần xuất hiện để độ hỗ trợ vẫn đúng).",
            "Tính local utility / subtree utility cho gốc bằng Fast Utility Counting."],
           size=19.5, bullet="#", gap=10, spacing=1.08)
        grid = [20.0, 10.0, 5.0, 2.5, 2.0]
        rows = [["minUtil", "A: item còn lại", "B: item còn lại", "C: item còn lại"]]
        for g in grid:
            rows.append([f"{mu(g)}%"] + [str(int(d.run(m, g)["items_after_twu_pruning"]))
                                          for m in MODELS])
        table(s, 790, 176, [140, 160, 160, 160], rows, size=19, row_h=50)
        init = int(d.H["items_initial"].iloc[0])
        tb(s, 790, 176 + 50 * len(rows) + 16, 620, 160,
           f"minUtil tính theo % tổng utility của CSDL. Ngưỡng càng cao thì cắt tỉa TWU "
           f"càng loại nhiều item: trong {init} item ban đầu, ở minUtil {mu(grid[0])}% "
           f"chỉ còn {int(d.run('A_signal', grid[0])['items_after_twu_pruning'])} (A), "
           f"{int(d.run('B_discriminative', grid[0])['items_after_twu_pruning'])} (B) và "
           f"{int(d.run('C_rare_discriminative', grid[0])['items_after_twu_pruning'])} (C).",
           size=16.5, color=MUTED, spacing=1.08)

    # -------------------------------------------------- 16 experiment design
    def s_exp(self):
        d, cs = self.d, self.d.cs
        s = self.slide("exp", "Thực nghiệm trên UJIIndoorLoc")
        tb(s, 31, 46, 1360, 90, "Thực nghiệm trên bộ dữ liệu UJIIndoorLoc", size=64,
           bold=True, shadow=True)
        kp = [(225, vi(cs["train_rows"]), "fingerprint huấn luyện"),
              (378, vi(cs["validation_rows"]), "fingerprint kiểm định"),
              (534, str(d.n_wap), "cột WAP (đặc trưng Wi-Fi)")]
        for y, v, lab in kp:
            tb(s, 70, y - 6, 250, 70, f"{{g:{v}}}", size=56)
            tb(s, 72, y + 70, 250, 30, lab, size=19)
        tb(s, 373, 188, 600, 34, "QUY TRÌNH THỰC NGHIỆM", size=26, bold=True)
        g = d.PO["selected_config"]
        steps = [(463, "01", "TIỀN XỬ LÝ\nRSSI", "100 → không phát hiện; RSSI → "
                                                  "1…100 (học trên train)"),
                 (668, "02", "CHUYỂN ĐỔI\nFINGERPRINT", f"thành "
                  f"{vi(d.hsum['database']['n_non_empty_transactions'])} utility "
                  f"transaction"),
                 (883, "03", "KHAI PHÁ\nHUIM VỚI EFIM", f"{len(d.H)} lần chạy, 3 mô hình "
                  f"p(i), minUtil {mu(d.H['rel_minutil_pct'].max())}% → "
                  f"{mu(d.H['rel_minutil_pct'].min())}%"),
                 (1089, "04", "ĐỊNH VỊ\nk-NN", f"V1: k = 1, Manhattan; V2: Bray-Curtis, "
                  f"k = {g['k']}"),
                 (1294, "05", "ĐÁNH GIÁ\nHIỆU NĂNG", "tòa nhà, tầng, sai số mét, thời "
                                                    "gian, bộ nhớ")]
        for cx, n, h, t in steps:
            tb(s, cx - 95, 270, 190, 60, f"{{g:{n}}}", size=46, align="c")
            tb(s, cx - 95, 388, 190, 64, h, size=19.5, bold=True, align="c", spacing=1.0)
            tb(s, cx - 84, 462, 168, 140, t, size=16.5, align="c", spacing=1.06)
        tb(s, 412, 676, 170, 50, "{g:KẾT QUẢ}", size=32)
        tb(s, 622, 652, 750, 100,
           f"Mọi tham số chọn trên trainingData; validationData chỉ dùng để đánh giá, "
           f"một lần cho mỗi điều kiện ({len(d.I)} điều kiện ở Version 1, 5 bước "
           f"S0–S4 ở Version 2).", size=19.5, anchor="m", spacing=1.08)

    # ------------------------------------------------ 17 minUtil vs #HUI
    def s_minutil(self):
        d = self.d
        s = self.slide("plain", "minUtil và số HUI")
        tb(s, 40, 40, 560, 230, "Trong HUIM, mỗi phần tử có một giá trị utility và "
           "thuật toán tìm những tập phần tử có tổng utility đạt ngưỡng {b:minUtil}.",
           size=27, color=GOLD, spacing=1.1)
        lines = []
        for m in MODELS:
            g = d.H[d.H["utility_model"] == m]
            zero = int((g["n_huis"] == 0).sum())
            first = g[g["n_huis"] > 0]["rel_minutil_pct"].max()
            lines.append(f"Model {MNAME[m]}: {zero} lần chạy cho 0 HUI; HUI đầu tiên "
                         f"xuất hiện ở {mu(first)}%")
        jump = []
        for m in MODELS:
            g = d.H[d.H["utility_model"] == m].sort_values("rel_minutil_pct")
            done = g[g["status"] == "COMPLETED"]
            jump.append(int(done["unique_waps_in_huis"].max()))
        lines += ["Chuyển pha rất gắt: từ “không có gì” sang “vượt ngân sách” chỉ trong "
                  "một bước ngưỡng.",
                  f"Số HUI khổng lồ nhưng chỉ đến từ {min(jump)}–{max(jump)} WAP riêng biệt "
                  f"(lần chạy xong)."]
        tb(s, 40, 268, 565, 500, lines, size=19, bullet="–", gap=9, spacing=1.08)
        p = ch_hui(d, 740, 700)
        chart_card(s, p, 632, 44, 770, 730)

    # ------------------------------------------------ 18 cost by p(i)
    def s_cost_p(self):
        d = self.d
        s = self.slide("plain", "p(i) quyết định chi phí")
        title(s, "p(i) quyết định {g:chi phí khai phá}", size=56)
        db = d.hsum["database"]
        tb(s, 48, 112, 1300, 30, f"Cùng ngưỡng minUtil = 2% tổng utility, EFIM chạy trên "
           f"{vi(db['n_non_empty_transactions'])} giao dịch", size=20, color=MUTED)
        p = ch_cost2(d, 880, 590)
        chart_card(s, p, 46, 160, 910, 610)
        rA, rB, rC = (d.run(m, 2.0) for m in MODELS)
        lo = rA["runtime_s"] / max(rB["runtime_s"], rC["runtime_s"])
        hi = rA["runtime_s"] / min(rB["runtime_s"], rC["runtime_s"])
        card(s, 986, 160, 408, 290, radius=0.06)
        tb(s, 1010, 186, 370, 70, f"{{g:{vn(lo, 1)}–{vn(hi, 1)}×}}", size=50)
        tb(s, 1010, 270, 360, 170, "Model A chậm hơn B và C ở cùng ngưỡng.", size=20,
           spacing=1.1)
        card(s, 986, 470, 408, 300, radius=0.06)
        tb(s, 1010, 494, 360, 60, f"{{g:{vi(rA['n_huis'])} vs {vi(rB['n_huis'])}}}",
           size=42)
        tb(s, 1010, 566, 362, 200, "Số HUI của A so với B: A có ít HUI hơn nhưng vẫn "
           "chậm hơn, vì p = 1 làm các cận trên ít phân hoá mục nên cắt tỉa kém.",
           size=18.5, spacing=1.1)

    # --------------------------------------------- 19 runtime / memory
    def s_runtime(self):
        d = self.d
        from scipy.stats import spearmanr
        s = self.slide("plain", "Runtime và bộ nhớ")
        tb(s, 40, 36, 600, 170, ["Runtime và bộ nhớ", "theo ngưỡng"], size=56,
           color=GOLD, spacing=1.0, shadow=True)
        lines = []
        rho = []
        for m in MODELS:
            g = d.H[d.H["utility_model"] == m]
            ab = int((g["status"] != "COMPLETED").sum())
            lines.append(f"Model {MNAME[m]}: {vn(g['runtime_s'].min())}–"
                         f"{vn(g['runtime_s'].max())} s; bộ nhớ đỉnh "
                         f"{vn(g['peak_memory_delta_mb'].min(), 1)}–"
                         f"{vn(g['peak_memory_delta_mb'].max(), 1)} MB; {ab} lần dừng "
                         f"(ABORTED)")
            rho.append(vn(spearmanr(g["search_nodes_visited"], g["runtime_s"])[0], 3))
        lines.append(f"Spearman ρ(số nút đã duyệt, runtime) = {', '.join(rho)} "
                     f"{self.R.cite('spearman1904')} → chi phí là chi phí TRÊN MỖI NÚT.")
        lines.append(f"Giới hạn mỗi lần chạy: {vi(d.t_lim)} s hoặc {vi(d.max_huis)} HUI; "
                     f"vượt thì dừng có kiểm soát và giữ kết quả một phần.")
        tb(s, 40, 214, 590, 560, lines, size=19, bullet="–", gap=10, spacing=1.08)
        p = ch_runtime_mem(d, 740, 700)
        chart_card(s, p, 652, 44, 750, 730)

    # ------------------------------------------- 20 HUIM vs controls
    def s_controls(self):
        d = self.d
        s = self.slide("plain", "HUIM vs đối chứng")
        title(s, "HUIM có hơn các {g:đối chứng} không?", size=56)
        b = d.base
        tb(s, 48, 112, 1340, 30, f"Độ chính xác joint (tòa nhà + tầng) trên validationData, "
           f"cùng số WAP; baseline 520 WAP = {pc(b['joint_building_floor_accuracy'])} "
           f"(k-NN Version 1)", size=19, color=MUTED)
        p = ch_union(d, 900, 590)
        chart_card(s, p, 46, 160, 930, 610)
        O5 = d.O5
        k = [(f"{O5['huim_beats_random']}/{O5['n_comparisons']}", "ngân sách WAP: HUIM "
              "tốt nhất vượt chọn ngẫu nhiên"),
             (f"{O5['huim_beats_top_mi']}/{O5['n_comparisons']}", "vượt top-n theo "
              "mutual information"),
             (f"{O5['huim_beats_top_support']}/{O5['n_comparisons']}", "vượt top-n theo "
              "số lần phát hiện")]
        for i, (v, t) in enumerate(k):
            y = 160 + i * 122
            card(s, 1004, y, 390, 108, radius=0.08)
            tb(s, 1024, y + 14, 120, 70, f"{{g:{v}}}", size=40, anchor="m")
            tb(s, 1150, y + 14, 230, 80, t, size=16.5, anchor="m", spacing=1.05)
        tb(s, 1004, 538, 390, 230,
           "{g:Kết luận trung thực:} HUIM union thắng ngẫu nhiên ở B và C nhưng thua "
           "top-n theo MI ở cả ba mô hình. Mọi cách chọn đều thấp hơn dùng đủ 520 WAP.",
           size=17.5, spacing=1.1)

    # ------------------------------------------- 21 feature reduction
    def s_featred(self):
        d = self.d
        s = self.slide("res_fr", "Feature Reduction và Accuracy")
        tb(s, 31, 26, 1300, 90, "Feature Reduction và Accuracy", size=62, bold=True,
           shadow=True)
        p = ch_featred(d, 790, 620)
        chart_card(s, p, 30, 125, 800, 640)
        O9 = d.O9
        best = O9["best_joint_at_or_below_that_budget_any_method"]
        sm = O9["smallest_control_subset_reaching_baseline"]
        tb(s, 860, 140, 540, 620, [
            "KHÔNG tập con nào do HUIM chọn đạt joint accuracy của baseline toàn bộ WAP.",
            "Building accuracy chịu được giảm chiều mạnh; Floor accuracy là ràng buộc chặt.",
            f"Tập HUIM lớn nhất có {O9['largest_huim_selected_set']} WAP; ở ngân sách đó "
            f"tốt nhất mọi phương pháp chỉ {pc(best['joint'])} ({best['condition']}).",
            f"Cần {sm['n_waps']} WAP ({sm['condition']}, {pc(sm['joint'])}) mới chạm "
            f"baseline — vượt xa trần mà HUIM cung cấp.",
            "Hình thoi = điểm union, nơi minUtil là núm điều khiển duy nhất."],
           size=20, bullet="–", gap=12, spacing=1.08)

    # ------------------------------------------- 22 time & memory of IPS
    def s_ipscost(self):
        d = self.d
        s = self.slide("res_cost", "Thời gian và bộ nhớ")
        tb(s, 31, 26, 900, 90, "Kết quả và Phân tích", size=62, bold=True, shadow=True)
        tb(s, 900, 36, 500, 40, "Đánh giá thời gian và bộ nhớ", size=27, bold=True,
           align="r")
        p = ch_ips_cost(d, 1290, 400)
        chart_card(s, p, 46, 112, 1348, 400)
        cb = d.cost.loc["baseline_all_520"]
        sm = d.O9["smallest_control_subset_reaching_baseline"]
        c2 = d.cost.loc[sm["condition"]]
        hc = d.D.COST[d.D.COST["family"].str.startswith("huim")]
        hb = hc.loc[hc["joint_building_floor_accuracy"].idxmax()]
        tb(s, 84, 600, 200, 80, "{g:Nhận xét}", size=29)
        tb(s, 330, 534, 930, 240, [
            f"Baseline 520 WAP: {vn(cb['predict_ms_per_fingerprint'], 3)} ms/fingerprint, "
            f"CSDL vân tay {vn(cb['db_memory_float64_mb'], 1)} MB (float64).",
            f"{sm['condition']} ({int(c2['n_features'])} WAP): "
            f"{vn(cb['predict_ms_per_fingerprint'] / c2['predict_ms_per_fingerprint'], 1)}× "
            f"nhanh hơn, {vn(c2['db_memory_float64_mb'] / cb['db_memory_float64_mb'] * 100, 0)}% "
            f"bộ nhớ, joint {pc(c2['joint_building_floor_accuracy'])}.",
            f"HUIM tốt nhất ({int(hb['n_features'])} WAP): "
            f"{vn(cb['predict_ms_per_fingerprint'] / hb['predict_ms_per_fingerprint'], 1)}× "
            f"nhanh hơn, {vn(hb['db_memory_float64_mb'] / cb['db_memory_float64_mb'] * 100, 0)}% "
            f"bộ nhớ, nhưng joint chỉ {pc(hb['joint_building_floor_accuracy'])}.",
            "Đo một lõi, lặp lại lấy trung vị; bộ nhớ CSDL = dòng × cột × kích thước phần "
            "tử. Bộ định vị Version 2 chưa được đo thời gian riêng."],
           size=16.5, bullet="–", gap=4, spacing=1.03)

    # -------------------------------------- 23 metres, Version 1
    def s_metres(self):
        d, R = self.d, self.R
        b, cc = d.base, d.coord
        s = self.slide("res_m", "Sai số mét Version 1")
        tb(s, 31, 26, 900, 90, "Kết quả và Phân tích", size=62, bold=True, shadow=True)
        tb(s, 860, 36, 540, 40, "Độ chính xác định vị theo mét — V1", size=26,
           bold=True, align="r")
        tb(s, 60, 214, 110, 80, ["{g:Nhận}", "{g:xét}"], size=29, spacing=1.0)
        tb(s, 191, 182, 560, 520, [
            f"Baseline V1: trung vị {vn(b['pos_median_error'])} m toạ độ chiếu = "
            f"{vn(b['pos_median_error_ground_m'])} m thực địa; trung bình "
            f"{vn(b['pos_mean_error'])} m chiếu = {vn(b['pos_mean_error_ground_m'])} m "
            f"thực địa.",
            f"Trung vị DƯỚI dải 7–8 m; trung bình TRÊN dải vì phân bố có đuôi dài "
            f"(P95 = {vn(b['pos_p95_error_ground_m'], 1)} m).",
            f"Toạ độ là EPSG:3857 Web Mercator — xác định bằng nghịch đảo phép chiếu, "
            f"rơi cách campus UJI {vn(cc['distance_to_uji_campus_m'], 1)} m.",
            f"Web Mercator KHÔNG bảo toàn khoảng cách: hệ số "
            f"{vn(cc['web_mercator_scale_factor_at_this_latitude'], 6)} → đọc thẳng từ "
            f"cột toạ độ lớn hơn thực tế khoảng "
            f"{vn((cc['web_mercator_scale_factor_at_this_latitude'] - 1) * 100, 0)}%."],
           size=18.5, bullet="–", gap=12, spacing=1.08)
        p = ch_v1_error(d, 600, 400)
        chart_card(s, p, 800, 104, 610, 420)
        card(s, 800, 540, 610, 225, radius=0.05)
        tb(s, 822, 552, 560, 30, "Sai số thực địa của một fingerprint", size=18.5,
           color=GOLD, bold=True)
        self.F.put(s, r"e=\sqrt{(x-\hat{x})^2+(y-\hat{y})^2}\;\times\;\cos\varphi",
                   822, 592, size=24, max_w=566, align="c", box_w=566)
        tb(s, 822, 652, 566, 50, f"φ = {vn(cc['inverted_centroid_lat_deg'], 4)}° ⇒ cos φ = "
           f"{vn(cc['ground_distance_correction_factor'], 6)}", size=16, color=SOFT)
        self.source(s, f"hệ số tỉ lệ Mercator sec φ — {R.who('snyder1987')}; báo cáo "
                       f"trung vị, trung bình, phân vị theo {R.who('adler2015')}",
                    822, 690, 566, size=13.5)

    # --------------------------------------------- 24 section divider
    def s_divider(self):
        d = self.d
        s0, sel = d.s0, d.sel
        s = self.slide("cover", "Version 2")
        tb(s, 46, 70, 680, 120, "{g:Version 2}", size=96, bold=True, shadow=True)
        tb(s, 46, 190, 650, 200, ["Tối ưu sai số", "định vị theo mét"], size=62,
           bold=True, spacing=0.98, shadow=True)
        tb(s, 50, 400, 600, 120, "Giữ nguyên dữ liệu, HUIM và giao thức chống rò rỉ — "
           "chỉ thay bộ ước lượng toạ độ.", size=24, color=SOFT, spacing=1.1)
        drop = (1 - sel["val_mean_ground_m"] / s0["val_mean_ground_m"]) * 100
        tb(s, 50, 605, 640, 50, f"{{g:{vn(s0['val_mean_ground_m'])} m → "
           f"{vn(sel['val_mean_ground_m'])} m  (−{vn(drop, 0)}%)}}", size=34)
        tb(s, 52, 662, 620, 60, f"sai số trung bình (m thực địa) trên "
           f"{vi(len(d.D.Pnew))} fingerprint kiểm định", size=19, color=SOFT)

    # ------------------------------------------- 25 why V1 error was large
    def s_why(self):
        d, R, cs = self.d, self.R, self.d.cs
        s = self.slide("plain", "Vì sao V1 sai số lớn?")
        title(s, "Vì sao {g:Version 1} sai số lớn?", size=56)
        n_ph = len(cs["PHONEID_validation_values"])
        n_pos = len(d.D.va.drop_duplicates(["LONGITUDE", "LATITUDE", "FLOOR"]))
        items = [("Chọn tham số sai mục tiêu",
                  "k = 1, Manhattan được chọn cho bài toán PHÂN LOẠI tầng, không phải "
                  "ước lượng toạ độ."),
                 ("Phép chia dev không giống tập kiểm định",
                  "Chia ngẫu nhiên 75/25 đặt fingerprint đo tại cùng điểm vào cả hai "
                  f"phía nên k = 1 thắng; validationData dùng thiết bị khác "
                  f"({n_ph - len(cs['PHONEID_overlap'])}/{n_ph} mã thiết bị không có trong "
                  f"train), tại {vi(n_pos)} vị trí."),
                 ("Phương sai cao của k = 1",
                  "Toạ độ dự đoán chính là một vị trí huấn luyện: láng giềng sai là sai "
                  "số trọn vẹn."),
                 ("Hoà khoảng cách",
                  "RSSI là số nguyên nên khoảng cách Manhattan thường bằng nhau; với k = 1 "
                  "cách phá hoà quyết định toạ độ."),
                 ("Lệch thiết bị",
                  f"Điện thoại khác nhau đo RSSI lệch mức, làm sai khoảng cách tính trên "
                  f"RSSI thô {R.cite('khalaj2017', 'kaem2004')}.")]
        for i, (h, t) in enumerate(items):
            y = 128 + i * 128
            card(s, 46, y, 1348, 112, radius=0.08)
            c = card(s, 66, y + 26, 60, 60, fill=GOLD, line=None, radius=0.5)
            tb(s, 66, y + 26, 60, 60, str(i + 1), size=28, bold=True, color=INK,
               align="c", anchor="m")
            tb(s, 150, y + 14, 1220, 32, h, size=21, bold=True, color=GOLD)
            tb(s, 150, y + 48, 1220, 60, t, size=18, spacing=1.05)

    # ------------------------------------------- 26 selection protocol
    def s_protocol(self):
        d = self.d
        s = self.slide("plain", "Giao thức chọn mô hình")
        title(s, "Giao thức chọn mô hình {g:không rò rỉ}", size=56)
        n_users = int(d.tr["USERID_n_unique"])
        k = d.n_folds
        tb(s, 48, 112, 1300, 30, f"GroupKFold({k}) theo USERID trên trainingData: "
           f"{n_users} người đo chia {k} nhóm; mỗi fold dev chỉ gồm người đo mà mô hình "
           f"chưa thấy", size=19, color=MUTED)
        card(s, 46, 160, 600, 480, radius=0.04)
        tb(s, 66, 172, 560, 30, f"{k} fold giữ lại người đo", size=19, color=GOLD,
           bold=True)
        cw, chh = int(396 / k) - 8, 50
        for r in range(k):
            tb(s, 66, 222 + r * 66, 90, chh, f"Fold {r + 1}", size=16, anchor="m")
            for c in range(k):
                dev = r == c
                card(s, 160 + c * (cw + 8), 222 + r * 66, cw, chh,
                     fill=GOLD if dev else HEAD_FILL, line=None, radius=0.15)
        tb(s, 66, 598, 560, 30, "Ô {g:vàng} = fold dev (chấm điểm để CHỌN) · ô xanh = "
           "huấn luyện", size=15, color=SOFT)
        rows = [["Bước", "Dữ liệu dùng"],
                ["Chuẩn hoá RSSI, p(i), khai phá HUIM, chọn WAP", "trainingData"],
                ["Chọn siêu tham số k-NN (Version 1)", "Chia ngẫu nhiên 75/25 của train"],
                ["Chọn mô hình định vị (Version 2)", f"{k} fold giữ lại người đo"],
                ["Huấn luyện bộ định vị cuối", "Toàn bộ trainingData"],
                ["Đánh giá", "validationData, 1 lần / bước"]]
        table(s, 672, 160, [410, 312], rows, size=16.5, row_h=52,
              highlight={3: RGBColor(0x2A, 0x4A, 0x2A)})
        P = d.PO["protocol"]
        tb(s, 672, 494, 722, 280, [
            "Tiêu chí chọn: sai số trung bình (m thực địa) trên các fold dev.",
            f"5 bước S0 → S4 nối tiếp; bước sau chỉ được nhận nếu giảm sai số dev ít "
            f"nhất {vn(P['min_gain_m'])} m.",
            "validationData chỉ chấm điểm mỗi bước một lần, SAU khi lựa chọn đã xong."],
           size=18, bullet="–", gap=8, spacing=1.08)
        tb(s, 46, 660, 600, 110, "Leave-users-out mô phỏng đúng tình huống kiểm định: "
           "người đo / thiết bị mới, chưa từng có trong dữ liệu huấn luyện.", size=17,
           color=SOFT, spacing=1.08)

    # ------------------------------------------- 27 V2 formulas
    def s_v2_formulas(self):
        d, R = self.d, self.R
        g = d.PO["protocol"]["grids"]
        cfg = d.PO["selected_config"]
        s = self.slide("plain", "Công thức Version 2")
        title(s, "Các cải tiến của Version 2 — {g:công thức và nguồn}", size=50)
        W2, H2 = 664, 300
        self.fcard(s, 46, 112, W2, H2, "1 · Biểu diễn tín hiệu (RSSI phát hiện được)",
                   [r"\mathrm{powed}=\left(\dfrac{\mathrm{RSSI}+105}{105}\right)^{e}",
                    r"\mathrm{exponential}=\dfrac{\exp\left((\mathrm{RSSI}+105)/24\right)}"
                    r"{\exp(105/24)}"],
                   f"Lưới: {', '.join(g['representations'])}; chọn {cfg['representation']}",
                   R.who("torres2015"), fsize=23)
        self.fcard(s, 730, 112, W2, H2, "2 · Khoảng cách Bray–Curtis",
                   [r"d_{\mathrm{BC}}(a,b)=\dfrac{\sum_{j}|a_j-b_j|}{\sum_{j}(a_j+b_j)}"],
                   f"Lưới: {', '.join(g['distances'])}; chuẩn hoá theo tổng cường độ "
                   f"của hai vector",
                   f"{R.who('bray1957')}; dùng cho vân tay Wi-Fi — {R.who('torres2015')}",
                   fsize=26)
        self.fcard(s, 46, 428, W2, H2, "3 · k-NN có trọng số nghịch đảo khoảng cách",
                   [r"\hat{x}=\dfrac{\sum_{i=1}^{k}w_i\,x_i}{\sum_{i=1}^{k}w_i},\qquad "
                    r"w_i=\dfrac{1}{d_i^{\,p}}"],
                   f"k từ {min(g['k'])} đến {max(g['k'])}, p từ {min(g['powers'])} đến "
                   f"{max(g['powers'])}; chọn k = {cfg['k']}, p = {int(cfg['power'])}",
                   f"k-NN trong không gian tín hiệu — {R.who('bahl2000')}; trọng số "
                   f"nghịch đảo khoảng cách mũ p — {R.who('shepard1968')}", fsize=25)
        self.fcard(s, 730, 428, W2, H2, "4 · Dò lệch thiết bị và ensemble",
                   [r"\delta^{*}=\arg\min_{\delta}\;\min_{j}\;d\left(r+\delta,\;f_j\right)",
                    r"\hat{x}=\frac{1}{2}\left(\hat{x}_{\mathrm{exp}}+\hat{x}_{\mathrm{pos}}\right)"],
                   f"δ ∈ {{−{cfg['offset_range']}, …, +{cfg['offset_range']}}} dB, bước "
                   f"{g['offset_step_db']} dB; chỉ dùng tín hiệu của chính truy vấn",
                   "thiết kế của bài tìm hiểu — chưa đối chiếu tài liệu về hiệu chỉnh "
                   "lệch thiết bị [CẦN BỔ SUNG NGUỒN]", fsize=23)
        tb(s, 46, 744, 1350, 40, "Phá hoà tất định: láng giềng sắp theo (khoảng cách, "
           "chỉ số dòng) — kết quả không phụ thuộc cách thư viện xử lý hoà.", size=16.5,
           color=SOFT)

    # ------------------------------------------- 28 stage results
    def s_stages(self):
        d = self.d
        A, sel, bv = d.AB, d.sel, d.best_val
        s = self.slide("plain", "Kết quả S0–S4")
        title(s, "Kết quả từng bước {g:S0 → S4}", size=56)
        p = ch_stages(d, 820, 590)
        chart_card(s, p, 46, 130, 840, 620)
        rows = [["Bước", "Dev TB", "KĐ TB", "KĐ trung vị", "Joint"]]
        for _, r in A.iterrows():
            rows.append([r["stage"], vn(r["dev_mean_ground_m"]), vn(r["val_mean_ground_m"]),
                         vn(r["val_median_ground_m"]),
                         pc(r["val_joint_building_floor_accuracy"])])
        table(s, 906, 130, [80, 100, 100, 120, 110], rows, size=16.5, row_h=44,
              highlight={len(rows) - 1: RGBColor(0x5A, 0x48, 0x14)})
        tb(s, 906, 400, 500, 30, "m thực địa · KĐ = validationData", size=14.5,
           color=MUTED, italic=True)
        tb(s, 906, 438, 490, 330, [
            "Phần lớn mức giảm đến từ S1–S2 (khoảng cách + k láng giềng có trọng số).",
            f"{{g:Trung thực:}} {bv['stage']} có sai số KĐ thấp nhất "
            f"({vn(bv['val_mean_ground_m'])} m), thấp hơn {sel['stage']} "
            f"({vn(sel['val_mean_ground_m'])} m). Giao thức chọn {sel['stage']} trên dev "
            f"nên giữ {sel['stage']}; chọn lại theo số KĐ là dùng tập kiểm tra để chọn mô "
            f"hình.",
            "Dò lệch và ensemble chủ yếu nâng độ chính xác tầng."],
           size=17, bullet="–", gap=8, spacing=1.06)

    # ------------------------------------------- 29 CDF
    def s_cdf(self):
        d, R = self.d, self.R
        BS, CDF = d.BS, d.CDF
        v1, v2 = CDF.iloc[0], CDF.iloc[1]
        s = self.slide("plain", "Phân bố sai số")
        title(s, "Phân bố sai số: {g:Version 1} so với {g:Version 2}", size=50)
        p = ch_cdf(d, 800, 560)
        chart_card(s, p, 46, 120, 820, 580)
        rows = [["m thực địa", "V1", "V2"]]
        for lab, k, f in (("Trung bình", "mean", vn), ("P25", "p25", vn),
                          ("Trung vị", "p50", vn), ("P75", "p75", vn),
                          ("P90", "p90", vn), ("P95", "p95", vn),
                          ("Lớn nhất", "max", vn),
                          ("≤ 5 m", "le5", lambda v: pc(v, 1)),
                          ("≤ 8 m", "le8", lambda v: pc(v, 1))):
            rows.append([lab, f(v1[k]), f(v2[k])])
        table(s, 890, 120, [190, 150, 150], rows, size=16.5, row_h=38)
        card(s, 890, 520, 504, 250, radius=0.05)
        tb(s, 910, 532, 470, 30, f"Bootstrap {vi(BS['B'])} lần (khoảng tin cậy 95%)",
           size=17, color=GOLD, bold=True)
        tb(s, 910, 568, 470, 170, [
            f"TB V2: {vn(BS['new'])} m ({vn(BS['new_ci'][0])}–{vn(BS['new_ci'][1])})",
            f"TB V1: {vn(BS['old'])} m ({vn(BS['old_ci'][0])}–{vn(BS['old_ci'][1])})",
            f"Giảm ghép cặp: {vn(BS['diff'])} m ({vn(BS['diff_ci'][0])}–"
            f"{vn(BS['diff_ci'][1])})",
            f"V2 sai ít hơn trên {pc(BS['better'], 1)} fingerprint"],
           size=16.5, bullet="–", gap=4, spacing=1.05)
        self.source(s, R.who("efron1979"), 910, 734, 470, size=13.5)
        tb(s, 46, 712, 820, 60, "Version 2 dịch cả phân bố sai số sang trái — không chỉ "
           "trung bình mà cả đuôi (P90, P95, sai số lớn nhất).", size=17, color=SOFT)

    # ------------------------------------------- 30 V2 on HUIM WAP sets
    def s_sets(self):
        d = self.d
        CN = d.CN.set_index("condition")
        s = self.slide("plain", "V2 trên tập WAP HUIM")
        title(s, "Version 2 trên các {g:tập WAP của HUIM}", size=54)
        p = ch_sets(d, 900, 600)
        chart_card(s, p, 46, 125, 920, 640)
        hu = CN.loc["huim_union__C_rare_discriminative__1.7321pct"]
        tm = CN.loc["control_top_mi__n91"]
        tb(s, 990, 135, 410, 620, [
            "Version 2 giảm sai số với MỌI tập WAP.",
            f"Nhưng tập HUIM vẫn trên {vn(CN.loc[[c for c in CN.index if c.startswith('huim_union')], 'opt_mean_ground_m'].min(), 0)} m"
            f" — xa dải 7–8 m.",
            f"Cùng {int(hu['n_waps'])} WAP: HUIM union C {vn(hu['study_mean_ground_m'])} → "
            f"{vn(hu['opt_mean_ground_m'])} m, top-n theo MI {vn(tm['study_mean_ground_m'])}"
            f" → {{g:{vn(tm['opt_mean_ground_m'])} m}}.",
            "{g:Kết luận RQ2 của Version 1 không đổi:} đối chứng đơn giản vẫn chọn WAP tốt "
            "hơn HUIM. Cải thiện của V2 đến từ bộ ước lượng toạ độ, không từ HUIM."],
           size=18.5, bullet="–", gap=12, spacing=1.08)

    # ------------------------------------------- 31 V1 vs V2 numbers
    def s_cmp_numbers(self):
        d = self.d
        s0, sel = d.s0, d.sel
        s = self.slide("plain", "So sánh V1–V2: số liệu")
        title(s, "So sánh {g:Version 1} và {g:Version 2}: số liệu", size=52)

        def delta(a, b, unit="m", pct=False):
            x = (b - a) * (100 if pct else 1)
            sign = "−" if x < 0 else "+"
            col = "G" if (x < 0) == (unit == "m") else "R"
            return f"{{{col}:{sign}{vn(abs(x))} {'pp' if pct else unit}}}"
        rows = [["Chỉ số (validationData)", "Version 1", "Version 2", "Thay đổi"]]
        for lab, k in (("Sai số trung bình (m thực địa)", "val_mean_ground_m"),
                       ("Sai số trung vị (m thực địa)", "val_median_ground_m"),
                       ("Phân vị 90 (m thực địa)", "val_p90_ground_m"),
                       ("Phân vị 95 (m thực địa)", "val_p95_ground_m"),
                       ("Sai số lớn nhất (m thực địa)", "val_max_ground_m"),
                       ("Sai số TB, m toạ độ chiếu", "val_mean_projected_m")):
            rows.append([lab, vn(s0[k]), vn(sel[k]), delta(s0[k], sel[k])])
        for lab, k in (("Tỉ lệ fingerprint sai ≤ 8 m", "val_share_within_8m"),
                       ("Độ chính xác tòa nhà", "val_building_accuracy"),
                       ("Độ chính xác tầng", "val_floor_accuracy"),
                       ("Độ chính xác joint (tòa nhà + tầng)",
                        "val_joint_building_floor_accuracy")):
            rows.append([lab, pc(s0[k]), pc(sel[k]), delta(s0[k], sel[k], "pp", True)])
        hu = d.CN.set_index("condition").loc["huim_union__C_rare_discriminative__1.7321pct"]
        rows.append([f"HUIM union C ({int(hu['n_waps'])} WAP): sai số TB (m)",
                     vn(hu["study_mean_ground_m"]),
                     vn(hu["opt_mean_ground_m"]),
                     delta(hu["study_mean_ground_m"], hu["opt_mean_ground_m"])])
        table(s, 46, 118, [610, 230, 230, 278], rows, size=18, row_h=50)
        tb(s, 46, 118 + 50 * len(rows) + 12, 1350, 40, f"Cùng {vi(len(d.D.Pnew))} "
           f"fingerprint kiểm định; V1 = k-NN k = 1 ({s0['stage']}), V2 = mô hình "
           f"{sel['stage']} do giao thức chọn trên dev. {{G:Xanh}} = tốt hơn.", size=16,
           color=MUTED)

    # ------------------------------------------- 32 V1 vs V2 method
    def s_cmp_method(self):
        d = self.d
        cfg = d.PO["selected_config"]
        s = self.slide("plain", "So sánh V1–V2: phương pháp")
        title(s, "So sánh {g:Version 1} và {g:Version 2}: phương pháp & trình bày",
              size=46)
        rows = [["Hạng mục", "Version 1", "Version 2"],
                ["Bộ ước lượng toạ độ", "k-NN k = 1, Manhattan, biểu diễn positive, "
                 "trọng số đều", f"{cfg['representation']} + Bray-Curtis, k = {cfg['k']}, "
                 f"w = 1/d{sup(cfg['power'])}, dò lệch ±{cfg['offset_range']} dB, "
                 f"ensemble với {cfg['ensemble_with']}"],
                ["Cách chọn tham số", "Chia ngẫu nhiên 75/25 trainingData; tiêu chí độ "
                 "chính xác phân loại", f"GroupKFold({d.n_folds}) theo USERID; tiêu chí "
                 f"sai số TB theo mét"],
                ["Phần HUIM (EFIM, 3 mô hình p)", "Như nhau", "Như nhau (không chạy lại)"],
                ["Kiểm chứng", "Unit test EFIM so với oracle vét cạn",
                 f"+ {d.vchk['huim_checks.json']['n_checks']} phép tính lại HUIM, "
                 f"{d.vchk['positioning_checks.json']['n_checks']} phép tính lại định vị, "
                 f"Excel tính tay"],
                ["Công thức", "Không ghi nguồn; Model B ghi nhầm là utility nội",
                 "Mỗi công thức ghi tác giả / nguồn có DOI đã kiểm tra"],
                ["Biểu đồ, bảng", "Ảnh chụp từ Word/PowerPoint, font khác nhau",
                 "Vẽ lại từ results/, in minUtil, mỗi biểu đồ kèm CSV"],
                ["Định dạng số", "Lẫn lộn (87.76% / 19.861 / 19,861)",
                 "Thống nhất kiểu Việt Nam: 1.234,56"],
                ["Số trang", f"{d.v1_pages} trang", "…"]]
        gf = table(s, 46, 112, [290, 500, 560], rows, size=17, row_h=68,
                   align=["l", "l", "l"])
        self._pages_cell = (gf, len(rows) - 1, 2)

    # ------------------------------------------- 33 corrections
    def fix_rows(self):
        d = self.d
        db = d.hsum["database"]
        return [["Slide V1", "Version 1 ghi", "Version 2 sửa thành"],
                ["8", "Model B: “internal utility đo bằng mutual information”",
                 "MI là {g:external} utility p(i); q(i,T) dùng chung cho cả 3 mô hình"],
                ["14", "Sắp xếp mục theo TWU {R:giảm} dần",
                 "Sắp theo TWU {g:tăng} dần (slide 15 của V1 đã ghi đúng)"],
                ["14", "“Xây dựng cây EFIM: quét CSDL 1 lần để tạo cây”",
                 "EFIM không dựng cây dữ liệu; nó duyệt cây liệt kê bằng DFS và chiếu CSDL"],
                ["14", "Xuất tập có utility {R:>} minUtil", "u(X) {g:≥} minUtil"],
                ["16", f"“{vi(db['n_non_empty_transactions'])} bản ghi kiểm luyện”; “lọc "
                 "nhiễu, làm đầy RSSI”", f"{vi(db['n_rows'])} fingerprint huấn luyện "
                 f"({vi(db['n_non_empty_transactions'])} giao dịch không rỗng); không làm "
                 "đầy giá trị nào"],
                ["6", f"“EFIM chạy trên {vi(db['n_rows'])} fingerprint”",
                 f"{vi(db['n_non_empty_transactions'])} giao dịch ({vi(db['n_empty_fingerprints_dropped'])} "
                 "fingerprint rỗng bị bỏ)"],
                ["11", "“Dòng CSV (dòng 3248)”", f"Dòng {vi(d.ex_line)} của file CSV = chỉ "
                 f"số pandas {vi(d.ex_index)}"],
                ["3", "“HUIM … nâng cao độ chính xác và ổn định của IPS”",
                 "Thực nghiệm cho thấy HUIM KHÔNG thắng đối chứng; nêu như câu hỏi nghiên cứu"]]

    def s_fixes(self):
        s = self.slide("plain", "Đính chính slide V1")
        title(s, "Đính chính nội dung slide {g:Version 1}", size=52)
        rows = self.fix_rows()
        table(s, 46, 112, [130, 560, 660], rows, size=18, row_h=70,
              align=["c", "l", "l"])
        tb(s, 46, 112 + 70 * len(rows) + 14, 1350, 40, f"Số slide theo bản Version 1 "
           f"({self.d.v1_pages} trang).", size=15, color=MUTED)

    # ------------------------------------------- 34 verification
    def s_verify(self):
        d = self.d
        s = self.slide("plain", "Kiểm chứng độc lập")
        title(s, "Kiểm chứng {g:độc lập}", size=58)
        tb(s, 48, 116, 1300, 30, "Ba lớp kiểm tra không dựa vào mã thí nghiệm — "
           "người đọc tự chạy lại được", size=20, color=MUTED)
        hc, pcs = d.vchk["huim_checks.json"], d.vchk["positioning_checks.json"]
        cards_ = [("HUIM", f"{hc['n_checks'] - hc['n_failed']}/{hc['n_checks']}",
                   "phép kiểm tra đạt", "Tính lại q, p, u(X) từ CSV gốc; kiểm tra mọi HUI "
                   "của ba lần khai phá chính; vét cạn mọi tổ hợp 1–3 WAP.",
                   "python src/verify_huim.py"),
                  ("Định vị", f"{pcs['n_checks'] - pcs['n_failed']}/{pcs['n_checks']}",
                   "phép kiểm tra đạt", "Tính lại độ chính xác và sai số chỉ bằng thư viện "
                   "chuẩn của Python; tìm lại láng giềng gần nhất trên toàn bộ trainingData.",
                   "python src/verify_positioning.py --brute 25")]
        if d.excel:
            cards_.append(("Excel", vi(d.excel), "ô kiểm tra KHỚP",
                           "Công thức Excel tính lại minUtil, utility của HUI và mọi dự "
                           "đoán của cả hai phiên bản.",
                           "verification/kiem_chung_thu_cong.xlsx"))
        w = (1348 - 32 * (len(cards_) - 1)) / len(cards_)
        for i, (h, v, sub, t, cmd) in enumerate(cards_):
            x = 46 + i * (w + 32)
            card(s, x, 176, w, 470, radius=0.05)
            tb(s, x + 26, 196, w - 52, 30, h.upper(), size=21, color=GOLD, bold=True)
            tb(s, x + 26, 240, w - 52, 80, f"{{g:{v}}}", size=58)
            tb(s, x + 26, 326, w - 52, 30, sub, size=18, color=SOFT)
            tb(s, x + 26, 380, w - 52, 220, t, size=18.5, spacing=1.1)
            tb(s, x + 26, 590, w - 52, 50, cmd, size=14.5, color=MUTED, font="Consolas")
        tb(s, 46, 676, 1350, 90, "Mọi con số trên các slide được src/make_deck_v2.py đọc "
           "từ results/ lúc dựng slide (không gõ tay); mỗi biểu đồ có file CSV chứa đúng "
           "các dòng đã vẽ trong figures/deck_v2/.", size=18, color=SOFT, spacing=1.1)

    # ------------------------------------------- 35 conclusions
    def s_conclusion(self):
        d, R = self.d, self.R
        db = d.hsum["database"]
        U = d.U
        s0, sel, BS = d.s0, d.sel, d.BS
        jac = d.O3["pairwise_overlap"]["A_signal vs B_discriminative"]["jaccard"]
        s = self.slide("plain", "Kết luận")
        tb(s, 40, 22, 600, 100, "Kết luận", size=76, bold=True, color=GOLD, shadow=True)
        tb(s, 44, 128, 1350, 640, [
            f"{{g:RQ1 — Mô hình hoá được:}} {vi(db['n_rows'])} fingerprint → "
            f"{vi(db['n_non_empty_transactions'])} utility transaction; EFIM chạy {len(d.H)} "
            f"cấu hình ngưỡng và khớp oracle vét cạn. “Hợp lý” là hợp lý CỦA ĐỊNH NGHĨA "
            f"utility, không phải tính chất của dữ liệu.",
            f"{{g:RQ2 — HUIM chưa thắng đối chứng:}} HUIM union thua top-n theo MI ở cả ba "
            f"mô hình ({' / '.join(pc(v) for v in U['huim_union'])} so với "
            f"{' / '.join(pc(v) for v in U['top_mi'])}); trần tối đa "
            f"{int(U['n_waps'].max())} WAP.",
            f"{{g:RQ3 — Chi phí:}} số HUI và bộ nhớ khai phá tăng vọt khi hạ minUtil, số WAP "
            f"gần như đứng yên; thời gian và bộ nhớ định vị giảm gần tuyến tính theo số WAP.",
            f"{{g:RQ4 — Version 2:}} sai số trung bình {vn(s0['val_mean_ground_m'])} → "
            f"{vn(sel['val_mean_ground_m'])} m (giảm {vn(BS['diff'])} m, KTC 95% "
            f"{vn(BS['diff_ci'][0])}–{vn(BS['diff_ci'][1])}), joint "
            f"{pc(s0['val_joint_building_floor_accuracy'])} → "
            f"{pc(sel['val_joint_building_floor_accuracy'])}, không dùng tập kiểm định để chọn.",
            f"Model B chọn tập WAP khác Model A (Jaccard {vn(jac, 3)} "
            f"{R.cite('jaccard1912')}), nhưng MI gần trùng tần suất phát hiện (ρ = "
            f"{vn(d.O6['spearman_support_vs_mi_over_all_detected_waps'], 4)}).",
            "{g:Đóng góp đúng mức:} khảo sát khả năng thích nghi HUIM cho bài toán chọn "
            "nhóm AP — có đối chứng, báo cáo cả kết quả âm, không tuyên bố tính mới."],
           size=21, bullet="–", gap=16, spacing=1.1)

    # ------------------------------------------- 36 future work
    def s_future(self):
        d = self.d
        s0, sel = d.s0, d.sel
        s = self.slide("stairs", "Hướng nghiên cứu")
        tb(s, 57, 58, 1000, 190, ["Kết luận và", "{g:Hướng nghiên cứu tiếp theo}"],
           size=72, bold=True, spacing=0.98, shadow=True)
        tb(s, 56, 296, 760, 480, [
            "HUIM mô hình hoá được vân tay Wi-Fi thành dữ liệu giao dịch để chọn nhóm AP.",
            "Nhưng nhóm AP do HUIM chọn chưa tốt hơn tiêu chí đơn giản (top-n theo MI).",
            f"Sai số định vị giảm {vn(s0['val_mean_ground_m'])} → "
            f"{vn(sel['val_mean_ground_m'])} m nhờ tối ưu bộ ước lượng toạ độ (V2), không "
            f"nhờ HUIM.",
            "Hướng tiếp: utility ngoài độc lập với tần suất phát hiện; top-k hoặc closed "
            "HUIM để lấy nhiều WAP hơn với cùng ngân sách; thêm một tập kiểm tra chưa từng "
            "dùng."], size=21, bullet="–", gap=14, spacing=1.1)

    # ------------------------------------------- 37-38 references
    def s_refs(self):
        R = self.R
        order = list(R.order)
        half = (len(order) + 1) // 2
        for k, chunk in enumerate((order[:half], order[half:])):
            s = self.slide("plain", f"Tài liệu tham khảo {k + 1}")
            title(s, f"Tài liệu {{g:tham khảo}} ({k + 1}/2)", size=50)
            paras = []
            for key in chunk:
                mark = "‡" if R.kind[key] == "new" else "†"
                paras.append(f"[{R.order.index(key) + 1}] {R.entry[key]} {mark}")
            tb(s, 46, 112, 1350, 620, paras, size=17, gap=11, spacing=1.05)
            tb(s, 46, 744, 1350, 40,
               f"† danh mục đã kiểm tra DOI qua Crossref của dự án (research/papers_table.csv)"
               f"   ‡ kiểm tra DOI qua Crossref ngày "
               f"{R.checked[8:10]}/{R.checked[5:7]}/{R.checked[:4]} "
               f"(research/doi_verification_formulas.json)", size=13, color=MUTED)

    def build(self):
        for f in (self.s_cover, self.s_new, self.s_what, self.s_ips, self.s_defs,
                  self.s_model, self.s_pdef, self.s_prep, self.s_step1, self.s_step2,
                  self.s_step3, self.s_mapping, self.s_efim, self.s_bounds,
                  self.s_inside, self.s_exp, self.s_minutil, self.s_cost_p,
                  self.s_runtime, self.s_controls, self.s_featred, self.s_ipscost,
                  self.s_metres, self.s_divider, self.s_why, self.s_protocol,
                  self.s_v2_formulas, self.s_stages, self.s_cdf, self.s_sets,
                  self.s_cmp_numbers, self.s_cmp_method, self.s_fixes, self.s_verify,
                  self.s_conclusion, self.s_future, self.s_refs):
            f()
            LOG.info("slide %2d  %s", len(self.slides), self.slides[-1][0])
        self.page_numbers()
        gf, r, c = self._pages_cell
        cell = gf.table.cell(r, c)
        p = cell.text_frame.paragraphs[0]
        for run in list(p.runs):
            p._p.remove(run._r)
        add_rich(p, f"{len(self.slides)} trang (thêm phần Version 2, so sánh, đính "
                    f"chính, tài liệu)", 17, WHITE)
        cp = self.prs.core_properties
        cp.title = "HUIM và IPS — Version 2"
        cp.author = self.d.v1_author or "HUIM-IPS"
        cp.subject = "High Utility Itemset Mining cho định vị trong nhà (UJIIndoorLoc)"
        return self.prs


# ---------------------------------------------------------------------- PDF
def export_pdf(pptx: Path, pdf: Path, png_dir: Path | None = None) -> None:
    """PowerPoint (COM) saves the PDF; optionally one PNG per slide for QA."""
    png = f"$p.Slides | ForEach-Object {{ $_.Export((Join-Path '{png_dir}' " \
          f"('s{{0:D2}}.png' -f $_.SlideIndex)), 'PNG', 1440, 810) }};" if png_dir else ""
    ps = (f"$app = New-Object -ComObject PowerPoint.Application; "
          f"$p = $app.Presentations.Open('{pptx}', $true, $false, $false); "
          f"$p.SaveAs('{pdf}', 32); {png} $p.Close(); $app.Quit()")
    if png_dir:
        png_dir.mkdir(parents=True, exist_ok=True)
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                       capture_output=True, text=True)
    if r.returncode != 0 or not pdf.exists():
        raise RuntimeError(f"PowerPoint export failed: {r.stderr.strip()}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--no-pdf", action="store_true", help="write the PPTX only")
    ap.add_argument("--png", type=Path, help="also export one PNG per slide here")
    a = ap.parse_args()
    if not V1_PDF.exists():
        raise SystemExit(f"Version 1 deck not found: {V1_PDF}")
    if not font_available():
        LOG.warning("Public Sans is not installed; PowerPoint will substitute a font")
    chart_style()
    bg = templates(V1_PDF, ASSETS, TEMPLATES)
    d = Deck()
    prs = Builder(d, bg).build()
    tmp = OUT_PPTX.with_suffix(".__tmp__.pptx")
    prs.save(tmp)
    tmp.replace(OUT_PPTX)
    LOG.info("wrote %s (%d slides)", OUT_PPTX, len(prs.slides))
    if not a.no_pdf:
        export_pdf(OUT_PPTX, OUT_PDF, a.png)
        LOG.info("wrote %s", OUT_PDF)


if __name__ == "__main__":
    main()
