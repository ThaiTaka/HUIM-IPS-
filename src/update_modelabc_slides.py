"""Refresh the two native charts in slides/HUIM_ModelABC_4slides.pptx from
results/, and label every HUIM bar with the minUtil threshold that produced it.

The deck was laid out by hand; this script does not rebuild it. It opens the
file, replaces the data behind the two charts (categories, values, embedded
worksheet) and leaves every shape, colour and font as it was. Series
formatting survives because python-pptx's replace_data() rewrites only the
<c:tx>, <c:cat> and <c:val> children of each series.

Why the thresholds have to be on the chart: minUtil is a percentage of the
model's total database utility, and that total differs between the models
(A 9,058,242; B 4,354,299; C 6,211,852). "minUtil = 2%" is therefore three
different absolute thresholds. Each category label carries both forms, so a
reader can check  minUtil% x total utility = absolute minUtil  against
results/metrics/huim_runs.csv, and from there that every HUI behind a bar has
u(X) >= that absolute value (see src/make_verification_kit.py).

Slide 3  runtime of EFIM at minUtil = 2% for each model
Slide 4  joint accuracy of the HUIM union vs the matched-budget controls

Run: python src/update_modelabc_slides.py
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData

sys.path.insert(0, str(Path(__file__).resolve().parent))
from utils import METRICS, SLIDES, get_logger

LOG = get_logger("abcslides", "slides.log")

DECK = SLIDES / "HUIM_ModelABC_4slides.pptx"
MODELS = ["A_signal", "B_discriminative", "C_rare_discriminative"]
LETTER = {"A_signal": "A", "B_discriminative": "B",
          "C_rare_discriminative": "C"}
RUNTIME_PCT = 2.0           # slide 3 compares the models at one threshold


def vn_num(x: float, d: int = 0) -> str:
    """Vietnamese number format used on the deck: 1.234,5"""
    s = f"{x:,.{d}f}"
    return s.replace(",", "\0").replace(".", ",").replace("\0", ".")


def vn_pct(p: float) -> str:
    return f"{p:g}".replace(".", ",") + "%"


def chart_on(slide):
    charts = [sh for sh in slide.shapes if getattr(sh, "has_chart", False)
              and sh.has_chart]
    if len(charts) != 1:
        raise RuntimeError(f"expected one chart on the slide, found {len(charts)}")
    return charts[0].chart


def runtime_chart(H: pd.DataFrame) -> tuple[CategoryChartData, list[str]]:
    cats, vals, notes = [], [], []
    for m in MODELS:
        r = H[(H["utility_model"] == m)
              & np.isclose(H["rel_minutil_pct"], RUNTIME_PCT)]
        if len(r) != 1:
            raise RuntimeError(f"no unique {m} run at {RUNTIME_PCT}%")
        r = r.iloc[0]
        head = "A (p=1)" if m == "A_signal" else LETTER[m]
        cats.append(f"{head}\nminUtil {vn_pct(RUNTIME_PCT)} = "
                    f"{vn_num(r['abs_minutil'])}")
        vals.append(round(float(r["runtime_s"]), 1))
        notes.append(f"{LETTER[m]}: {vn_pct(RUNTIME_PCT)} x "
                     f"{vn_num(r['total_db_utility'])} = "
                     f"{vn_num(r['abs_minutil'])} -> {vn_num(r['n_huis'])} HUI, "
                     f"{float(r['runtime_s']):.1f} s ({r['run_tag']})")
    cd = CategoryChartData(number_format="0")
    cd.categories = cats
    cd.add_series("Thời gian (s)", vals)
    return cd, notes


def accuracy_chart(H: pd.DataFrame, I: pd.DataFrame
                   ) -> tuple[CategoryChartData, list[str]]:
    """HUIM union at the lowest COMPLETED minUtil of each model, against the
    controls at exactly the same number of WAPs."""
    cats, notes = [], []
    series = {"HUIM union": [], "Top theo MI": [], "Top theo tần suất": [],
              "Ngẫu nhiên (TB 3 seed)": []}
    for m in MODELS:
        done = H[(H["utility_model"] == m) & (H["status"] == "COMPLETED")
                 & (H["n_huis"] > 0)]
        run = done.loc[done["rel_minutil_pct"].idxmin()]
        cond = f"huim_union__{run['run_tag']}"
        u = I[I["condition"] == cond]
        if len(u) != 1:
            raise RuntimeError(f"{cond} missing from ips_results.csv")
        u = u.iloc[0]
        n = int(u["n_waps_selected"])
        mi = I[I["condition"] == f"control_top_mi__n{n}"]
        sp = I[I["condition"] == f"control_top_support__n{n}"]
        rd = I[(I["family"] == "control_random") & (I["budget_n"] == n)]
        if not (len(mi) == 1 and len(sp) == 1 and len(rd) == 3):
            raise RuntimeError(f"controls at n={n} incomplete")
        series["HUIM union"].append(round(u["joint_building_floor_accuracy"] * 100, 2))
        series["Top theo MI"].append(round(mi.iloc[0]["joint_building_floor_accuracy"] * 100, 2))
        series["Top theo tần suất"].append(round(sp.iloc[0]["joint_building_floor_accuracy"] * 100, 2))
        series["Ngẫu nhiên (TB 3 seed)"].append(round(rd["joint_building_floor_accuracy"].mean() * 100, 2))
        cats.append(f"{LETTER[m]} ({n} WAP)\nminUtil {vn_pct(run['rel_minutil_pct'])}"
                    f" = {vn_num(run['abs_minutil'])}")
        notes.append(f"{LETTER[m]}: {vn_pct(run['rel_minutil_pct'])} x "
                     f"{vn_num(run['total_db_utility'])} = "
                     f"{vn_num(run['abs_minutil'])}; {vn_num(run['n_huis'])} HUI "
                     f"-> {n} WAP ({cond})")
    cd = CategoryChartData(number_format="0.0")
    cd.categories = cats
    for name, vals in series.items():
        cd.add_series(name, vals)
    return cd, notes


def set_note(slide, marker: str, lines: list[str]) -> None:
    """Append (or refresh) a block of threshold lines in the speaker notes."""
    tf = slide.notes_slide.notes_text_frame
    text = tf.text.split("\n" + marker)[0].rstrip()
    tf.text = text + "\n" + marker + "\n" + "\n".join(lines)


def main() -> None:
    H = pd.read_csv(METRICS / "huim_runs.csv")
    I = pd.read_csv(METRICS / "ips_results.csv")
    prs = Presentation(str(DECK))
    s3, s4 = prs.slides[2], prs.slides[3]

    cd, n3 = runtime_chart(H)
    chart_on(s3).replace_data(cd)
    set_note(s3, "Ngưỡng minUtil (minUtil% x tổng utility CSDL = minUtil tuyệt đối):", n3)

    cd, n4 = accuracy_chart(H, I)
    chart_on(s4).replace_data(cd)
    set_note(s4, "Ngưỡng minUtil (minUtil% x tổng utility CSDL = minUtil tuyệt đối):", n4)

    tmp = DECK.with_name(DECK.stem + ".__tmp__.pptx")
    prs.save(str(tmp))
    os.replace(tmp, DECK)
    for ln in n3 + n4:
        LOG.info("%s", ln)
    LOG.info("updated %s", DECK)


if __name__ == "__main__":
    main()
