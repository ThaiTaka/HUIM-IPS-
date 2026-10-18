"""BAOCAO/so_lieu/kiem_chung_mean.xlsx and kiem_tra_mean.py.

The workbook recomputes every reported mean error with Excel formulas only:
    err = SQRT((pred_x - true_x)^2 + (pred_y - true_y)^2) * ground factor
    mean = AVERAGE(err)
next to the value the pipeline reported, and one worked example rebuilds a
predicted position from its cells: x = SUMPRODUCT(weight, cell x).

Run after src/baocao/export_data.py.
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "BAOCAO" / "so_lieu"
HEAD = Font(name="Times New Roman", bold=True)
BODY = Font(name="Times New Roman")
BLUE = Font(name="Times New Roman", color="1F4E9A")
GREEN = PatternFill("solid", fgColor="E2EFDA")
GREY = PatternFill("solid", fgColor="EDEDED")

KIT_PY = '''"""Tính lại sai số trung bình từ các file du_doan/*.csv, chỉ dùng thư viện chuẩn.

    python kiem_tra_mean.py

Với mỗi tập test: sai số từng mẫu = sqrt((x_dự đoán - x_thật)^2 + (y_dự đoán - y_thật)^2)
nhân hệ số đổi ra mét thực địa; trung bình cộng được so với 01_ket_qua_tong_hop.csv.
"""
import csv
import math
from pathlib import Path

HERE = Path(__file__).resolve().parent
tong_hop = {(r["dataset"], r["tap_test"]): r
            for r in csv.DictReader(open(HERE / "01_ket_qua_tong_hop.csv", encoding="utf-8"))}
loi = 0
for (ds, tap), r in tong_hop.items():
    he_so = float(r["he_so_met"])
    sai_so, dung = [], 0
    for d in csv.DictReader(open(HERE / "du_doan" / f"{ds}__{tap}.csv", encoding="utf-8")):
        e = math.hypot(float(d["pred_x"]) - float(d["true_x"]),
                       float(d["pred_y"]) - float(d["true_y"])) * he_so
        sai_so.append(e)
        dung += f"{d['true_building']}_{d['true_floor']}" == d["pred_building_floor"]
    tb = sum(sai_so) / len(sai_so)
    s = sorted(sai_so)
    n = len(s)
    tv = s[n // 2] if n % 2 else (s[n // 2 - 1] + s[n // 2]) / 2
    khop = (abs(round(tb, 4) - float(r["sai_so_tb_m"])) < 1e-4
            and abs(round(tv, 4) - float(r["trung_vi_m"])) < 1e-4
            and abs(round(dung / n, 6) - float(r["dung_toa_tang"])) < 1e-5)
    loi += not khop
    print(f"{r['ten']:<22} {tap:<15} n={n:<5} TB tính lại {tb:8.4f} m | báo cáo {float(r['sai_so_tb_m']):8.4f} m"
          f" | trung vị {tv:7.4f} m | đúng tòa+tầng {dung / n:.4f} | {'KHỚP' if khop else 'LỆCH'}")
print("Tất cả khớp." if not loi else f"{loi} tập test lệch.")
'''


def sheet_name(ds: str, t: str) -> str:
    return f"{ds}_{t}"[:31]


def main() -> None:
    R = pd.read_csv(OUT / "01_ket_qua_tong_hop.csv")
    data = json.load(open(OUT / "du_lieu_bao_cao.json", encoding="utf-8"))
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Tong_hop"
    ws0.append(["Dataset", "Tập test", "Số mẫu", "Hệ số mét", "TB tính bằng công thức (m)",
                "TB báo cáo (m)", "Chênh lệch (m)", "Đúng tòa+tầng (công thức)",
                "Đúng tòa+tầng báo cáo", "Kết luận"])
    for i, r in enumerate(R.itertuples(), start=2):
        sn = sheet_name(r.dataset, r.tap_test)
        P = pd.read_csv(OUT / "du_doan" / f"{r.dataset}__{r.tap_test}.csv")
        ws = wb.create_sheet(sn)
        ws.append(["Dòng nguồn", "x thật", "y thật", "x dự đoán", "y dự đoán",
                   "Sai số (m) = công thức", "Tòa_tầng thật", "Tòa_tầng dự đoán", "Đúng (1/0)",
                   "", "Hệ số mét", r.he_so_met])
        n = len(P)
        for j, d in enumerate(P.itertuples(), start=2):
            ws.append([int(d.src_row), d.true_x, d.true_y, d.pred_x, d.pred_y,
                       f"=SQRT((D{j}-B{j})^2+(E{j}-C{j})^2)*$L$1",
                       f"{d.true_building}_{d.true_floor}", d.pred_building_floor,
                       f"=IF(G{j}=H{j},1,0)"])
        last = n + 1
        ws["K3"], ws["L3"] = "TB sai số (m)", f"=AVERAGE(F2:F{last})"
        ws["K4"], ws["L4"] = "Trung vị (m)", f"=MEDIAN(F2:F{last})"
        ws["K5"], ws["L5"] = "Đúng tòa+tầng", f"=AVERAGE(I2:I{last})"
        ws["K6"], ws["L6"] = "TB báo cáo (m)", r.sai_so_tb_m
        ws["K7"], ws["L7"] = "Chênh lệch (m)", "=L3-L6"
        for c in ("L3", "L4", "L5", "L7"):
            ws[c].fill = GREEN
        ws["L1"].font = ws["L6"].font = BLUE
        for c in range(1, 13):
            ws.cell(1, c).font = HEAD
            ws.column_dimensions[get_column_letter(c)].width = 16
        ws.freeze_panes = "A2"
        ws0.append([r.ten, r.tap_test, r.so_mau, r.he_so_met, f"='{sn}'!L3", r.sai_so_tb_m,
                    f"=E{i}-F{i}", f"='{sn}'!L5", r.dung_toa_tang,
                    f'=IF(ABS(G{i})<0.00005,"KHỚP","LỆCH")'])
        for c in ("E", "G", "H", "J"):
            ws0[f"{c}{i}"].fill = GREEN

    ex = data["worked_example"]
    gf = data["datasets"]["uji"]["tests"]["validationData"]["he_so_met"]
    ws = wb.create_sheet("Vi_du_UJI_mot_truy_van")
    ws.append([f"Truy vấn: dòng {ex['src_row']} của validationData (UJIIndoorLoc). "
               "Toạ độ dự đoán = tổng có trọng số của tâm các ô HUIM chọn."])
    ws.append([])
    ws.append(["Ô (chỉ số)", "x tâm ô", "y tâm ô", "Trọng số a_c"])
    for c in ex["cells"]:
        ws.append([c["index"], c["x"], c["y"], c["weight"]])
    k = 4 + len(ex["cells"]) - 1
    rows = [("Tổng trọng số", f"=SUM(D4:D{k})", None),
            ("x dự đoán = SUMPRODUCT(a, x)", f"=SUMPRODUCT(D4:D{k},B4:B{k})", ex["pred_x"]),
            ("y dự đoán = SUMPRODUCT(a, y)", f"=SUMPRODUCT(D4:D{k},C4:C{k})", ex["pred_y"]),
            ("x thật", ex["true_x"], None), ("y thật", ex["true_y"], None),
            ("Hệ số mét", gf, None)]
    ws.append([])
    ws.append(["Đại lượng", "Công thức", "Giá trị pipeline"])
    base = k + 2
    for i, (a, b, c) in enumerate(rows):
        ws.append([a, b, c])
    r_px, r_py, r_tx, r_ty, r_gf = base + 2, base + 3, base + 4, base + 5, base + 6
    ws.append(["Sai số (m)", f"=SQRT((B{r_px}-B{r_tx})^2+(B{r_py}-B{r_ty})^2)*B{r_gf}", ex["err_m"]])
    for row in ws.iter_rows():
        for c in row:
            if c.font != HEAD:
                c.font = BODY
    for c in ("A3", "B3", "C3", "D3", f"A{base}", f"B{base}", f"C{base}"):
        ws[c].font = HEAD
    for col, w in zip("ABCD", (34, 30, 22, 16)):
        ws.column_dimensions[col].width = w

    for c in range(1, 11):
        ws0.cell(1, c).font = HEAD
        ws0.cell(1, c).alignment = Alignment(wrap_text=True, vertical="top")
        ws0.column_dimensions[get_column_letter(c)].width = 18
    ws0.row_dimensions[1].height = 45
    note = wb.create_sheet("Huong_dan", 0)
    for line in [
        "SỔ KIỂM CHỨNG SAI SỐ TRUNG BÌNH — ĐỊNH VỊ CHỈ BẰNG HUIM",
        "",
        "Mỗi sheet <dataset>_<tập test> chứa toàn bộ dự đoán của tập test đó, chép từ du_doan/*.csv.",
        "Cột F tính sai số từng mẫu bằng công thức: SQRT((x dự đoán - x thật)^2 + (y dự đoán - y thật)^2) x hệ số mét (ô L1).",
        "L3 = AVERAGE(F) là sai số trung bình tính lại; L6 là con số pipeline báo cáo; L7 là chênh lệch.",
        "Hệ số mét: UJIIndoorLoc dùng toạ độ Web Mercator (EPSG:3857), mét thực địa = mét chiếu x 0,766126; các dataset khác dùng mét trực tiếp (hệ số 1).",
        "Sheet Tong_hop gom kết quả mọi tập test; cột Kết luận ghi KHỚP khi chênh lệch dưới 0,00005 m.",
        "Sheet Vi_du_UJI_mot_truy_van dựng lại toạ độ dự đoán của một truy vấn từ các ô HUIM chọn bằng SUMPRODUCT.",
        "Ô màu xanh lá là kết quả công thức; chữ xanh dương là giá trị nhập từ file.",
    ]:
        note.append([line])
        note.cell(note.max_row, 1).font = HEAD if note.max_row == 1 else BODY
    note.column_dimensions["A"].width = 130
    wb.save(OUT / "kiem_chung_mean.xlsx")
    (OUT / "kiem_tra_mean.py").write_text(KIT_PY, encoding="utf-8")
    print("written kiem_chung_mean.xlsx, kiem_tra_mean.py")


if __name__ == "__main__":
    main()
